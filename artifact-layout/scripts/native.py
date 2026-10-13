"""Native format adapters for the shared layout model and geometry."""

import hashlib
import importlib.util
import io
import json
import re
import subprocess
from functools import lru_cache
from pathlib import Path

from primitives import DIRECTIONS, GRAPHICS, heat_color

from render import (
    bundle_image,
    chart_domain,
    chart_geometry,
    chart_svg,
    load_asset,
    numeric_cell,
)

# Keep the readable context with its visual mark. Prose and data tables still
# flow normally, while a figure/metric/quotation is one editorial unit.
NATIVE_UNITS = GRAPHICS | {
    "bar-chart",
    "line-chart",
    "big-number",
    "change-indicator",
    "pull-quote",
    "figure",
    "callout",
    "source-note",
}


def all_blocks(report):
    return [block for section in report["sections"] for block in section["blocks"]]


def prepare(report, office_dir, formats):
    """Fail before writing if required engines or base assets are unavailable."""
    needed = {"docx": ("docx", "blank.docx"), "pdf": ("typst", "blank.typ")}
    for name in set(formats) & needed.keys():
        package, asset = needed[name]
        if importlib.util.find_spec(package) is None:
            raise ImportError(
                f"Install {package} through office-files setup for {name}"
            )
        if not (office_dir / "assets" / asset).is_file():
            raise ValueError(
                f"Missing Office baseline asset: {asset}; provide --office-dir"
            )
    if "docx" in formats:
        for block in all_blocks(report):
            if block.get("image", "").lower().endswith(".svg"):
                raise ValueError(
                    "Word figures require PNG/JPEG; provide a raster copy of this SVG"
                )
        has_charts = any(
            block["type"] in {"bar-chart", "line-chart", "small-multiples"} | GRAPHICS
            for block in all_blocks(report)
        )
        if has_charts and importlib.util.find_spec("PIL") is None:
            raise ImportError(
                "Word chart images require Pillow; use office-files setup --charts"
            )


def literal(value):
    return "text(" + json.dumps(str(value), ensure_ascii=False) + ")"


@lru_cache(maxsize=8)
def font_reference(family):
    result = subprocess.run(
        ["fc-match", "--format=%{file}\n%{index}", family],
        capture_output=True,
        text=True,
        check=True,
    )
    file_name, face_index = result.stdout.strip().splitlines()
    path = Path(file_name)
    if not path.is_file():
        raise ValueError(f"No rendering font for {family}")
    return str(path), int(face_index)


def chart_png(block, theme, output, domain=None, label_size=30):
    """Draw the exact shared geometry at 2x resolution for editable Word."""
    from PIL import Image, ImageDraw, ImageFont

    geometry = chart_geometry(block, theme, domain, label_size)
    scale = 2
    image = Image.new(
        "RGB",
        (round(geometry["width"] * scale), round(geometry["height"] * scale)),
        theme["background"],
    )
    draw = ImageDraw.Draw(image)
    for x1, y1, x2, y2, color in geometry["lines"]:
        draw.line((x1 * scale, y1 * scale, x2 * scale, y2 * scale), fill=color, width=2)
    for x, y, width, height, color in geometry["rects"]:
        draw.rectangle(
            (x * scale, y * scale, (x + width) * scale, (y + height) * scale),
            fill=color,
        )
    points = [(x * scale, y * scale) for x, y in geometry["points"]]
    if len(points) > 1:
        draw.line(points, fill=theme["accent"], width=8)
    for x, y in points:
        draw.ellipse((x - 8, y - 8, x + 8, y + 8), fill=theme["accent"])
    anchors = {"start": "ls", "end": "rs", "middle": "ms"}
    for x, y, value, anchor, color, size in geometry["texts"]:
        family = (
            theme["cjk"] if re.search(r"[\u2e80-\u9fff]", str(value)) else theme["body"]
        )
        file_name, face_index = font_reference(family)
        font = ImageFont.truetype(file_name, round(size * scale), index=face_index)
        draw.text(
            (x * scale, y * scale),
            str(value),
            fill=color,
            font=font,
            anchor=anchors[anchor],
        )
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    data = buffer.getvalue()
    path = (
        output / "assets" / ("chart-" + hashlib.sha256(data).hexdigest()[:20] + ".png")
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def render_word(report, output, source_dir, office_dir):
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    theme = load_asset("themes.json")[report["theme"]]
    recipe = load_asset("recipes.json")[report["recipe"]]
    document = Document(office_dir / "assets" / "blank.docx")
    for paragraph in list(document.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)
    for style in document.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            style.font.name = theme["body"]
            style.font.color.rgb = RGBColor.from_string(theme["ink"][1:])
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), theme["cjk"])
            for attribute in list(fonts.attrib):
                if attribute.endswith(("Theme", "theme")):
                    del fonts.attrib[attribute]
    for name, size, color in [
        ("Layout Metric", 32, theme["accent"]),
        ("Layout Quote", 18, theme["ink"]),
        ("Layout Meta", 10, theme["muted"]),
    ]:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color[1:])
        style.paragraph_format.line_spacing = Pt(size + 5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = name == "Layout Metric"
    document.styles["Title"].font.name = theme["display"]
    background = OxmlElement("w:background")
    background.set(qn("w:color"), theme["background"][1:])
    document.element.insert(0, background)
    available = (210 - 36) / 25.4 * 72

    def configure(section, count=1):
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(18)
        section.top_margin = section.bottom_margin = Mm(18)
        cols = section._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            section._sectPr.append(cols)
        cols.set(qn("w:num"), str(count))
        cols.set(qn("w:space"), "360")

    configure(document.sections[-1])
    current_columns = 1

    def columns(count):
        nonlocal current_columns
        if count != current_columns:
            configure(document.add_section(WD_SECTION_START.CONTINUOUS), count)
            current_columns = count

    def paragraph(container, value, style="Normal"):
        return container.add_paragraph(str(value), style)

    def ordered_sequence():
        numbering = document.part.numbering_part.element
        base_id = document.styles["List Number"].element.pPr.numPr.numId.val
        abstract = numbering.xpath(f'./w:num[@w:numId="{base_id}"]/w:abstractNumId')[0]
        instance = numbering.add_num(int(abstract.get(qn("w:val"))))
        override = instance.add_lvlOverride(0)
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        return instance.numId

    def ordered_paragraph(container, value, sequence):
        p = paragraph(container, value, "List Number")
        properties = p._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_ilvl().val = 0
        properties.get_or_add_numId().val = sequence
        return p

    def note(container, block):
        for key, style in [("caption", "Caption"), ("source", "Source")]:
            if block.get(key):
                paragraph(container, block[key], style)

    def title(container, block):
        if block.get("title"):
            paragraph(container, block["title"], "Heading 2")

    def layout_grid(container, groups, width, proportions=None):
        proportions = proportions or [1] * len(groups)
        table = container.add_table(rows=1, cols=len(groups))
        table.autofit = False
        for index, (group, weight) in enumerate(zip(groups, proportions)):
            cell = table.cell(0, index)
            cell_width = width * weight / sum(proportions)
            cell.width = Pt(cell_width)
            table.columns[index].width = Pt(cell_width)
            for block in group:
                draw_block(cell, block, max(1, cell_width - 12))
            initial = cell.paragraphs[0]
            if not initial.text and not initial._p.xpath(".//w:drawing"):
                initial._p.getparent().remove(initial._p)
        paragraph(container, "")

    def metric(container, item):
        p = paragraph(container, item["value"], "Layout Metric")
        if item.get("unit"):
            run = p.add_run(" " + str(item["unit"]))
            run.font.size = Pt(13)
        p = paragraph(container, item["label"])
        p.runs[0].bold = True
        for key in ["delta", "period", "note"]:
            if item.get(key):
                paragraph(container, item[key], "Layout Meta")
        note(container, item)

    def image(container, path, width):
        p = paragraph(container, "", "Figure")
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width / 72))

    def draw_chart(container, block, width, domain=None):
        title(container, block)
        context = " · ".join(
            str(block[key]) for key in ["unit", "period"] if block.get(key)
        )
        if context:
            paragraph(container, context, "Caption")
        image(
            container,
            chart_png(block, theme, output, domain, 640 * 10.5 / width),
            width,
        )
        note(container, block)

    def draw_block(container, block, width):
        before = len(container.paragraphs)
        draw_block_content(container, block, width)
        if block["type"] in NATIVE_UNITS:
            group = container.paragraphs[before:]
            for index, p in enumerate(group):
                p.paragraph_format.keep_together = True
                p.paragraph_format.keep_with_next = index < len(group) - 1

    def draw_block_content(container, block, width):
        kind = block["type"]
        if kind == "heading":
            paragraph(container, block["text"], f"Heading {block.get('level', 2) - 1}")
            return
        if kind == "standfirst":
            p = paragraph(container, block["text"])
            p.runs[0].font.color.rgb = RGBColor.from_string(theme["muted"][1:])
            return
        if kind == "metric-sparkline":
            metric(container, {**block, "source": None})
            image(
                container,
                chart_png(block, theme, output, label_size=640 * 10.5 / width),
                width,
            )
            note(container, block)
            return
        if kind in {"target-comparison", "progress"}:
            title(container, block)
            p = paragraph(container, block["label"])
            p.runs[0].bold = True
            draw_chart(container, {**block, "title": None}, width)
            return
        if kind in {"bar-chart", "line-chart"} | GRAPHICS:
            draw_chart(container, block, width, block.get("_domain"))
            return
        if kind == "image-text":
            layout_grid(
                container,
                [
                    [{**block, "type": "figure", "title": None}],
                    [
                        {
                            "type": "prose",
                            "title": block.get("title"),
                            "paragraphs": block["paragraphs"],
                        }
                    ],
                ],
                width,
            )
            return
        if kind == "metric-strip":
            layout_grid(
                container,
                [[{**item, "type": "big-number"}] for item in block["items"]],
                width,
            )
            return
        title(container, block)
        if kind == "prose":
            for value in block["paragraphs"]:
                paragraph(container, value)
        elif kind == "big-number":
            metric(container, block)
        elif kind == "change-indicator":
            p = paragraph(
                container,
                DIRECTIONS[block["direction"]] + " " + block["value"],
                "Layout Quote",
            )
            p.runs[0].font.color.rgb = RGBColor.from_string(theme["accent"][1:])
            p = paragraph(container, block["label"])
            p.runs[0].bold = True
            paragraph(container, block["comparison"], "Layout Meta")
            note(container, block)
        elif kind == "pull-quote":
            paragraph(container, block["text"], "Layout Quote")
            if block.get("attribution"):
                paragraph(container, block["attribution"], "Caption")
            note(container, block)
        elif kind in {"callout", "source-note"}:
            p = paragraph(
                container,
                block["text"],
                "Source" if kind == "source-note" else "Normal",
            )
            if kind == "callout":
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), theme["tint"][1:])
                p._p.get_or_add_pPr().append(shading)
            note(container, block)
        elif kind in {"takeaways", "numbered-list"}:
            sequence = ordered_sequence() if kind == "numbered-list" else None
            for value in block["items"]:
                if kind == "numbered-list":
                    ordered_paragraph(container, value, sequence)
                else:
                    paragraph(container, value, "List Bullet")
            note(container, block)
        elif kind == "timeline":
            table = container.add_table(rows=0, cols=2)
            table.autofit = False
            for index, column in enumerate(table.columns):
                column.width = Pt(width * (0.25 if index == 0 else 0.75))
            for item in block["items"]:
                row = table.add_row()
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for index, value in enumerate([item["label"], item["text"]]):
                    cell = row.cells[index]
                    cell.width = Pt(width * (0.25 if index == 0 else 0.75))
                    p = cell.paragraphs[0]
                    p.style = "Normal"
                    p.text = value
                    if index == 0:
                        p.runs[0].bold = True
            note(container, block)
            paragraph(container, "")
        elif kind == "process-steps":
            sequence = ordered_sequence()
            for item in block["items"]:
                p = ordered_paragraph(container, item["label"], sequence)
                p.runs[0].bold = True
                p.paragraph_format.keep_with_next = True
                p = paragraph(container, item["text"])
                p.paragraph_format.left_indent = Pt(18)
            note(container, block)
        elif kind == "figure":
            name = bundle_image((source_dir / block["image"]).resolve(), output)
            image(container, output / name, width)
            note(container, block)
        elif kind == "small-multiples":
            domain = chart_domain(block)
            for offset in range(0, len(block["charts"]), 2):
                panels = block["charts"][offset : offset + 2]
                layout_grid(
                    container,
                    [[{**panel, "_domain": domain}] for panel in panels],
                    width,
                )
            note(container, block)
        elif kind in {"data-table", "comparison-table", "heatmap-table"}:
            table = container.add_table(rows=1, cols=len(block["headers"]))
            table.autofit = False
            for column in table.columns:
                column.width = Pt(width / len(block["headers"]))
            header = table.rows[0]
            header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for row_index, values in enumerate([block["headers"]] + block["rows"]):
                row = header if row_index == 0 else table.add_row()
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for index, (cell, value) in enumerate(zip(row.cells, values)):
                    cell.width = Pt(width / len(block["headers"]))
                    p = cell.paragraphs[0]
                    p.style = "Table Header" if row_index == 0 else "Table Text"
                    p.text = str(value)
                    properties = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for side in ["top", "bottom", "left", "right"]:
                        margin = OxmlElement("w:" + side)
                        margin.set(qn("w:w"), "100")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    properties.append(margins)
                    borders = OxmlElement("w:tcBorders")
                    border = OxmlElement("w:bottom")
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "8" if row_index == 0 else "4")
                    border.set(
                        qn("w:color"), theme["ink" if row_index == 0 else "rule"][1:]
                    )
                    borders.append(border)
                    properties.append(borders)
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), theme["tint"][1:])
                        properties.append(shading)
                    if kind == "heatmap-table" and row_index:
                        shading = OxmlElement("w:shd")
                        shading.set(
                            qn("w:fill"), heat_color(block, index, value, theme)[1:]
                        )
                        cell._tc.get_or_add_tcPr().append(shading)
                    p.paragraph_format.keep_with_next = row_index == 0
                    p.alignment = (
                        WD_ALIGN_PARAGRAPH.RIGHT
                        if row_index and numeric_cell(block, index, value)
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
            note(container, block)
            if kind == "heatmap-table":
                paragraph(
                    container,
                    f"{block.get('scaleLabel', 'Color scale')}: {block['domain'][0]} → {block['domain'][1]}",
                    "Caption",
                )
            paragraph(container, "")
        else:
            raise ValueError(f"Unknown component: {kind}")

    if report.get("meta"):
        paragraph(document, report["meta"], "Layout Meta")
    paragraph(document, report["title"], "Title")
    if report.get("subtitle"):
        paragraph(document, report["subtitle"])
    for section in report["sections"]:
        columns(1)
        if section.get("title"):
            paragraph(document, section["title"], "Heading 1")
        layout = section.get("layout", recipe["defaultLayout"])
        if layout == "two-column":
            for block in section["blocks"]:
                spanning = block.get("span") == "all"
                columns(1 if spanning else 2)
                draw_block(
                    document, block, available if spanning else (available - 18) / 2
                )
        elif layout in {"sidebar", "modular"}:
            group = []

            def flush(group, layout):
                if not group:
                    return
                if layout == "sidebar":
                    groups = [
                        [b for b in group if b.get("placement") != "aside"],
                        [b for b in group if b.get("placement") == "aside"],
                    ]
                    layout_grid(document, groups, available, [2, 1])
                else:
                    for offset in range(0, len(group), 2):
                        layout_grid(
                            document,
                            [[b] for b in group[offset : offset + 2]],
                            available,
                        )
                group.clear()

            for block in section["blocks"]:
                if block.get("span") == "all":
                    flush(group, layout)
                    draw_block(document, block, available)
                else:
                    group.append(block)
            flush(group, layout)
        else:
            for block in section["blocks"]:
                draw_block(document, block, available)
    columns(1)
    if report.get("footer"):
        paragraph(document, report["footer"], "Source")
    target = output / "report.docx"
    document.save(target)
    return target


def render_pdf(report, output, source_dir, office_dir):
    import typst

    theme = load_asset("themes.json")[report["theme"]]
    recipe = load_asset("recipes.json")[report["recipe"]]
    available = (210 - 36) / 25.4 * 72

    def paragraphs(values):
        return "\n\n".join("#" + literal(value) for value in values)

    def title(block):
        return (
            "#heading(level: 2, " + literal(block["title"]) + ")\n\n"
            if block.get("title")
            else ""
        )

    def notes(block):
        values = [block[key] for key in ["caption", "source"] if block.get(key)]
        return (
            "\n\n"
            + "\n\n".join(
                f'#block[#set text(size: 10pt, fill: rgb("{theme["muted"]}"))\n#{literal(value)}]'
                for value in values
            )
            if values
            else ""
        )

    def image(block):
        name = bundle_image((source_dir / block["image"]).resolve(), output)
        return "#image(" + json.dumps(name) + ", width: 100%)" + notes(block)

    def metric(item):
        value = "#" + literal(item["value"])
        if item.get("unit"):
            value += " #text(size: 13pt)[#" + literal(item["unit"]) + "]"
        lines = [
            f'#block[#set text(size: 32pt, fill: rgb("{theme["accent"]}"))\n{value}]',
            "#strong(" + literal(item["label"]) + ")",
        ]
        lines.extend(
            "#" + literal(item[key])
            for key in ["delta", "period", "note"]
            if item.get(key)
        )
        return "\n\n".join(lines) + notes(item)

    def chart(block, domain=None, width=available):
        svg = chart_svg(block, theme, domain, 640 * 10.5 / width)
        name = "assets/chart-" + hashlib.sha256(svg.encode()).hexdigest()[:20] + ".svg"
        target = output / name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(svg, encoding="utf-8")
        context = " · ".join(
            str(block[key]) for key in ["unit", "period"] if block.get(key)
        )
        unit = "#" + literal(context) + "\n\n" if context else ""
        return (
            title(block)
            + unit
            + "#image("
            + json.dumps(name)
            + ", width: 100%)"
            + notes(block)
        )

    def grid(groups, columns="(1fr, 1fr)", width=available, weights=None):
        weights = weights or [1] * len(groups)
        usable = width - 16 * (len(groups) - 1)
        return (
            f"#grid(columns: {columns}, gutter: 16pt,\n"
            + ",\n".join(
                "[\n"
                + "\n\n".join(
                    draw_block(b, usable * weight / sum(weights)) for b in group
                )
                + "\n]"
                for group, weight in zip(groups, weights)
            )
            + ")"
        )

    def draw_block(block, width=available):
        content = draw_block_content(block, width)
        if block["type"] in NATIVE_UNITS:
            return "#layout-unit[\n" + content + "\n]"
        return content

    def draw_block_content(block, width=available):
        kind = block["type"]
        if kind == "heading":
            return (
                "#heading(level: "
                + str(block.get("level", 2) - 1)
                + ", "
                + literal(block["text"])
                + ")"
            )
        if kind == "standfirst":
            return (
                '#block[#set text(fill: rgb("'
                + theme["muted"]
                + '"))\n#'
                + literal(block["text"])
                + "]"
            )
        if kind == "metric-sparkline":
            return (
                metric({**block, "source": None})
                + "\n\n"
                + chart(
                    {**block, "period": None, "unit": None, "title": None}, width=width
                )
            )
        if kind in {"target-comparison", "progress"}:
            return (
                title(block)
                + "#strong("
                + literal(block["label"])
                + ")\n\n"
                + chart({**block, "title": None}, width=width)
            )
        if kind in {"bar-chart", "line-chart"} | GRAPHICS:
            return chart(block, width=width)
        if kind == "image-text":
            return grid(
                [
                    [{**block, "type": "figure", "title": None}],
                    [
                        {
                            "type": "prose",
                            "title": block.get("title"),
                            "paragraphs": block["paragraphs"],
                        }
                    ],
                ],
                width=width,
            )
        if kind == "metric-strip":
            columns = "(" + ",".join("1fr" for _ in block["items"]) + ",)"
            return grid(
                [[{**item, "type": "big-number"}] for item in block["items"]],
                columns,
                width=width,
            )
        content = title(block)
        if kind == "prose":
            content += paragraphs(block["paragraphs"])
        elif kind == "big-number":
            content += metric(block)
        elif kind == "change-indicator":
            content += (
                f'#block[#set text(size: 18pt, fill: rgb("{theme["accent"]}"))\n#{literal(DIRECTIONS[block["direction"]] + " " + block["value"])}]\n\n#strong({literal(block["label"])})\n\n#{literal(block["comparison"])}'
                + notes(block)
            )
        elif kind in {"callout", "source-note", "pull-quote"}:
            text = "#" + literal(block["text"])
            if kind == "callout":
                content += f'#block(width: 100%, inset: 12pt, fill: rgb("{theme["tint"]}"))[{text}]'
            elif kind == "pull-quote":
                content += (
                    "#block(inset: (left: 12pt))[#set text(size: 18pt)\n" + text + "]"
                )
                if block.get("attribution"):
                    content += "\n\n#" + literal(block["attribution"])
            else:
                content += "#block[#set text(size: 10pt)\n" + text + "]"
            content += notes(block)
        elif kind in {"takeaways", "numbered-list"}:
            content += (
                ("#enum(" if kind == "numbered-list" else "#list(")
                + ", ".join(literal(item) for item in block["items"])
                + ")"
                + notes(block)
            )
        elif kind == "timeline":
            content += "\n\n".join(
                "#grid(columns: (1fr, 3fr), gutter: 12pt, strong("
                + literal(item["label"])
                + "), "
                + literal(item["text"])
                + ")"
                for item in block["items"]
            ) + notes(block)
        elif kind == "process-steps":
            content += (
                "#enum("
                + ", ".join(
                    "[\n#strong("
                    + literal(item["label"])
                    + ")\n\n#"
                    + literal(item["text"])
                    + "\n]"
                    for item in block["items"]
                )
                + ")"
                + notes(block)
            )
        elif kind == "figure":
            content += image(block)
        elif kind == "small-multiples":
            domain = chart_domain(block)
            panels = [
                "[\n" + chart(panel, domain, (width - 16) / 2) + "\n]"
                for panel in block["charts"]
            ]
            content += (
                "#grid(columns: (1fr, 1fr), gutter: 16pt, "
                + ",".join(panels)
                + ")"
                + notes(block)
            )
        elif kind in {"data-table", "comparison-table", "heatmap-table"}:
            headers = (
                "table.header("
                + ", ".join(
                    "strong(" + literal(value) + ")" for value in block["headers"]
                )
                + ")"
            )
            cells = [headers]
            cells.extend(
                "table.cell(align: "
                + ("right" if numeric_cell(block, index, value) else "start")
                + (
                    ', fill: rgb("' + heat_color(block, index, value, theme) + '")'
                    if kind == "heatmap-table"
                    else ""
                )
                + ", "
                + literal(value)
                + ")"
                for row in block["rows"]
                for index, value in enumerate(row)
            )
            columns = "(" + ",".join("1fr" for _ in block["headers"]) + ",)"
            content += (
                f'#table(columns: {columns}, stroke: (x: none, y: .5pt + rgb("{theme["rule"]}")), inset: 5pt, '
                + ",\n".join(cells)
                + ")"
                + notes(block)
            )
            if kind == "heatmap-table":
                content += (
                    "\n\n#block[#set text(size: 10pt)\n#"
                    + literal(
                        f"{block.get('scaleLabel', 'Color scale')}: {block['domain'][0]} → {block['domain'][1]}"
                    )
                    + "]"
                )
        else:
            raise ValueError(f"Unknown component: {kind}")
        return content

    def section_content(section):
        layout = section.get("layout", recipe["defaultLayout"])
        blocks, result, group = section["blocks"], [], []

        def flush():
            if not group:
                return
            if layout == "two-column":
                rendered = [draw_block(b, (available - 18) / 2) for b in group]
                units = [
                    "[\n" + content + "\n]"
                    for b, content in zip(group, rendered)
                    if b["type"] in NATIVE_UNITS
                ]
                arguments = "(" + ",".join(units) + (",)" if units else ")")
                result.append(
                    "#balanced-columns(units: "
                    + arguments
                    + ")[\n"
                    + "\n\n".join(rendered)
                    + "\n]"
                )
            elif layout == "sidebar":
                result.append(
                    grid(
                        [
                            [b for b in group if b.get("placement") != "aside"],
                            [b for b in group if b.get("placement") == "aside"],
                        ],
                        "(2fr, 1fr)",
                        weights=[2, 1],
                    )
                )
            elif layout == "modular":
                for offset in range(0, len(group), 2):
                    result.append(grid([[b] for b in group[offset : offset + 2]]))
            else:
                result.extend(draw_block(b) for b in group)
            group.clear()

        for block in blocks:
            if block.get("span") == "all":
                flush()
                result.append(draw_block(block))
            else:
                group.append(block)
        flush()
        return "\n\n".join(result)

    preamble = (office_dir / "assets" / "blank.typ").read_text(encoding="utf-8")
    preamble += f'\n#set page("a4", margin: 18mm, fill: rgb("{theme["background"]}"))\n#set text(font: ("{theme["body"]}", "{theme["cjk"]}"), fill: rgb("{theme["ink"]}"))\n'
    # Typst columns fill the remaining page height by default. Balance short
    # groups using their measured height; keep long groups in normal page flow.
    preamble += """
#let layout-unit(body) = context layout(size => {
  let measured = measure(block(width: size.width, body))
  // Keep ordinary units together; let content taller than a full A4 text
  // area flow instead of overflowing an unbreakable box.
  block(width: 100%, breakable: measured.height > 261mm, body)
})
#let balanced-columns(body, units: ()) = context layout(size => {
  let column-width = (size.width - 18pt) / 2
  let measured = measure(block(width: column-width, body))
  let tallest = 0pt
  for unit in units {
    tallest = calc.max(tallest, measure(block(width: column-width, unit)).height)
  }
  let target = calc.max(measured.height / 2 + 24pt, tallest + 4pt)
  if target < size.height {
    block(height: target, breakable: false)[#columns(2, gutter: 18pt, body)]
  } else {
    columns(2, gutter: 18pt, body)
  }
})
"""
    parts = []
    if report.get("meta"):
        parts.append("#block[#set text(size: 10pt)\n#" + literal(report["meta"]) + "]")
    parts.append("#doc-title[#" + literal(report["title"]) + "]")
    if report.get("subtitle"):
        parts.append("#" + literal(report["subtitle"]))
    for section in report["sections"]:
        if section.get("title"):
            parts.append("#heading(level: 1, " + literal(section["title"]) + ")")
        parts.append(section_content(section))
    if report.get("footer"):
        parts.append("#" + literal(report["footer"]))
    source = output / "report.typ"
    source.write_text(preamble + "\n\n" + "\n\n".join(parts), encoding="utf-8")
    target = output / "report.pdf"
    typst.compile(str(source), output=str(target), root=str(output))
    return target


def render_native(report, output, source_dir, office_dir, formats):
    results = {}
    if "docx" in formats:
        results["docx"] = str(render_word(report, output, source_dir, office_dir))
    if "pdf" in formats:
        results["pdf"] = str(render_pdf(report, output, source_dir, office_dir))
    return results
