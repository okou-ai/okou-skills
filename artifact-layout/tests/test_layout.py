"""Behavioral coverage for layout composition, literal content, and native output."""

import hashlib
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

PACKAGE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PACKAGE / "scripts"))

from catalog import lookup
from preview import build_preview
from render import ASSETS, chart_domain, chart_geometry, load_asset, render_report


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.model = {
            "version": 1,
            "title": "Literal content",
            "language": "en",
            "theme": "editorial",
            "recipe": "editorial-research",
            "sections": [
                {
                    "layout": "single",
                    "blocks": [
                        {"type": "prose", "paragraphs": ["The supplied explanation."]},
                    ],
                }
            ],
        }

    def run_model(self, model=None, formats=("html",)):
        source = self.directory / "report.json"
        source.write_text(json.dumps(model or self.model, ensure_ascii=False))
        return render_report(
            source, self.directory / "output", formats, PACKAGE.parent / "office-files"
        )

    def test_catalog_retains_complete_reference_and_explicit_coverage(self):
        catalog = load_asset("catalog.json")
        entries = catalog["entries"]
        self.assertEqual(len(entries), 121)
        self.assertEqual(len({entry["id"] for entry in entries}), 121)
        self.assertEqual(sum(entry["tier"] == "核心" for entry in entries), 109)
        self.assertEqual(len(catalog["sources"]), 24)
        source_ids = {source["id"] for source in catalog["sources"]}
        self.assertTrue(all(set(entry["sources"]) <= source_ids for entry in entries))
        advanced = next(entry for entry in entries if entry["id"] == "C29")
        self.assertEqual(advanced["rendering"]["coverage"], "reference-only")

    def test_lookup_is_bilingual_and_scoped(self):
        english = lookup(query="big numbers")["entries"]
        chinese = lookup(query="大数字")["entries"]
        self.assertIn("N01", {entry["id"] for entry in english})
        self.assertIn("N01", {entry["id"] for entry in chinese})
        self.assertTrue(
            all(entry["group"] == "L" for entry in lookup(group="L")["entries"])
        )
        with self.assertRaises(ValueError):
            lookup(ids=["UNKNOWN"])

    def test_html_keeps_text_literal_and_supplied_metric_context(self):
        attack = '<script>alert("x")</script> #set page(height: 1pt) 中文'
        self.model["title"] = attack
        self.model["sections"][0]["blocks"] = [
            {"type": "prose", "paragraphs": [attack]},
            {
                "type": "big-number",
                "value": "42%",
                "label": "Adoption",
                "period": "2026",
                "source": "Supplied source",
            },
        ]
        result = self.run_model()
        document = Path(result["html"]).read_text()
        self.assertNotIn("<script>alert", document)
        self.assertIn("&lt;script&gt;", document)
        self.assertIn("42%", document)
        self.assertIn("Supplied source", document)
        self.assertIn("2026", document)

    def test_shape_errors_do_not_create_an_output(self):
        self.model["sections"][0]["blocks"] = [
            {"type": "bar-chart", "labels": ["A", "B"], "values": [1]},
        ]
        with self.assertRaisesRegex(ValueError, "same length"):
            self.run_model()
        self.assertFalse((self.directory / "output").exists())

    def test_nonfinite_chart_and_unknown_component_are_rejected(self):
        self.model["sections"][0]["blocks"] = [
            {"type": "line-chart", "labels": ["A"], "values": [float("nan")]}
        ]
        with self.assertRaisesRegex(ValueError, "finite"):
            self.run_model()
        self.model["sections"][0]["blocks"] = [{"type": "imaginary-layout"}]
        with self.assertRaisesRegex(ValueError, "component"):
            self.run_model()

    def test_asset_is_bundled_under_its_content_hash(self):
        original = ASSETS / "examples" / "field-study.png"
        self.model["sections"][0]["blocks"] = [
            {"type": "figure", "image": str(original), "alt": 'Literal "image"'}
        ]
        result = self.run_model()
        name = (
            "image-" + hashlib.sha256(original.read_bytes()).hexdigest()[:20] + ".png"
        )
        self.assertEqual(
            (self.directory / "output" / "assets" / name).read_bytes(),
            original.read_bytes(),
        )
        self.assertIn(
            'alt="Literal &quot;image&quot;"', Path(result["html"]).read_text()
        )

    def test_small_multiples_compare_values_on_one_scale(self):
        panels = [
            {"type": "bar-chart", "labels": ["A"], "values": [10]},
            {"type": "bar-chart", "labels": ["A"], "values": [100]},
        ]
        domain = chart_domain({"charts": panels})
        theme = load_asset("themes.json")["editorial"]
        first = chart_geometry(panels[0], theme, domain)["rects"][0][2]
        second = chart_geometry(panels[1], theme, domain)["rects"][0][2]
        self.assertAlmostEqual(second / first, 10)

    def test_signed_bars_use_the_same_zero_baseline(self):
        block = {
            "type": "bar-chart",
            "labels": ["Decrease", "Increase"],
            "values": [-20, 20],
        }
        geometry = chart_geometry(block, load_asset("themes.json")["mono"])
        zero = geometry["lines"][0][0]
        left, right = geometry["rects"]
        self.assertAlmostEqual(left[0] + left[2], zero)
        self.assertAlmostEqual(right[0], zero)

    def test_long_chart_label_wraps_without_losing_text(self):
        label = "Long customer segment / 长期客户分组及其说明"
        block = {"type": "bar-chart", "labels": [label], "values": [42]}
        geometry = chart_geometry(block, load_asset("themes.json")["editorial"])
        label_parts = [text[2] for text in geometry["texts"] if text[0] == 0]
        self.assertGreater(len(label_parts), 1)
        self.assertEqual("".join(label_parts), label)

    def test_preview_has_actual_pages_for_every_recipe_and_theme(self):
        site = self.directory / "preview"
        build_preview(site)
        for recipe in load_asset("recipes.json"):
            for theme in load_asset("themes.json"):
                path = site / "reports" / f"{recipe}-{theme}" / "index.html"
                self.assertTrue(path.is_file())
                self.assertIn("ILLUSTRATIVE CONTENT", path.read_text())
        self.assertNotIn("THEME_BUTTONS", (site / "index.html").read_text())


NATIVE_AVAILABLE = all(
    importlib.util.find_spec(package) for package in ["docx", "typst", "PIL", "pypdf"]
)


@unittest.skipUnless(NATIVE_AVAILABLE, "Native integration dependencies are required")
class NativeTests(unittest.TestCase):
    def test_all_recipes_render_editable_word_and_actual_pdf(self):
        from docx import Document
        from pypdf import PdfReader

        for recipe in load_asset("recipes.json"):
            with (
                self.subTest(recipe=recipe),
                tempfile.TemporaryDirectory() as temporary,
            ):
                source = ASSETS / "examples" / (recipe + ".json")
                result = render_report(
                    source, temporary, ("docx", "pdf"), PACKAGE.parent / "office-files"
                )
                word = Document(result["docx"])
                self.assertTrue(word.paragraphs)
                pdf = PdfReader(result["pdf"])
                self.assertGreater(len(pdf.pages), 0)
                text = "".join(page.extract_text() for page in pdf.pages)
                self.assertIn("LAYOUT STUDY", text)
                self.assertIn("ILLUSTRATIVE CONTENT", text)

    def test_mixed_script_and_typst_control_text_remain_literal(self):
        from docx import Document
        from pypdf import PdfReader

        text = (
            "中文分析与 English 42% — #set page(height: 1pt) <script>literal</script>"
        )
        report = {
            "version": 1,
            "title": "混合语言 / Mixed script",
            "language": "zh-CN",
            "theme": "warm-editorial",
            "recipe": "editorial-research",
            "sections": [
                {
                    "layout": "two-column",
                    "blocks": [{"type": "prose", "paragraphs": [text] * 12}],
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "model.json"
            source.write_text(json.dumps(report, ensure_ascii=False))
            result = render_report(
                source,
                Path(temporary) / "output",
                ("docx", "pdf"),
                PACKAGE.parent / "office-files",
            )
            word = Document(result["docx"])
            self.assertEqual(sum(p.text == text for p in word.paragraphs), 12)
            columns = word.element.xpath(".//w:cols[@w:num='2']")
            self.assertTrue(columns)
            pdf = PdfReader(result["pdf"])
            extracted = "".join(page.extract_text() for page in pdf.pages)
            self.assertIn("#set page(height: 1pt)", " ".join(extracted.split()))
            self.assertIn("中文", extracted)
            self.assertEqual(len(pdf.pages), 1)

    def test_word_svg_rejection_precedes_output(self):
        report = {
            "version": 1,
            "title": "SVG",
            "language": "en",
            "theme": "mono",
            "recipe": "editorial-research",
            "sections": [
                {
                    "blocks": [
                        {
                            "type": "figure",
                            "image": "figure.svg",
                            "alt": "A source vector",
                        }
                    ]
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            (directory / "figure.svg").write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" />'
            )
            source = directory / "model.json"
            source.write_text(json.dumps(report))
            with self.assertRaisesRegex(ValueError, "raster copy"):
                render_report(
                    source,
                    directory / "output",
                    ("docx",),
                    PACKAGE.parent / "office-files",
                )
            self.assertFalse((directory / "output").exists())

    def test_long_column_flow_keeps_all_content_across_pages(self):
        from pypdf import PdfReader

        paragraphs = [
            f"PARA{index:02d} "
            + "The source text continues with readable body type. " * 12
            for index in range(30)
        ]
        report = {
            "version": 1,
            "title": "Long column flow",
            "language": "en",
            "theme": "editorial",
            "recipe": "editorial-research",
            "sections": [
                {
                    "layout": "two-column",
                    "blocks": [{"type": "prose", "paragraphs": paragraphs}],
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "model.json"
            source.write_text(json.dumps(report))
            result = render_report(
                source,
                Path(temporary) / "output",
                ("pdf",),
                PACKAGE.parent / "office-files",
            )
            pdf = PdfReader(result["pdf"])
            text = "".join(page.extract_text() for page in pdf.pages)
            self.assertGreater(len(pdf.pages), 1)
            for index in range(30):
                self.assertIn(f"PARA{index:02d}", text)


if __name__ == "__main__":
    unittest.main()
