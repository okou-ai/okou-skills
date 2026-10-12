"""Verify the ranked expressions by rendering real HTML and native artifacts."""

# ruff: noqa: E402 -- Imports follow the local script path bootstrap.

import copy
import importlib.util
import json
import math
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

PACKAGE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PACKAGE / "scripts"))

from catalog import lookup
from primitives import heat_color
from render import chart_geometry, load_asset, render_report
from specimens import build_specimens, ranked_reports, resolve_images


class Top30Tests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.theme = load_asset("themes.json")["corporate"]

    def render(self, report, name="case", formats=("html",)):
        model = copy.deepcopy(report)
        resolve_images(model)
        source = self.directory / (name + ".json")
        source.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
        return render_report(source, self.directory / name, formats)

    def test_priority_is_ordered_complete_and_each_expression_renders(self):
        entries = lookup(top30=True)["entries"]
        self.assertEqual(len(entries), 30)
        self.assertEqual(len({entry["id"] for entry in entries}), 30)
        self.assertEqual(
            [entry["priority"]["rank"] for entry in entries], list(range(1, 31))
        )
        self.assertTrue(
            all(
                entry["rendering"]["formats"] == ["html", "docx", "pdf"]
                for entry in entries
            )
        )
        for entry, report in ranked_reports():
            with self.subTest(expression=entry["id"]):
                result = self.render(report, entry["id"])
                document = Path(result["html"]).read_text()
                self.assertIn(entry["component"], document)
                self.assertIn("ILLUSTRATIVE CONTENT", document)

    def test_grouped_signed_values_share_a_zero_baseline(self):
        chart = {
            "type": "grouped-bar-chart",
            "labels": ["A"],
            "series": [
                {"name": "Prior", "values": [-20]},
                {"name": "Current", "values": [40]},
            ],
        }
        geometry = chart_geometry(chart, self.theme)
        negative, positive = geometry["rects"]
        self.assertAlmostEqual(negative[0] + negative[2], positive[0])
        self.assertAlmostEqual(positive[2], negative[2] * 2)
        labels = "".join(item[2] for item in geometry["texts"])
        self.assertIn("A / Prior", labels)
        self.assertIn("A / Current", labels)

    def test_stacks_keep_absolute_totals_and_normalize_each_category(self):
        data = {
            "labels": ["A", "B"],
            "series": [
                {"name": "One", "values": [20, 40]},
                {"name": "Two", "values": [30, 60]},
            ],
        }
        raw = chart_geometry({"type": "stacked-bar-chart", **data}, self.theme)
        percentages = chart_geometry(
            {"type": "percent-stacked-bar-chart", **data}, self.theme
        )
        raw_bars = [rect for rect in raw["rects"] if rect[3] == 24]
        pct_bars = [rect for rect in percentages["rects"] if rect[3] == 24]
        self.assertAlmostEqual(
            sum(rect[2] for rect in raw_bars[2:]),
            2 * sum(rect[2] for rect in raw_bars[:2]),
        )
        self.assertAlmostEqual(sum(rect[2] for rect in pct_bars[:2]), 290)
        self.assertAlmostEqual(sum(rect[2] for rect in pct_bars[2:]), 290)
        labels = "".join(item[2] for item in percentages["texts"])
        self.assertIn("20 (40.0%)", labels)
        self.assertIn("60 (60.0%)", labels)

    def test_progress_and_target_keep_their_denominators(self):
        progress = chart_geometry(
            {"type": "progress", "value": 18, "total": 24}, self.theme
        )
        self.assertAlmostEqual(progress["rects"][1][2] / progress["rects"][0][2], 0.75)
        target = chart_geometry(
            {"type": "target-comparison", "actual": 72, "target": 85, "maximum": 100},
            self.theme,
        )
        self.assertAlmostEqual(target["lines"][0][0], 10 + 620 * 0.85)
        self.assertIn("Target 85", [item[2] for item in target["texts"]])

    def test_flat_sparkline_remains_finite(self):
        geometry = chart_geometry(
            {"type": "metric-sparkline", "values": [5, 5, 5]}, self.theme
        )
        self.assertTrue(
            all(
                math.isfinite(coordinate)
                for point in geometry["points"]
                for coordinate in point
            )
        )
        self.assertEqual(len({point[1] for point in geometry["points"]}), 1)

    def test_heatmap_keeps_exact_values_and_a_fixed_scale(self):
        report = next(
            report for entry, report in ranked_reports() if entry["id"] == "B03"
        )
        block = report["sections"][0]["blocks"][0]
        low = heat_color(block, 1, 0, self.theme)
        high = heat_color(block, 1, 100, self.theme)
        self.assertEqual(low, self.theme["background"])
        self.assertNotEqual(low, high)
        self.assertEqual(high, heat_color(block, 1, 200, self.theme))
        html = Path(self.render(report)["html"]).read_text()
        for value in (82, 64, 71, 68, 73, 80):
            self.assertIn(f">{value}</td>", html)
        self.assertIn("0 → 100", html)

    def test_invalid_new_shapes_fail_before_output(self):
        cases = [
            {"type": "heading", "text": "A", "level": 1},
            {"type": "standfirst", "text": ""},
            {
                "type": "change-indicator",
                "label": "A",
                "value": "+2%",
                "comparison": "Prior",
                "direction": "good",
            },
            {
                "type": "target-comparison",
                "label": "A",
                "actual": 110,
                "target": 80,
                "maximum": 100,
            },
            {"type": "progress", "label": "A", "value": 1, "total": 0},
            {
                "type": "metric-sparkline",
                "label": "A",
                "value": "5",
                "period": "2026",
                "values": [float("nan"), 5],
            },
            {
                "type": "grouped-bar-chart",
                "labels": ["A"],
                "series": [{"name": "B", "values": [1, 2]}],
            },
            {
                "type": "stacked-bar-chart",
                "labels": ["A"],
                "series": [{"name": "B", "values": [-1]}],
            },
            {
                "type": "percent-stacked-bar-chart",
                "labels": ["A"],
                "series": [{"name": "B", "values": [0]}],
            },
            {
                "type": "heatmap-table",
                "headers": ["A"],
                "rows": [[5]],
                "domain": [0, 100],
                "heatColumns": [1],
            },
            {"type": "process-steps", "items": [{"label": "A"}]},
            {"type": "numbered-list", "items": [5]},
        ]
        base = ranked_reports()[0][1]
        for index, block in enumerate(cases):
            with self.subTest(component=block["type"]):
                report = copy.deepcopy(base)
                report["sections"] = [{"layout": "single", "blocks": [block]}]
                with self.assertRaises(ValueError):
                    self.render(report, f"invalid{index}")
                self.assertFalse((self.directory / f"invalid{index}").exists())

    def test_downloadable_models_are_portable_and_use_all_four_themes(self):
        markup, links = build_specimens(self.directory)
        self.assertEqual(markup.count('data-term="'), 30)
        self.assertEqual(
            len(list(self.directory.glob("expressions/*/index.html"))), 120
        )
        for identity in ["I01", "I02", "N04", "C06"]:
            for theme, paths in links[identity]["themes"].items():
                model = self.directory / paths["model"]
                report = json.loads(model.read_text())
                self.assertEqual(report["theme"], theme)
                for section in report["sections"]:
                    for block in section["blocks"]:
                        if "image" in block:
                            self.assertFalse(Path(block["image"]).is_absolute())
                            self.assertTrue((model.parent / block["image"]).is_file())
                render_report(
                    model, self.directory / "portable" / f"{identity}-{theme}"
                )
                archive = self.directory / paths["bundle"]
                extracted = self.directory / "unpacked" / f"{identity}-{theme}"
                with zipfile.ZipFile(archive) as bundle:
                    bundle.extractall(extracted)
                render_report(
                    extracted / "report.json",
                    self.directory / "from-zip" / f"{identity}-{theme}",
                )


@unittest.skipUnless(
    all(importlib.util.find_spec(name) for name in ["docx", "typst", "PIL", "pypdf"]),
    "Native engines not installed",
)
class Top30NativeTests(unittest.TestCase):
    setUp = Top30Tests.setUp
    render = Top30Tests.render

    def test_ordered_blocks_restart_independent_sequences(self):
        from docx import Document

        reports = {entry["id"]: report for entry, report in ranked_reports()}
        report = reports["T08"]
        report["sections"][0]["blocks"] += reports["D02"]["sections"][0]["blocks"]
        word = Document(self.render(report, formats=("docx",))["docx"])
        items = [
            paragraph
            for paragraph in word.paragraphs
            if paragraph.style.name == "List Number"
        ]
        sequences = [item._p.pPr.numPr.numId.val for item in items]
        self.assertEqual(len(set(sequences[:4])), 1)
        self.assertEqual(len(set(sequences[4:])), 1)
        self.assertNotEqual(sequences[0], sequences[4])
        for sequence in set(sequences):
            starts = word.part.numbering_part.element.xpath(
                f'./w:num[@w:numId="{sequence}"]/w:lvlOverride/w:startOverride/@w:val'
            )
            self.assertEqual(starts, ["1"])

    def test_every_expression_produces_editable_word_and_real_pdf(self):
        from docx import Document
        from pypdf import PdfReader

        for entry, report in ranked_reports():
            with self.subTest(expression=entry["id"]):
                files = self.render(report, entry["id"], ("docx", "pdf"))
                word = Document(files["docx"])
                pdf = PdfReader(files["pdf"])
                self.assertGreaterEqual(len(pdf.pages), 1)
                self.assertIn(
                    entry["id"],
                    " ".join(paragraph.text for paragraph in word.paragraphs),
                )
                self.assertIn(
                    entry["id"], " ".join(page.extract_text() for page in pdf.pages)
                )
                if entry["id"] == "B03":
                    cells = word.tables[0].rows[1].cells
                    self.assertEqual(cells[1].text, "82")
                    self.assertTrue(cells[1]._tc.xpath("./w:tcPr/w:shd"))
                    self.assertEqual(word.styles["Table Text"].font.size.pt, 11)
