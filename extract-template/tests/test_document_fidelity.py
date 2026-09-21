"""Regression tests for document reversal. Requires pandoc 3.x, Writer, Poppler,
PyMuPDF, Pillow and python-docx. Run: python3 -m unittest discover -s extract-template/tests -p 'test_*.py' -v
"""
from contextlib import redirect_stdout
import hashlib
import io
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from xml.etree import ElementTree as ET
import zipfile

import docx
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageChops
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "docx" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from docx_layout import Styles, choose_body, choose_section, section_xml, W, NS
from docx_slots import read_slots
from build_reference import build
from verify_reference import main as verify
from inspect_docx import main as inspect
from rendered_probe import MARKERS, check_rendered_probe


def call(*args):
    return subprocess.run([str(a) for a in args], capture_output=True, text=True, timeout=300)


def rewrite(path, edits):
    with zipfile.ZipFile(path) as z:
        members = [(item, z.read(item.filename)) for item in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for item, contents in members:
            z.writestr(item, edits.pop(item.filename, contents))
        for name, contents in edits.items():
            z.writestr(name, contents)


def document(path, cover=False, white=False, control=False):
    d = docx.Document()
    section = d.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(3 if cover else 1)
    d.styles["Normal"].font.name = "Liberation Sans"
    d.styles["Normal"].font.size = Pt(11)
    footer = section.footer.paragraphs[0]
    for text in ("w", "ww.example.co", "m"):
        footer.add_run(text)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "1"
    run.append(text); field.append(run); footer._p.append(field)
    if cover:
        parent = d.styles.add_style("CoverFrame", WD_STYLE_TYPE.PARAGRAPH)
        parent.base_style = d.styles["Normal"]
        frame = OxmlElement("w:framePr"); frame.set(qn("w:vAnchor"), "page"); frame.set(qn("w:y"), "5119")
        parent.element.get_or_add_pPr().append(frame)
        info = d.styles.add_style("Info", WD_STYLE_TYPE.PARAGRAPH); info.base_style = parent
        for number in range(8):
            d.add_paragraph(f"Contact {number}", "Info")
        d.styles["Heading 1"].base_style = parent
        d.styles["Heading 1"].font.color.rgb = RGBColor.from_string("FFFFFF")
        d.add_paragraph("20XX", "Heading 1")
        section = d.add_section(WD_SECTION_START.NEW_PAGE)
        section.top_margin = Inches(1)
        d.add_paragraph("Body heading", "Heading 2")
    elif white:
        d.styles["Heading 1"].font.color.rgb = RGBColor.from_string("FFFFFF")
    p = d.add_paragraph("Body prose should retain its own margins and typography. " * 9)
    d.add_paragraph("Second paragraph containing additional ordinary prose. " * 6)
    if control:
        heading = d.add_paragraph(style="Title")
        r = heading.add_run("Green control title")
        r.font.color.rgb = RGBColor.from_string("008844")
        r.font.name = "Liberation Sans"; r.font.size = Pt(24)
        heading._p.remove(r._r)
        sdt = OxmlElement("w:sdt"); pr = OxmlElement("w:sdtPr")
        showing = OxmlElement("w:showingPlcHdr"); pr.append(showing)
        placeholder = OxmlElement("w:placeholder"); link = OxmlElement("w:docPart")
        link.set(qn("w:val"), "TitlePlaceholder"); placeholder.append(link); pr.append(placeholder)
        pr.append(OxmlElement("w:text"))
        content = OxmlElement("w:sdtContent"); content.append(r._r)
        sdt.append(pr); sdt.append(content); heading._p.append(sdt)
    p = d.add_paragraph()
    for text in ("Executive ", "s", "ummary\u00a0"):
        p.add_run(text)
    d.add_paragraph("", "Subtitle")
    table = d.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Repeatable row"
    d.save(path)
    if control:
        with zipfile.ZipFile(path) as z:
            rels = z.read("word/_rels/document.xml.rels").decode()
            types = z.read("[Content_Types].xml").decode()
        rels = rels.replace("</Relationships>", '<Relationship Id="rIdGlossary" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/glossaryDocument" Target="glossary/document.xml"/></Relationships>')
        types = types.replace("</Types>", '<Override PartName="/word/glossary/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.glossary+xml"/></Types>')
        glossary = f'<w:glossaryDocument xmlns:w="{W[1:-1]}"><w:docParts><w:docPart><w:docPartPr><w:name w:val="TitlePlaceholder"/></w:docPartPr><w:docPartBody><w:p><w:r><w:t>Green control title</w:t></w:r></w:p></w:docPartBody></w:docPart></w:docParts></w:glossaryDocument>'
        rewrite(path, {"word/_rels/document.xml.rels": rels.encode(), "[Content_Types].xml": types.encode(), "word/glossary/document.xml": glossary.encode()})
    return path


class DocumentFidelity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="document-fidelity-")
        cls.base = Path(cls.temp.name)
        cls.plain = document(cls.base / "plain.docx")
        cls.cover = document(cls.base / "cover.docx", cover=True)
        cls.control = document(cls.base / "control.docx", control=True)
        cls.white = document(cls.base / "white.docx", white=True)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def tempdir(self):
        path = Path(tempfile.mkdtemp(dir=self.base))
        return path

    def build(self, source, **kwargs):
        target = self.tempdir() / "reference.docx"
        with redirect_stdout(io.StringIO()):
            build(source, target, **kwargs)
        return target

    def test_cover_paragraph_count_does_not_choose_body_or_section(self):
        with zipfile.ZipFile(self.cover) as z:
            source, styles = z.read("word/document.xml"), Styles(z.read("word/styles.xml"))
        self.assertEqual(choose_body(source, styles), "Normal")
        self.assertEqual(choose_section(source, styles, "Normal"), 2)
        selected = ET.fromstring(section_xml(source, 2))
        self.assertEqual(selected.find(W + "pgMar").get(W + "top"), "1440")
        self.assertIsNotNone(selected.find(W + "footerReference"))

    def test_unsafe_cover_heading_is_not_silently_packaged(self):
        with self.assertRaisesRegex(ValueError, "heading 1 inherits framePr"):
            self.build(self.cover)

    def test_explicit_body_mapping_preserves_source_and_resolves_header_inheritance(self):
        before = hashlib.sha256(self.cover.read_bytes()).hexdigest()
        target = self.tempdir() / "reference.docx"
        output = io.StringIO()
        with redirect_stdout(output):
            build(self.cover, target, {"__body__": "Normal", "heading 2": "Heading1"}, 2)
        self.assertNotIn("source style not found", output.getvalue())
        self.assertEqual(hashlib.sha256(self.cover.read_bytes()).hexdigest(), before)
        with zipfile.ZipFile(target) as z:
            styles = Styles(z.read("word/styles.xml"))
            self.assertIsNone(styles.property(styles.resolve("Body Text"), "pPr", "framePr"))
            root = ET.fromstring(z.read("word/document.xml"))
            self.assertEqual(root.find("w:body/w:sectPr/w:pgMar", NS).get(W + "top"), "1440")
            self.assertIsNotNone(root.find("w:body/w:sectPr/w:footerReference", NS))
        with redirect_stdout(io.StringIO()):
            self.assertEqual(verify(str(target), render_dir=str(self.tempdir())), 0)

    def test_invalid_body_mapping_fails_before_output(self):
        with self.assertRaisesRegex(ValueError, "Unknown source style"):
            self.build(self.plain, mapping={"__body__": "NoSuchStyle"})
        with self.assertRaisesRegex(ValueError, "inherits framePr"):
            self.build(self.cover, mapping={"__body__": "Info"})

    def test_suggested_map_never_turns_body_copy_into_a_heading(self):
        source = self.tempdir() / "custom.docx"
        d = docx.Document()
        for name in ("Memo Title", "Section Head", "Body Copy"):
            style = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = d.styles["Normal"]
            style.font.size = Pt(20 if name == "Memo Title" else 14 if name == "Section Head" else 11)
        d.add_paragraph("Memo", "Memo Title")
        d.add_paragraph("Scope", "Section Head")
        for _ in range(3):
            d.add_paragraph("Ordinary body prose. " * 20, "Body Copy")
        table = d.add_table(rows=3, cols=1)
        for cell in table.column_cells(0):
            cell.text = "Table cell"
        d.save(source)
        output = io.StringIO()
        with redirect_stdout(output):
            inspect(source)
        suggestion = re.search(r"--map '([^']+)'", output.getvalue()).group(1)
        self.assertIn("__body__=body copy", suggestion)
        self.assertNotIn("body copy=Heading", suggestion)

    def test_invalid_section_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "--section must be between"):
            self.build(self.plain, section=3)

    def test_ambiguous_body_requires_mapping(self):
        xml = f'<w:document xmlns:w="{W[1:-1]}"><w:body><w:p><w:r><w:t>Equal</w:t></w:r></w:p><w:p><w:pPr><w:pStyle w:val="Other"/></w:pPr><w:r><w:t>Equal</w:t></w:r></w:p><w:sectPr/></w:body></w:document>'
        styles = Styles(f'<w:styles xmlns:w="{W[1:-1]}"><w:style w:styleId="Normal"><w:name w:val="Normal"/></w:style><w:style w:styleId="Other"><w:name w:val="Other"/></w:style></w:styles>')
        with self.assertRaisesRegex(ValueError, "ambiguous"):
            choose_body(xml, styles)
        self.assertEqual(choose_body(xml, styles, {"__body__": "Other"}), "Other")

    def test_slots_include_footer_split_runs_and_field_marking(self):
        records = read_slots(self.control)
        footer = next(row for row in records if row["part"].startswith("word/footer") and row["text"])
        self.assertEqual(footer["text"], "www.example.com1")
        self.assertEqual([run["text"] for run in footer["runs"][:3]], ["w", "ww.example.co", "m"])
        self.assertTrue(footer["runs"][-1]["field"])
        summary = next(row for row in records if row["text"] == "Executive summary\u00a0")
        self.assertEqual(len(summary["runs"]), 3)
        self.assertEqual(len({run["path"] for run in summary["runs"]}), 3)
        self.assertTrue(any(row["row"] for row in records if row["text"] == "Repeatable row"))
        self.assertTrue(any(not row["runs"] and row["style"] == "Subtitle" for row in records))

    def test_slots_link_the_actual_glossary_placeholder(self):
        runs = [run for row in read_slots(self.control) for run in row["runs"] if run.get("placeholder_name")]
        self.assertEqual(len(runs), 1)
        self.assertTrue(runs[0]["showing_placeholder"])
        self.assertEqual(runs[0]["placeholder"]["part"], "word/glossary/document.xml")
        self.assertEqual(runs[0]["placeholder"]["text"], "Green control title")

    def test_nested_textboxes_and_fields_crossing_paragraphs_keep_unique_addresses(self):
        source = self.tempdir() / "nested.docx"
        source.write_bytes(self.plain.read_bytes())
        with zipfile.ZipFile(source) as z:
            root = ET.fromstring(z.read("word/document.xml"))
        body = root.find(W + "body")
        final_section = body.find(W + "sectPr")
        body.remove(final_section)
        outer = ET.SubElement(body, W + "p")
        drawing = ET.SubElement(ET.SubElement(outer, W + "r"), W + "pict")
        box = ET.SubElement(drawing, W + "txbxContent")
        inner = ET.SubElement(ET.SubElement(box, W + "p"), W + "r")
        ET.SubElement(inner, W + "t").text = "Textbox value"
        begin = ET.SubElement(ET.SubElement(body, W + "p"), W + "r")
        ET.SubElement(begin, W + "fldChar", {W + "fldCharType": "begin"})
        cached = ET.SubElement(ET.SubElement(body, W + "p"), W + "r")
        ET.SubElement(cached, W + "t").text = "Cached field value"
        ET.SubElement(cached, W + "fldChar", {W + "fldCharType": "end"})
        body.append(final_section)
        rewrite(source, {"word/document.xml": ET.tostring(root)})
        runs = [run for row in read_slots(source) for run in row["runs"]]
        self.assertEqual(sum(run["text"] == "Textbox value" for run in runs), 1)
        self.assertTrue(next(run for run in runs if run["text"] == "Cached field value")["field"])

    def test_plain_reference_passes_rendered_probe(self):
        target = self.build(self.plain)
        with redirect_stdout(io.StringIO()):
            self.assertEqual(verify(str(target), render_dir=str(self.tempdir())), 0)

    def test_white_on_white_heading_fails_rendered_probe(self):
        target = self.build(self.white)
        output = io.StringIO()
        with redirect_stdout(output):
            self.assertEqual(verify(str(target), render_dir=str(self.tempdir())), 1)
        self.assertIn("Heading one", output.getvalue())

    def test_white_text_on_dark_background_is_visible(self):
        pdf = self.tempdir() / "dark.pdf"
        with pymupdf.open() as d:
            page = d.new_page(width=612, height=792)
            page.draw_rect(page.rect, fill=(0, 0, 0))
            for row, marker in enumerate(MARKERS):
                page.insert_text((72, 72 + row * 24), marker, color=(1, 1, 1))
            d.save(pdf)
        self.assertEqual(check_rendered_probe(pdf, 612, 792), [])

    def test_word_cover_and_all_pages_share_a_flattened_render(self):
        destination = self.tempdir()
        result = call("node", ROOT / "scripts/render-document.mjs", "--input", self.control, "--out", destination)
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(result.stdout)
        self.assertIn("none", call("pdfinfo", data["pdf"]).stdout.split("Form:", 1)[1].splitlines()[0])
        with pymupdf.open(data["pdf"]) as d:
            spans = [span for block in d[0].get_text("dict")["blocks"] if "lines" in block
                     for line in block["lines"] for span in line["spans"]]
            title = next(span for span in spans if "Green control title" in span["text"])
            self.assertEqual(title["color"], 0x008844)
            self.assertAlmostEqual(title["size"], 24, delta=0.1)
        cover = self.tempdir() / "cover.png"
        result = call("node", ROOT / "scripts/cover-page.mjs", "--input", self.control, "--out", cover)
        self.assertEqual(result.returncode, 0, result.stderr)
        with Image.open(data["images"][0]) as image, Image.open(cover) as cover_image:
            self.assertEqual(image.size, (1000, 1295))
            self.assertIsNone(ImageChops.difference(image.convert("RGB"), cover_image.convert("RGB")).getbbox())
        self.assertFalse(list(destination.glob(".prt-document-*")))

    def test_native_pdf_forms_are_allowed_but_word_form_guard_rejects_them(self):
        path = self.tempdir() / "form.pdf"
        with pymupdf.open() as d:
            page = d.new_page(width=612, height=792)
            widget = pymupdf.Widget(); widget.field_name = "name"; widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
            widget.rect = pymupdf.Rect(72, 72, 220, 100); widget.field_value = "Example"
            page.add_widget(widget); d.save(path)
        result = call("node", ROOT / "scripts/cover-page.mjs", "--input", path, "--out", path.with_suffix(".png"))
        self.assertEqual(result.returncode, 0, result.stderr)
        code = f'import {{ describePdf }} from {json.dumps((ROOT / "scripts/document-render.mjs").as_uri())}; describePdf("pdfinfo", process.argv[1], true);'
        result = call("node", "--input-type=module", "-e", code, path)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Word export must contain no form fields", result.stderr)

    def test_mixed_pdf_sizes_keep_each_pages_proportions(self):
        source = self.tempdir() / "mixed.pdf"
        with pymupdf.open() as d:
            for width, height in [(612, 792), (792, 612)]:
                d.new_page(width=width, height=height).insert_text((72, 72), "Page proportion probe")
            d.save(source)
        out = self.tempdir()
        result = call("node", ROOT / "scripts/render-document.mjs", "--input", source, "--out", out)
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(result.stdout)
        self.assertEqual(data["pageCount"], 2)
        sizes = []
        for name in data["images"]:
            with Image.open(name) as image:
                sizes.append(image.size)
        self.assertEqual(sizes, [(1000, 1295), (1000, 773)])
        repeated = call("node", ROOT / "scripts/render-document.mjs", "--input", source, "--out", out)
        self.assertNotEqual(repeated.returncode, 0)
        self.assertIn("must be empty", repeated.stderr)


if __name__ == "__main__":
    unittest.main()
