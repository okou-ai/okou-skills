#!/usr/bin/env python3
"""Create an ordinary XLSX from a JSON specification; never edit a source XLSX."""

import argparse
from copy import copy
from datetime import date
import json
import math
from pathlib import Path
import re
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table


COLORS = ["245C7B", "D19A45", "3C8D7B", "94689A", "7B8794"]
STYLE_FIELDS = {"font", "font_size", "bold", "italic", "color", "fill", "format", "horizontal", "vertical", "wrap", "indent", "border"}
BLOCK_FIELDS = {"start_cell", "columns", "rows", "table", "title"}


class PreservingTranslator(Translator):
    @classmethod
    def translate_range(cls, range_str, rdelta, cdelta):
        # Translator only shifts the last endpoint of Data!A1:Data!B2. Reject
        # unsupported qualifier placement, ignoring quoted names/table columns.
        quoted, brackets, after_colon = False, 0, False
        for char in range_str:
            if char == "'" and not brackets:
                quoted = not quoted
            elif not quoted:
                if char == "[":
                    brackets += 1
                elif char == "]":
                    brackets -= 1
                elif not brackets and char == ":":
                    after_colon = True
                elif not brackets and char == "!" and after_colon:
                    raise TranslatorError("Column formulas need one sheet qualifier before a range, e.g. 'Data'!A2:B3; repeated endpoint qualifiers and 3-D sheet ranges are unsupported")
        # openpyxl 3.1.5 drops the sheet qualifier on named/structured tokens.
        # Keep its A1 translation while preserving a name's worksheet scope.
        qualifier, _ = cls.strip_ws_name(range_str)
        translated = super().translate_range(range_str, rdelta, cdelta)
        return qualifier + translated if qualifier and not translated.startswith(qualifier) else translated


def fields(value, allowed, where):
    if not isinstance(value, dict) or set(value) - set(allowed):
        raise ValueError(f"{where}: expected an object with fields {', '.join(sorted(allowed))}")


def bounds(ref):
    if not isinstance(ref, str) or not re.fullmatch(r"\$?[A-Z]+\$?[1-9][0-9]*(?::\$?[A-Z]+\$?[1-9][0-9]*)?", ref):
        raise ValueError(f"Use an uppercase A1 cell/range without a sheet qualifier: {ref!r}")
    result = range_boundaries(ref)
    x1, y1, x2, y2 = result
    if x2 < x1 or y2 < y1 or x2 > 16384 or y2 > 1048576:
        raise ValueError(f"Invalid XLSX range: {ref}")
    return result


def color(value):
    if not isinstance(value, str) or not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        raise ValueError(f"Colors must be six-digit RGB hex strings: {value!r}")
    return value.upper()


def style_cell(cell, attributes, formats):
    fields(attributes, STYLE_FIELDS, f"{cell.coordinate} style")
    font = copy(cell.font)
    for key, attr in (("font", "name"), ("font_size", "sz"), ("bold", "b"), ("italic", "i")):
        if key in attributes:
            setattr(font, attr, attributes[key])
    if "color" in attributes:
        font.color = color(attributes["color"])
    cell.font = font
    if "fill" in attributes:
        cell.fill = PatternFill("solid", fgColor=color(attributes["fill"]))
    alignment = copy(cell.alignment)
    for key, attr in (("horizontal", "horizontal"), ("vertical", "vertical"), ("wrap", "wrap_text"), ("indent", "indent")):
        if key in attributes:
            setattr(alignment, attr, attributes[key])
    cell.alignment = alignment
    if "format" in attributes:
        cell.number_format = formats.get(attributes["format"], attributes["format"])
    if "border" in attributes:
        edge = Side(style="thin", color=color(attributes["border"]))
        cell.border = Border(bottom=edge)


def put(cell, entry, styles, formats, base=None):
    options = entry if isinstance(entry, dict) else {"value": entry}
    fields(options, STYLE_FIELDS | {"value", "formula", "type", "style"}, f"cell {cell.coordinate}")
    if "formula" in options and "value" in options:
        raise ValueError(f"{cell.coordinate}: use value or formula, not both")
    value = options.get("formula", options.get("value"))
    if "formula" in options and (not isinstance(value, str) or not value.startswith("=")):
        raise ValueError(f"{cell.coordinate}: formula must begin with =")
    if value is not None and (not isinstance(value, (str, bool, int, float)) or isinstance(value, float) and not math.isfinite(value)):
        raise ValueError(f"{cell.coordinate}: value must be a finite JSON scalar")
    if options.get("type") == "date":
        if value is not None:
            if not isinstance(value, str):
                raise ValueError(f"{cell.coordinate}: date values must be ISO date strings or null")
            value = date.fromisoformat(value)
    elif options.get("type") not in (None, "text"):
        raise ValueError(f"{cell.coordinate}: supported explicit types are date and text")
    elif options.get("type") == "text" and value is not None and not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False)
    if "value" in options or "formula" in options:
        cell.value = value
        if options.get("type") == "text" and value is not None:
            cell.data_type = "s"
    defaults = dict(styles["normal"])
    if cell.data_type == "f":
        defaults["wrap"] = False
    style_cell(cell, defaults, formats)
    style_cell(cell, base or {}, formats)
    if "style" in options:
        if options["style"] not in styles:
            raise ValueError(f"Unknown style: {options['style']}")
        style_cell(cell, styles[options["style"]], formats)
    style_cell(cell, {k: v for k, v in options.items() if k in STYLE_FIELDS}, formats)
    if options.get("type") == "date" and "format" not in options and "format" not in (base or {}):
        cell.number_format = formats["date"]


def overlaps(first, second):
    x1, y1, x2, y2 = first
    a, b, c, d = second
    return not (x2 < a or c < x1 or y2 < b or d < y1)


def block_layout(spec, where, reusable):
    """Resolve a block before writing, including its reserved blank cells."""
    fields(spec, BLOCK_FIELDS, where)
    x0, y0, xend, yend = bounds(spec.get("start_cell", "A1"))
    if x0 != xend or y0 != yend:
        raise ValueError(f"{where}: start_cell must be one cell")
    columns, rows = spec.get("columns", []), spec.get("rows", [])
    if not isinstance(columns, list) or not isinstance(rows, list) or any(not isinstance(row, list) for row in rows):
        raise ValueError(f"{where}: columns and rows must be arrays; each row must be an array")
    for column in columns:
        fields(column, {"header", "width", "type", "formula"} | STYLE_FIELDS, f"{where} column")
        if "type" in column and column["type"] not in ("date", "text"):
            raise ValueError(f"{where}: column type must be date or text")
        if "formula" in column:
            if not isinstance(column["formula"], str) or not column["formula"].startswith("="):
                raise ValueError(f"{where}: column formula must begin with =")
            if "type" in column:
                raise ValueError(f"{where}: a formula column uses format, not a literal value type")
    input_count = sum("formula" not in column for column in columns)
    if columns and any(len(row) not in (len(columns), input_count) for row in rows):
        raise ValueError(f"{where}: each row needs {len(columns)} cells, or {input_count} input cells with formula columns omitted")
    title = None
    if "title" in spec:
        title = spec["title"] if isinstance(spec["title"], dict) else {"value": spec["title"]}
        fields(title, {"value", "style"} | STYLE_FIELDS, f"{where} title")
        if not isinstance(title.get("value"), str) or not title["value"]:
            raise ValueError(f"{where}: title requires nonempty literal text")
        title = {**title, "type": "text"}
    width = len(columns) if columns else max([len(row) for row in rows] + [int(title is not None)])
    height = int(title is not None) + int(bool(columns)) + len(rows)
    rectangle = (x0, y0, x0 + width - 1, y0 + height - 1) if width and height else None
    if reusable and rectangle is None:
        raise ValueError(f"{where}: an empty block needs columns or a title; rows may be []")
    if rectangle and (rectangle[2] > 16384 or rectangle[3] > 1048576):
        raise ValueError(f"{where}: block extends beyond the XLSX worksheet limits")
    return {"spec": spec, "where": where, "columns": columns, "rows": rows, "title": title,
            "rectangle": rectangle, "x0": x0, "header_row": y0 + int(title is not None),
            "data_row": y0 + int(title is not None) + int(bool(columns))}


def sheet_widths(sheet, layouts, overrides):
    if not isinstance(overrides, dict):
        raise ValueError(f"{sheet.title}: column_widths must be an object")
    for key in overrides:
        if not re.fullmatch(r"[A-Z]+", key):
            raise ValueError(f"Invalid column_widths key: {key}")
        bounds(key + "1")
    widths, explicit = {}, {}
    for layout in layouts:
        for index, column in enumerate(layout["columns"], layout["x0"]):
            key = get_column_letter(index)
            widths.setdefault(key, 16)
            if "width" in column:
                width = column["width"]
                if isinstance(width, bool) or not isinstance(width, (int, float)) or not 0 < width <= 255:
                    raise ValueError(f"Invalid width for column {key}")
                if key in explicit and explicit[key] != width and key not in overrides:
                    raise ValueError(f"{sheet.title}: conflicting widths for column {key}; set one shared column_widths override")
                explicit[key] = widths[key] = width
    for key, width in {**widths, **overrides}.items():
        if isinstance(width, bool) or not isinstance(width, (int, float)) or not 0 < width <= 255:
            raise ValueError(f"Invalid width for column {key}")
        sheet.column_dimensions[key].width = width


def write_block(sheet, layout, styles, formats, theme, table_names):
    columns, rows = layout["columns"], layout["rows"]
    x0, header_row, data_row = layout["x0"], layout["header_row"], layout["data_row"]
    title = layout["title"]
    title_merge = None
    if title is not None:
        put(sheet.cell(header_row - 1, x0), title, styles, formats, styles["title"])
        title_merge = f"{get_column_letter(x0)}{header_row - 1}:{get_column_letter(layout['rectangle'][2])}{header_row - 1}"
    for index, column in enumerate(columns, x0):
        put(sheet.cell(header_row, index), column.get("header", ""), styles, formats)
        style_cell(sheet.cell(header_row, index), styles["header"], formats)
    translators = {}
    for index, column in enumerate(columns, x0):
        if "formula" in column:
            try:
                translators[index] = PreservingTranslator(column["formula"], origin=f"{get_column_letter(index)}{data_row}")
                translators[index].translate_formula(row_delta=0, col_delta=0)
            except (TokenizerError, TranslatorError) as error:
                raise ValueError(f"{layout['where']}: invalid column formula: {error}") from error
    for offset, values in enumerate(rows, data_row):
        if columns and len(values) != len(columns):
            inputs = iter(values)
            values = [None if "formula" in column else next(inputs) for column in columns]
        for index, value in enumerate(values, x0):
            base = {"fill": theme.get("stripe", "F1F5F8") if (offset - header_row) % 2 else "FFFFFF"}
            column = columns[index - x0] if columns else {}
            base.update({k: v for k, v in column.items() if k in STYLE_FIELDS})
            if index in translators:
                if value is not None and (not isinstance(value, dict) or set(value) - (STYLE_FIELDS | {"style"})):
                    raise ValueError(f"{sheet.title}!{get_column_letter(index)}{offset}: formula-column cells must be null or style-only objects; omit formula columns for compact rows")
                try:
                    formula = column["formula"] if offset == data_row else translators[index].translate_formula(f"{get_column_letter(index)}{offset}")
                except TranslatorError as error:
                    raise ValueError(f"{layout['where']}: cannot translate formula to row {offset}: {error}") from error
                value = {**(value or {}), "formula": formula}
            elif "type" in column:
                value = {"type": column["type"], **(value if isinstance(value, dict) else {"value": value})}
            put(sheet.cell(offset, index), value, styles, formats, base)
    table_name = layout["spec"].get("table")
    if table_name:
        if not columns or not rows or any(not isinstance(c.get("header"), str) or not c["header"] for c in columns):
            raise ValueError(f"{layout['where']}: a table needs column headers and at least one data row; omit table for empty data")
        if not isinstance(table_name, str) or table_name.casefold() in table_names:
            raise ValueError(f"Invalid or duplicate workbook table name: {table_name!r}")
        if len({column["header"].casefold() for column in columns}) != len(columns):
            raise ValueError(f"{layout['where']}: structured-table headers must be unique")
        table_names.add(table_name.casefold())
        ref = f"{get_column_letter(x0)}{header_row}:{get_column_letter(x0 + len(columns) - 1)}{data_row + len(rows) - 1}"
        sheet.add_table(Table(displayName=table_name, ref=ref))
    return title_merge


def merged_ranges(sheet, refs):
    occupied = []
    for ref in refs:
        x1, y1, x2, y2 = bounds(ref)
        if any(not (x2 < a or c < x1 or y2 < b or d < y1) for a, b, c, d in occupied):
            raise ValueError(f"{sheet.title}: overlapping merges: {ref}")
        for row in sheet.iter_rows(min_row=y1, max_row=y2, min_col=x1, max_col=x2):
            for cell in row:
                if cell.coordinate != sheet.cell(y1, x1).coordinate and cell.value is not None:
                    raise ValueError(f"{sheet.title}!{ref}: merge would erase {cell.coordinate}; move its value first")
        anchor = sheet.cell(y1, x1)
        if anchor.data_type == "f" and anchor.alignment.wrap_text:
            raise ValueError(f"{sheet.title}!{ref}: Calc drops wrapping on merged formula cells. Use an unmerged wrapped cell, or a separate short numeric/formula result and literal label.")
        sheet.merge_cells(ref)
        occupied.append((x1, y1, x2, y2))


def text_width(text):
    return sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in text)


def fit_literal_rows(sheet, explicit_heights):
    # A conservative initial height, not a substitute for reviewing the Calc PDF.
    merges = {sheet.cell(r.min_row, r.min_col).coordinate: r for r in sheet.merged_cells.ranges}
    for row in sheet:
        height = 21
        for cell in row:
            if cell.value is not None:
                height = max(height, (cell.font.sz or 11) * 1.4 + 6)
            if not isinstance(cell.value, str) or cell.data_type == "f" or not cell.alignment.wrap_text:
                continue
            merged = merges.get(cell.coordinate)
            columns = range(merged.min_col, merged.max_col + 1) if merged else [cell.column]
            widths = []
            for col in columns:
                dimension = sheet.column_dimensions.get(get_column_letter(col))
                widths.append(dimension.width if dimension is not None else sheet.sheet_format.defaultColWidth or 13)
            width = max(1, sum(widths) - 2 - 3 * (cell.alignment.indent or 0))
            lines = sum(max(1, math.ceil(text_width(part) * (cell.font.sz or 11) / 11 / width)) for part in cell.value.split("\n"))
            height = max(height, lines * (cell.font.sz or 11) * 1.4 + 6)
        if str(row[0].row) not in explicit_heights:
            sheet.row_dimensions[row[0].row].height = height


def add_charts(workbook, sheet, definitions, theme, formats):
    occupied = []
    for spec in definitions:
        fields(spec, {"type", "title", "data_sheet", "data", "categories", "titles_from_data", "anchor", "labels", "number_format", "legend", "colors"}, "chart")
        kind = spec.get("type", "column")
        if kind not in ("column", "bar", "line"):
            raise ValueError("Chart type must be column, bar or line; use Python for other chart types")
        source = workbook[spec.get("data_sheet", sheet.title)]
        data = bounds(spec["data"])
        categories = bounds(spec["categories"])
        headers = spec.get("titles_from_data", True)
        if not isinstance(headers, bool) or not isinstance(spec.get("labels", False), bool):
            raise ValueError("titles_from_data and labels must be JSON booleans")
        if categories[0] != categories[2] or categories[3] - categories[1] + 1 != data[3] - data[1] + 1 - int(headers):
            raise ValueError(f"{sheet.title} chart: categories must be one column matching the number of data rows")
        box = bounds(spec["anchor"])
        x1, y1, x2, y2 = box
        if any(not (x2 < a or c < x1 or y2 < b or d < y1) for a, b, c, d in occupied):
            raise ValueError(f"{sheet.title}: chart rectangles overlap at {spec['anchor']}; leave a row/column gutter")
        occupied.append(box)
        chart = LineChart() if kind == "line" else BarChart()
        if kind == "line":
            chart.smooth = False
        else:
            chart.type = "bar" if kind == "bar" else "col"
            chart.grouping = "clustered"
            chart.overlap = 0
        chart.title = spec.get("title")
        chart.add_data(Reference(source, min_col=data[0], min_row=data[1], max_col=data[2], max_row=data[3]), titles_from_data=headers)
        chart.set_categories(Reference(source, min_col=categories[0], min_row=categories[1], max_col=categories[2], max_row=categories[3]))
        legend = spec.get("legend", "b" if len(chart.series) > 1 else None)
        if legend is None:
            chart.legend = None
        elif legend in ("b", "t", "l", "r"):
            chart.legend.position = legend
        else:
            raise ValueError("Chart legend must be null, b, t, l or r")
        chart.y_axis.numFmt = formats.get(spec.get("number_format", "decimal"), spec.get("number_format", "#,##0.00"))
        if kind == "bar":
            chart.x_axis.scaling.orientation = "maxMin"
            chart.x_axis.tickLblPos = "low"
            chart.y_axis.crosses = "max"
        palette = spec.get("colors", theme.get("chart_colors", COLORS))
        if not isinstance(palette, list) or not palette:
            raise ValueError("Chart colors must be a nonempty array")
        for i, series in enumerate(chart.series):
            hue = color(palette[i % len(palette)])
            series.graphicalProperties.solidFill = hue
            series.graphicalProperties.line.solidFill = hue
            if kind == "line":
                series.smooth = False
            series.dLbls = DataLabelList(showVal=bool(spec.get("labels", False)), showSerName=False, showCatName=False, showLegendKey=False, showPercent=False, showBubbleSize=False, showLeaderLines=False)
        # Explicit cell boundaries survive Calc's anchor conversion; no cm width guess.
        chart.anchor = TwoCellAnchor(editAs="twoCell", _from=AnchorMarker(col=x1 - 1, row=y1 - 1), to=AnchorMarker(col=x2, row=y2))
        sheet.add_chart(chart)
    return occupied


def author(spec, output):
    fields(spec, {"version", "theme", "styles", "sheets"}, "workbook")
    if spec.get("version") != 1 or not isinstance(spec.get("sheets"), list) or not spec["sheets"]:
        raise ValueError("Use version: 1 and a nonempty sheets array")
    theme = spec.get("theme", {})
    fields(theme, {"font", "font_size", "accent", "stripe", "currency", "chart_colors"}, "theme")
    currency = theme.get("currency", "¥")
    if not isinstance(currency, str) or '"' in currency:
        raise ValueError("theme.currency must be a string without double quotes")
    formats = {"money": f'"{currency}"#,##0.00;[Red]-"{currency}"#,##0.00', "integer": "#,##0", "decimal": "#,##0.00", "percent": "0.0%", "date": "yyyy-mm-dd", "text": "@"}
    styles = {
        "normal": {"font": theme.get("font", "Calibri"), "font_size": theme.get("font_size", 11), "color": "243342", "vertical": "center", "wrap": True},
        "header": {"bold": True, "color": "FFFFFF", "fill": theme.get("accent", COLORS[0]), "wrap": True},
        "title": {"bold": True, "font_size": 18, "color": theme.get("accent", COLORS[0])},
        "total": {"bold": True, "fill": "E8EFF3", "border": "B4C7D4"},
        "note": {"font_size": 10, "color": "536575", "wrap": True},
    }
    if not isinstance(spec.get("styles", {}), dict):
        raise ValueError("styles must be an object mapping names to style objects")
    for name, attributes in spec.get("styles", {}).items():
        fields(attributes, STYLE_FIELDS, f"style {name}")
        styles[name] = {**styles.get(name, {}), **attributes}
    book = Workbook()
    book.remove(book.active)
    names = set()
    for sheet_spec in spec["sheets"]:
        if not isinstance(sheet_spec, dict):
            raise ValueError("Each sheets entry must be an object")
        name = sheet_spec.get("name")
        if not isinstance(name, str) or not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name) or name.casefold() in names:
            raise ValueError(f"Invalid or duplicate worksheet name: {name!r}")
        names.add(name.casefold())
        book.create_sheet(name)
    table_names = set()
    for sheet_spec in spec["sheets"]:
        fields(sheet_spec, {"name", "start_cell", "columns", "rows", "blocks", "cells", "column_widths", "row_heights", "merges", "freeze", "filter", "table", "charts", "print"}, "sheet")
        sheet = book[sheet_spec["name"]]
        sheet.sheet_view.showGridLines = False
        sheet.sheet_format.defaultColWidth = 14
        reusable = "blocks" in sheet_spec
        if reusable:
            if set(sheet_spec) & BLOCK_FIELDS:
                raise ValueError(f"{sheet.title}: blocks cannot be combined with sheet-level start_cell, columns, rows or table")
            if not isinstance(sheet_spec["blocks"], list):
                raise ValueError(f"{sheet.title}: blocks must be an array")
            layouts = [block_layout(block, f"{sheet.title} block {index}", True) for index, block in enumerate(sheet_spec["blocks"], 1)]
        else:
            layouts = [block_layout({k: v for k, v in sheet_spec.items() if k in BLOCK_FIELDS}, sheet.title, False)]
        rectangles = []
        for layout in layouts:
            rectangle = layout["rectangle"]
            if rectangle:
                if any(overlaps(rectangle, other) for other in rectangles):
                    raise ValueError(f"{layout['where']}: blocks overlap, including their title/header and blank data cells")
                rectangles.append(rectangle)
        sheet_widths(sheet, layouts, sheet_spec.get("column_widths", {}))
        title_merges = []
        for layout in layouts:
            title_merge = write_block(sheet, layout, styles, formats, theme, table_names)
            if title_merge:
                title_merges.append(title_merge)
        for address, entry in sheet_spec.get("cells", {}).items():
            x, y, last_x, last_y = bounds(address)
            if x != last_x or y != last_y:
                raise ValueError(f"cells keys must be single cell addresses: {address}")
            if reusable and any(overlaps((x, y, x, y), rectangle) for rectangle in rectangles):
                raise ValueError(f"{sheet.title}!{address}: cells cannot overwrite a block; put its value/style in the block row")
            put(sheet.cell(y, x), entry, styles, formats)
        explicit_heights = sheet_spec.get("row_heights", {})
        for key, height in explicit_heights.items():
            if not key.isdigit() or not 1 <= int(key) <= 1048576 or not isinstance(height, (int, float)) or not 0 < height <= 409:
                raise ValueError(f"Invalid row height: {key}")
            sheet.row_dimensions[int(key)].height = height
        merges = sheet_spec.get("merges", [])
        if reusable:
            for ref in merges:
                if any(overlaps(bounds(ref), rectangle) for rectangle in rectangles):
                    raise ValueError(f"{sheet.title}!{ref}: explicit merges cannot intersect a block; use its title or move the merge outside")
        merged_ranges(sheet, [*title_merges, *merges])
        for layout in layouts:
            if layout["title"] is not None:
                title = layout["title"]
                attributes = {**styles["normal"], **styles["title"], **styles.get(title.get("style"), {}),
                              **{k: v for k, v in title.items() if k in STYLE_FIELDS}}
                for col in range(layout["x0"], layout["rectangle"][2] + 1):
                    style_cell(sheet.cell(layout["header_row"] - 1, col), attributes, formats)
        fit_literal_rows(sheet, explicit_heights)
        if sheet_spec.get("freeze"):
            x, y, last_x, last_y = bounds(sheet_spec["freeze"])
            if x != last_x or y != last_y:
                raise ValueError("freeze must be a single cell")
            sheet.freeze_panes = sheet_spec["freeze"]
        if sheet_spec.get("filter"):
            bounds(sheet_spec["filter"])
            sheet.auto_filter.ref = sheet_spec["filter"]
        rectangles = add_charts(book, sheet, sheet_spec.get("charts", []), theme, formats)
        page = sheet_spec.get("print", {})
        fields(page, {"area", "orientation", "title_rows", "break_rows"}, "print")
        max_col = max([sheet.max_column] + [box[2] for box in rectangles])
        max_row = max([sheet.max_row] + [box[3] for box in rectangles])
        area = page.get("area", f"A1:{get_column_letter(max_col)}{max_row}")
        bounds(area)
        sheet.print_area = area
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = page.get("orientation", "landscape" if max_col > 6 else "portrait")
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = sheet.page_margins.right = 0.3
        sheet.page_margins.top = sheet.page_margins.bottom = 0.4
        if page.get("title_rows"):
            sheet.print_title_rows = page["title_rows"]
        for row in page.get("break_rows", []):
            sheet.row_breaks.append(Break(id=row))
        sheet.oddFooter.center.text = "Page &P of &N"
    output = Path(output).resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must have the .xlsx extension")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + ".writing.xlsx")
    book.save(temporary)
    temporary.replace(output)
    return {"status": "AUTHORED_NOT_VERIFIED", "output": str(output), "sheets": book.sheetnames}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", required=True, help="JSON specification; copy assets/workbook.json")
    parser.add_argument("--output", required=True, help="New generated .xlsx path")
    args = parser.parse_args()
    try:
        result = author(json.loads(Path(args.spec).read_text(encoding="utf-8")), args.output)
    except (ValueError, TypeError, KeyError, OSError) as error:
        parser.exit(1, f"author_workbook: {error}\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
