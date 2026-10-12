#!/usr/bin/env python3
"""Check common baseline-style regressions; never claim visual acceptance."""

import argparse
import json
import re

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

SIZES = {
    "Normal": 12,
    "List Bullet": 12,
    "List Number": 12,
    "Title": 18,
    "Heading 1": 13,
    "Heading 2": 13,
    "Heading 3": 13,
    "Table Text": 11,
    "Table Header": 11,
    "Table Title": 11,
    "Caption": 10,
    "Source": 10,
}


def is_number(value):
    value = str(value).strip().replace("−", "-").replace("\u00a0", " ")
    return bool(
        re.fullmatch(r"\(?[+\-]?[¥￥$€£]?\s*\d+(?:[,. ]\d+)*(?:\s*%)?\)?", value)
    )


def inherited(obj, name):
    while obj is not None:
        value = getattr(obj.font, name)
        if value is not None:
            return value
        obj = getattr(obj, "base_style", None)
    return None


def effective(paragraph, run, name):
    direct = getattr(run.font, name)
    if direct is not None:
        return direct
    character = inherited(run.style, name)
    return character if character is not None else inherited(paragraph.style, name)


def paragraph_property(paragraph, name):
    value = getattr(paragraph.paragraph_format, name)
    style = paragraph.style
    while value is None and style is not None:
        value = getattr(style.paragraph_format, name)
        style = style.base_style
    return value


def check_document(document):
    errors = []

    def check_paragraph(paragraph, label, table=False):
        if (
            paragraph._p.xpath(".//w:drawing")
            and paragraph_property(paragraph, "line_spacing_rule")
            == WD_LINE_SPACING.EXACTLY
        ):
            errors.append(f"{label}: image paragraph uses exact line spacing")
        expected = 11 if table else SIZES.get(paragraph.style.name)
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = effective(paragraph, run, "size")
            if expected and (size is None or abs(size.pt - expected) > 0.05):
                errors.append(f"{label}: {paragraph.style.name} must use {expected} pt")
                break

    def check_container(container, label):
        for index, paragraph in enumerate(container.paragraphs, 1):
            check_paragraph(paragraph, f"{label} paragraph {index}")
        for number, table in enumerate(container.tables, 1):
            for index, row in enumerate(table.rows):
                location = f"{label} table {number} row {index + 1}"
                header = bool(
                    row._tr.xpath(
                        "./w:trPr/w:tblHeader[not(@w:val='0') and not(@w:val='false') and not(@w:val='off')]"
                    )
                )
                if index == 0 and not header:
                    errors.append(f"{location}: header is not repeated")
                if row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY:
                    errors.append(f"{location}: exact row height can clip content")
                seen = set()
                for cell in row.cells:
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    for paragraph in cell.paragraphs:
                        check_paragraph(paragraph, location, table=True)
                        for run in paragraph.runs:
                            if (
                                run.text.strip()
                                and header
                                and not effective(paragraph, run, "bold")
                            ):
                                errors.append(f"{location}: header is not bold")
                                break
                        if (
                            not header
                            and is_number(paragraph.text)
                            and paragraph_property(paragraph, "alignment")
                            != WD_ALIGN_PARAGRAPH.RIGHT
                        ):
                            errors.append(
                                f"{location}: numeric cell is not right aligned"
                            )
                    # Nested tables are uncommon in the baseline; inspect their paragraphs too.
                    if cell.tables:
                        check_container(cell, location)

    check_container(document, "body")
    for index, section in enumerate(document.sections, 1):
        for name in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            container = getattr(section, name)
            if not container.is_linked_to_previous:
                check_container(container, f"section {index} {name}")
    elements = [
        part.element for part in document.part.package.parts if hasattr(part, "element")
    ]
    for element in elements:
        for fonts in element.xpath(".//w:rFonts"):
            if any(key.endswith(("Theme", "theme")) for key in fonts.attrib):
                errors.append("font configuration: theme font override remains")
                break
            latin, east = fonts.get(qn("w:ascii")), fonts.get(qn("w:eastAsia"))
            if latin and latin == east and "CJK" in latin:
                errors.append(
                    "font configuration: Latin and CJK slots use the same CJK font"
                )
                break
    return list(dict.fromkeys(errors))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file")
    args = parser.parse_args()
    errors = check_document(Document(args.file))
    print(json.dumps({"errors": errors, "visual_review": "still required"}, indent=2))
    raise SystemExit(bool(errors))
