#!/usr/bin/env python3
"""Recalculate ordinary XLSX files and check independently supplied expectations.

Quick checks do not recalculate or permit delivery. Verify checks fresh results;
render_workbook.py prepares a candidate-bound visual review. No source workbook
is overwritten. Run accept again before delivery.
"""

import argparse
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import posixpath
import re
import shutil
import subprocess
import sys
import zipfile

from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries


NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RISKY_PARTS = ("xl/externalLinks/", "xl/pivot", "xl/slicer", "xl/activeX/", "xl/embeddings/", "xl/queryTables/", "xl/ctrlProps/", "xl/webextensions/", "customXml/")
LIMITATIONS = [
    "Expectations must be calculated independently from the declared raw data; "
    "the script cannot prove their independence or completeness.",
    "Reference checks detect broken references and recalculation changes, not "
    "every valid-looking but wrong range. Review ranges affected by edits.",
    "Formatting, chart appearance, unsupported Excel features and native Excel "
    "compatibility need separate review. LibreOffice may rewrite them.",
    "This is a local snapshot, not tamper-proof evidence. Run accept immediately "
    "before delivering the exact recorded workbook.",
]


def fingerprint(path):
    path = Path(path).resolve(strict=True)
    return {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def formula_key(formula):
    """Ignore case/redundant sheet quotes, but retain strings and intersections."""
    parts = []
    for token in Tokenizer(formula).items:
        value = token.value
        if token.subtype != "TEXT":
            if token.subtype == "RANGE" and "!" in value:
                qualifier, address = value.rsplit("!", 1)
                qualifier = unquote_sheet(qualifier).replace("'", "''")
                value = f"'{qualifier}'!{address}"
            value = value.upper()
        parts.append(value)
    return "".join(parts)


def unquote_sheet(value):
    if value.startswith("'") and value.endswith("'"):
        value = value[1:-1]
    return value.replace("''", "'")


def check_references(expression, sheets, location):
    for token in Tokenizer("=" + expression.lstrip("=")).items:
        if token.subtype == "ERROR":
            raise ValueError(f"Broken reference/error in {location}: {token.value}")
        if token.subtype != "RANGE" or "!" not in token.value:
            continue
        qualifier = unquote_sheet(token.value.rsplit("!", 1)[0])
        if "[" in qualifier or "]" in qualifier:
            raise ValueError(f"External reference needs its native engine: {location}")
        for sheet in qualifier.split(":"):
            if sheet.casefold() not in {name.casefold() for name in sheets}:
                raise ValueError(f"Unknown sheet {sheet!r} referenced by {location}")


def workbook_snapshot(path):
    workbook = load_workbook(path, data_only=False)
    snapshot = {"sheets": workbook.sheetnames, "formulas": {}, "tables": {}, "charts": {}, "defined_names": {}}
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                location = f"{sheet.title}!{cell.coordinate}"
                if cell.data_type == "e":
                    raise ValueError(f"Error cell {location}: {cell.value}")
                if cell.data_type == "f":
                    if not isinstance(cell.value, str):
                        raise ValueError(f"Array/data-table formula requires native Excel review: {location}")
                    check_references(cell.value, workbook.sheetnames, location)
                    snapshot["formulas"][location] = formula_key(cell.value)
        for table in sheet.tables.values():
            bounds = range_boundaries(table.ref)
            if any(value is None or value < 1 for value in bounds) or bounds[2] > 16384 or bounds[3] > 1048576:
                raise ValueError(f"Invalid table range {sheet.title}!{table.ref}")
            snapshot["tables"][f"{sheet.title}!{table.name}"] = table.ref.replace("$", "").upper()
        for index, chart in enumerate(sheet._charts, 1):
            refs = [node.text for node in chart._write().iter() if node.tag.rsplit("}", 1)[-1] == "f" and node.text]
            for ref in refs:
                check_references(ref, workbook.sheetnames, f"{sheet.title} chart {index}")
            snapshot["charts"][f"{sheet.title}!{index}"] = sorted(formula_key("=" + ref) for ref in refs)
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read("xl/workbook.xml"), parser=ET.XMLParser(resolve_entities=False, no_network=True))
        for node in root.findall(f"{{{NS}}}definedNames/{{{NS}}}definedName"):
            key = f"{node.get('localSheetId', 'workbook')}!{node.get('name')}"
            check_references(node.text or "", workbook.sheetnames, f"defined name {key}")
            snapshot["defined_names"][key] = formula_key("=" + (node.text or ""))
    workbook.close()
    return snapshot


def read_expectations(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, dict) or set(data) - {"calculation_basis", "range_review", "cells", "tables"}:
        raise ValueError("Expectations fields: calculation_basis, range_review, cells, optional tables.")
    for field in ("calculation_basis", "range_review"):
        if not isinstance(data.get(field), str) or not data[field].strip():
            raise ValueError(f"Expectations need a nonempty {field} explanation.")
    if not isinstance(data.get("cells"), list) or not data["cells"]:
        raise ValueError("Expectations need nonempty independently derived cell values, including for value-only workbooks.")
    seen = set()
    for item in data["cells"]:
        if not isinstance(item, dict) or set(item) - {"sheet", "cell", "value", "abs_tol", "rel_tol"} or not {"sheet", "cell", "value"} <= set(item):
            raise ValueError("Each cell expectation needs sheet, cell, value; optional abs_tol/rel_tol.")
        if not isinstance(item["sheet"], str) or not item["sheet"] or not isinstance(item["cell"], str) or not re.fullmatch(r"[A-Z]+[1-9][0-9]*", item["cell"]):
            raise ValueError("Cell expectations need a sheet and a single uppercase A1 coordinate.")
        column, row, _, _ = range_boundaries(item["cell"])
        if column > 16384 or row > 1048576:
            raise ValueError(f"Expected cell is outside Excel's grid: {item['cell']}")
        key = (item["sheet"], item["cell"])
        if key in seen:
            raise ValueError(f"Duplicate expected cell: {key}")
        seen.add(key)
        value = item["value"]
        if value is not None and not isinstance(value, (str, bool, int, float)):
            raise ValueError("Expected values must be JSON strings, numbers, booleans or null.")
        if isinstance(value, (int, float)) and not math.isfinite(value):
            raise ValueError("Expected numbers must be finite.")
        for field in ("abs_tol", "rel_tol"):
            tol = item.get(field, 0)
            if isinstance(tol, bool) or not isinstance(tol, (int, float)) or not math.isfinite(tol) or tol < 0:
                raise ValueError("Numeric tolerances must be finite nonnegative numbers.")
    if not isinstance(data.get("tables", []), list):
        raise ValueError("tables must be a list of sheet/name/ref objects.")
    for table in data.get("tables", []):
        if not isinstance(table, dict) or set(table) != {"sheet", "name", "ref"} or any(not isinstance(v, str) or not v for v in table.values()):
            raise ValueError("Each table expectation needs nonempty sheet, name and ref.")
    return data


def check_supported_package(path):
    with zipfile.ZipFile(path) as package:
        if any(name.startswith(RISKY_PARTS) or name in {"xl/vbaProject.bin", "xl/connections.xml"} for name in package.namelist()):
            raise ValueError("Workbook has external links or advanced parts that need a native Excel preservation/recalculation workflow.")


def check_expectation_targets(snapshot, expectations, data):
    """Validate the intended checks without trusting any existing value cache."""
    formula_count = len(snapshot["formulas"])
    if formula_count and not data:
        raise ValueError("Formula workbooks require --data bindings for the raw data used to derive expectations.")
    if formula_count and not any(f"{item['sheet']}!{item['cell']}" in snapshot["formulas"] for item in expectations["cells"]):
        raise ValueError("A formula workbook must independently check at least one formula result; cover every requested/affected key result.")
    for item in expectations["cells"]:
        if item["sheet"] not in snapshot["sheets"]:
            raise ValueError(f"Expected sheet is absent: {item['sheet']}")
    for table in expectations.get("tables", []):
        key = f"{table['sheet']}!{table['name']}"
        if snapshot["tables"].get(key) != table["ref"].replace("$", "").upper():
            raise ValueError(f"Table range mismatch: {key}; expected {table['ref']}, found {snapshot['tables'].get(key)}")


def formula_coverage(snapshot, expectations, cell_results=()):
    """Count independent formula checks; quick checks have no compared values."""
    expected = {f"{item['sheet']}!{item['cell']}" for item in expectations["cells"]}
    checked = {f"{item['sheet']}!{item['cell']}": item["matches"] for item in cell_results}
    counts = ("all_formula_cells", "expected_formula_cells", "checked_formula_cells", "matched_formula_cells")
    by_sheet = {sheet: dict.fromkeys(counts, 0) for sheet in snapshot["sheets"]}
    for location in snapshot["formulas"]:
        sheet = by_sheet[location.rsplit("!", 1)[0]]
        sheet["all_formula_cells"] += 1
        sheet["expected_formula_cells"] += location in expected
        sheet["checked_formula_cells"] += location in checked
        sheet["matched_formula_cells"] += checked.get(location, False)
    total = {key: sum(sheet[key] for sheet in by_sheet.values()) for key in counts}
    for count in [total, *by_sheet.values()]:
        count["formula_cells_without_expectations"] = count["all_formula_cells"] - count["expected_formula_cells"]
        count["unchecked_formula_cells"] = count["all_formula_cells"] - count["checked_formula_cells"]
    missing_sheets = [name for name, count in by_sheet.items() if count["all_formula_cells"] and not count["expected_formula_cells"]]
    return {
        **total,
        "by_sheet": by_sheet,
        "formula_sheets_without_expectations": missing_sheets,
        "key_result_coverage": "not_established_by_counts",
        "missing_key_result_risk": (
            "Formula cells without independent expectations may include requested or affected key results; "
            "review those results and sheets before delivery."
            if total["formula_cells_without_expectations"] else
            "All formula cells have expectations; counts cannot establish their independence or coverage of non-formula key results."
            if total["all_formula_cells"] else
            "No formula cells; formula coverage does not establish coverage of non-formula key results."
        ),
    }


def quick(args):
    source = Path(args.input).resolve(strict=True)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("This checker supports ordinary .xlsx only; use a native workflow for other formats.")
    bindings = [fingerprint(path) for path in [source, args.expectations, *args.data, *args.resource]]
    check_supported_package(source)
    expectations = read_expectations(args.expectations)
    snapshot = workbook_snapshot(source)
    check_expectation_targets(snapshot, expectations, args.data)
    for binding in bindings:
        if fingerprint(binding["path"]) != binding:
            raise ValueError(f"An input changed during quick checks: {binding['path']}")
    result = {
        "status": "QUICK_CHECK_PASSED_NOT_VERIFIED", "delivery_allowed": False,
        "input": bindings[0], "sheet_count": len(snapshot["sheets"]),
        "formula_count": len(snapshot["formulas"]), "expected_cell_count": len(expectations["cells"]),
        "formula_coverage": formula_coverage(snapshot, expectations),
        "checks": "Workbook structure, supported parts, error cells, references and expectation inputs.",
        "next": "Run verify for fresh independent values, render_workbook.py for visual review, then accept.",
        "limitations": "No recalculation, expected-value comparison or visual review was performed.",
    }
    if args.out:
        output = Path(args.out).resolve()
        if str(output) in {binding["path"] for binding in bindings}:
            raise ValueError("The quick report must not overwrite an input or resource.")
        output.parent.mkdir(parents=True, exist_ok=True)
        write_json(output, result)
    print(json.dumps(result, ensure_ascii=False))
    return 0


def fresh_calculation_copy(source, target):
    """Drop old formula caches without round-tripping through openpyxl."""
    check_supported_package(source)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/workbook.xml":
                root = ET.fromstring(content, parser=ET.XMLParser(resolve_entities=False, no_network=True))
                calc = root.find(f"{{{NS}}}calcPr")
                if calc is None:
                    calc = ET.SubElement(root, f"{{{NS}}}calcPr")
                calc.attrib.update({"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"})
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                root = ET.fromstring(content, parser=ET.XMLParser(resolve_entities=False, no_network=True))
                for cell in root.iter(f"{{{NS}}}c"):
                    if cell.find(f"{{{NS}}}f") is not None:
                        for cached in list(cell):
                            if cached.tag in {f"{{{NS}}}v", f"{{{NS}}}is"}:
                                cell.remove(cached)
                        cell.attrib.pop("t", None)
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            output.writestr(item, content)


def empty_string_caches(path):
    """openpyxl reads a legitimate <v/> string result as None; inspect the XML."""
    empty = set()
    with zipfile.ZipFile(path) as package:
        parser = ET.XMLParser(resolve_entities=False, no_network=True)
        workbook = ET.fromstring(package.read("xl/workbook.xml"), parser=parser)
        rels = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"), parser=parser)
        targets = {node.get("Id"): node.get("Target") for node in rels}
        for sheet in workbook.findall(f"{{{NS}}}sheets/{{{NS}}}sheet"):
            rel = sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = targets[rel]
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            root = ET.fromstring(package.read(part), parser=parser)
            for cell in root.iter(f"{{{NS}}}c"):
                if cell.find(f"{{{NS}}}f") is None:
                    continue
                cached = cell.find(f"{{{NS}}}v")
                location = f"{sheet.get('name')}!{cell.get('r')}"
                if cached is None or (not cached.text and cell.get("t") != "str"):
                    raise ValueError(f"Formula has no fresh cached result: {location}")
                if not cached.text and cell.get("t") == "str":
                    empty.add(location)
    return empty


def check_values(candidate, expectations, formulas):
    workbook = load_workbook(candidate, data_only=True)
    empty_strings = empty_string_caches(candidate) if formulas else set()
    results = []
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e":
                    raise ValueError(f"Recalculated error {sheet.title}!{cell.coordinate}: {cell.value}")
                location = f"{sheet.title}!{cell.coordinate}"
                if location in formulas and cell.value is None and location not in empty_strings:
                    raise ValueError(f"Formula has no fresh cached result: {sheet.title}!{cell.coordinate}")
    for item in expectations["cells"]:
        if item["sheet"] not in workbook.sheetnames:
            raise ValueError(f"Expected sheet is absent: {item['sheet']}")
        actual = workbook[item["sheet"]][item["cell"]].value
        if f"{item['sheet']}!{item['cell']}" in empty_strings:
            actual = ""
        expected = item["value"]
        if isinstance(expected, (int, float)) and not isinstance(expected, bool):
            matches = isinstance(actual, (int, float)) and not isinstance(actual, bool) and math.isfinite(actual) and math.isclose(actual, expected, abs_tol=item.get("abs_tol", 0), rel_tol=item.get("rel_tol", 0))
        else:
            matches = type(actual) is type(expected) and actual == expected
        result = {**item, "actual": actual.isoformat() if hasattr(actual, "isoformat") else actual, "matches": matches}
        results.append(result)
    workbook.close()
    return results


def verify(args):
    source = Path(args.input).resolve(strict=True)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("This checker supports ordinary .xlsx only; use a native workflow for other formats.")
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=False)
    report = {"status": "UNVERIFIED", "checked_at": datetime.now(timezone.utc).isoformat(), "limitations": LIMITATIONS, "bindings": [], "failures": []}
    try:
        report["bindings"] = [fingerprint(path) for path in [source, args.expectations, __file__, *args.data, *args.resource]]
        check_supported_package(source)
        expectations = read_expectations(args.expectations)
        before = workbook_snapshot(source)
        report["structure"] = before
        report["range_review"] = expectations["range_review"]
        report["calculation_basis"] = expectations["calculation_basis"]
        formula_count = len(before["formulas"])
        report["formula_count"] = formula_count
        report["formula_coverage"] = formula_coverage(before, expectations)
        check_expectation_targets(before, expectations, args.data)
        if formula_count:
            engine = shutil.which("soffice")
            if engine is None:
                raise ValueError("Recalculation unavailable. Install libreoffice-calc (not only libreoffice-writer).")
            staged = out / "recalculation-input"
            staged.mkdir()
            fresh_calculation_copy(source, staged / source.name)
            calculated = out / "recalculated"
            calculated.mkdir()
            profile = out / "libreoffice-profile"
            command = [engine, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(calculated), str(staged / source.name)]
            process = subprocess.run(command, capture_output=True, text=True, timeout=120, check=False)
            report["engine"] = {"command": command, "returncode": process.returncode, "stdout": process.stdout, "stderr": process.stderr}
            candidate = calculated / source.name
            if process.returncode or not candidate.is_file() or not candidate.stat().st_size:
                raise ValueError("LibreOffice produced no fresh workbook. Check engine output and install libreoffice-calc; cached input values are not accepted.")
            after = workbook_snapshot(candidate)
            for key in before:
                if before[key] != after[key]:
                    raise ValueError(f"Recalculation changed workbook {key}; inspect preservation and use a compatible engine before delivery.")
            report["mode"] = "recalculated"
        else:
            candidate = source
            report["mode"] = "values_only_no_formulas"
        report["candidate"] = fingerprint(candidate)
        report["bindings"].append(report["candidate"])
        report["cell_results"] = check_values(candidate, expectations, before["formulas"])
        report["formula_coverage"] = formula_coverage(before, expectations, report["cell_results"])
        for result in report["cell_results"]:
            if not result["matches"]:
                report["failures"].append(f"{result['sheet']}!{result['cell']}: expected {result['value']!r}, got {result['actual']!r}")
        for binding in report["bindings"]:
            if fingerprint(binding["path"]) != binding:
                raise ValueError(f"An input/output changed during verification: {binding['path']}")
        if not report["failures"]:
            report["status"] = "VERIFIED"
    except (ValueError, OSError, KeyError, TypeError, subprocess.TimeoutExpired, zipfile.BadZipFile, ET.XMLSyntaxError) as error:
        report["failures"].append(str(error))
    write_json(out / "verification.json", report)
    print(json.dumps({"status": report["status"], "report": str(out / "verification.json"), "failures": report["failures"]}, ensure_ascii=False))
    return 0 if report["status"] == "VERIFIED" else 1


def load_verified_report(qa):
    qa = Path(qa).resolve(strict=True)
    report_path = qa / "verification.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    if report.get("status") != "VERIFIED" or not report.get("bindings") or not report.get("candidate"):
        raise ValueError("Workbook verification is incomplete or failed; no delivery acceptance.")
    for binding in report["bindings"]:
        if fingerprint(binding["path"]) != binding:
            raise ValueError(f"The workbook, expectations, checker or declared data/resource changed after verification: {binding['path']}. Verify into a new directory.")
    if fingerprint(report["candidate"]["path"]) != report["candidate"]:
        raise ValueError("The candidate changed after verification. Verify into a new directory.")
    return report


def validate_preview(qa, report, *, require_review=True):
    """Validate exact-candidate pages, without claiming native Excel rendering."""
    qa = Path(qa).resolve(strict=True)
    preview_dir = qa / "preview"
    inspection_path, review_path = preview_dir / "inspection.json", preview_dir / "review.json"
    if not inspection_path.is_file() or not review_path.is_file():
        raise ValueError("Visual review is missing. Run render_workbook.py --qa QA, inspect its gallery and complete preview/review.json.")
    preview = json.loads(inspection_path.read_text(encoding="utf-8"))
    review = json.loads(review_path.read_text(encoding="utf-8"))
    if preview.get("kind") != "office-workbook-preview" or preview.get("schema_version") != 1 or preview.get("status") != "needs-review":
        raise ValueError("The workbook preview is incomplete or unsupported; render it again.")
    expected_inputs = {
        "candidate": report["candidate"], "verification": fingerprint(qa / "verification.json"),
        "renderer": fingerprint(Path(__file__).with_name("render_workbook.py")),
        "gallery_builder": fingerprint(Path(__file__).with_name("page_gallery.py")),
        "pdf": fingerprint(preview_dir / "preview.pdf"),
    }
    if preview.get("inputs") != expected_inputs:
        raise ValueError("The candidate, verification, renderer or preview PDF changed; prepare and review a current preview.")
    if review.get("kind") != "office-workbook-review" or review.get("schema_version") != 1 or review.get("inspection_sha256") != fingerprint(inspection_path)["sha256"]:
        raise ValueError("Visual review does not match the current workbook preview.")
    pages = preview.get("pages", [])
    if not pages or len(pages) != preview.get("page_count") or [page["number"] for page in pages] != list(range(1, len(pages) + 1)):
        raise ValueError("Workbook preview must contain every rendered page exactly once.")
    reviewed = review.get("pages", [])
    if len(reviewed) != len(pages) or [page.get("number") for page in reviewed] != [page["number"] for page in pages]:
        raise ValueError("Visual review must cover every preview page exactly once and in order.")
    gallery_dir = preview_dir / "gallery"
    gallery = json.loads((gallery_dir / "manifest.json").read_text(encoding="utf-8"))
    if gallery.get("kind") != "office-page-gallery" or gallery.get("inspection_sha256") != fingerprint(inspection_path)["sha256"]:
        raise ValueError("The gallery does not match the current workbook preview.")
    for page, observed in zip(pages, reviewed):
        name = f"page-{page['number']:03d}.png"
        if page["image"]["path"] != name or fingerprint(preview_dir / name)["sha256"] != page["image"]["sha256"] or observed.get("image_sha256") != page["image"]["sha256"]:
            raise ValueError(f"Preview page {page['number']} changed; render and review the current workbook.")
        if fingerprint(gallery_dir / name)["sha256"] != page["image"]["sha256"]:
            raise ValueError(f"Gallery page {page['number']} no longer matches the reviewed workbook preview.")
        if require_review and (observed.get("status") != "pass" or not isinstance(observed.get("observations"), str) or not observed["observations"].strip()):
            raise ValueError(f"Page {page['number']} needs a passed visual review with observations on legibility, clipping, number formats and any charts.")
    return {"inspection": fingerprint(inspection_path), "review": fingerprint(review_path),
            "pdf": expected_inputs["pdf"], "page_count": len(pages),
            "limitations": preview.get("limitations", [])}


def accept(args):
    qa = Path(args.qa).resolve(strict=True)
    acceptance = qa / "acceptance.json"
    acceptance.unlink(missing_ok=True)
    report = load_verified_report(qa)
    visual = validate_preview(qa, report)
    result = {"status": "READY_TO_DELIVER", "accepted_at": datetime.now(timezone.utc).isoformat(), "candidate": report["candidate"], "verification": fingerprint(qa / "verification.json"), "visual_review": visual, "mode": report["mode"], "limitations": LIMITATIONS + visual["limitations"]}
    write_json(acceptance, result)
    print(json.dumps(result, ensure_ascii=False))
    return 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    for name, action, help_text in (
        ("quick", quick, "Check structure and expectation inputs; no recalculation or delivery acceptance."),
        ("verify", verify, "Check fresh calculation and independent cell expectations."),
    ):
        check = commands.add_parser(name, help=help_text)
        check.add_argument("--input", required=True)
        check.add_argument("--expectations", required=True)
        check.add_argument("--data", action="append", default=[], help="Raw data used for independent expected values; repeatable.")
        check.add_argument("--resource", action="append", default=[], help="Author/expectation scripts and other dependencies; repeatable.")
        check.add_argument("--out", required=name == "verify", help="New, nonexistent QA directory." if name == "verify" else "Optional quick-report JSON file.")
        check.set_defaults(func=action)
    acceptance = commands.add_parser("accept", help="Recheck all bindings immediately before delivery.")
    acceptance.add_argument("--qa", required=True)
    acceptance.set_defaults(func=accept)
    args = parser.parse_args()
    try:
        return args.func(args)
    except (ValueError, OSError, KeyError, TypeError, zipfile.BadZipFile, ET.XMLSyntaxError) as error:
        print(json.dumps({"status": "UNVERIFIED", "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
