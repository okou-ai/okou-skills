"""Small native authoring helpers for the blank baseline; no report outline."""

import hashlib
import json
import time
from pathlib import Path

from check_docx import check_document, is_number
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches

ASSETS = Path(__file__).resolve().parent.parent / "assets"


def literal(value):
    """User/source text is literal content, never executable Typst markup."""
    return "text(" + json.dumps(str(value), ensure_ascii=False) + ")"


class OfficeDocument:
    """Feed content once; emit editable Word and/or an independent native PDF.

    `word` and `preamble` remain accessible for language/layout adaptations.
    For richer documents, use the native libraries and blank assets directly.
    """

    def __init__(self):
        self.word = Document(ASSETS / "blank.docx")
        self.preamble = (ASSETS / "blank.typ").read_text()
        self.parts = []
        self.media = {}
        # The asset's one empty paragraph is useful when opening it in Word,
        # but should not add whitespace before the first authored element.
        paragraph = self.word.paragraphs[0]._p
        paragraph.getparent().remove(paragraph)

    def title(self, text):
        self.word.add_paragraph(str(text), "Title")
        self.parts.append(f"#doc-title[#{literal(text)}]")

    def heading(self, text, level=1):
        if level not in (1, 2, 3):
            raise ValueError("Use heading levels 1–3")
        self.word.add_paragraph(str(text), f"Heading {level}")
        self.parts.append(f"#heading(level: {level}, {literal(text)})")

    def paragraph(self, text):
        self.word.add_paragraph(str(text), "Normal")
        self.parts.append(f"#{literal(text)}")

    def list(self, items, numbered=False):
        items = list(items)
        for item in items:
            self.word.add_paragraph(
                str(item), "List Number" if numbered else "List Bullet"
            )
        native = "enum" if numbered else "list"
        self.parts.append(
            f"#{native}(" + ", ".join(literal(item) for item in items) + ")"
        )

    def table(self, headers, rows, *, numeric_columns=None, widths=None, title=None):
        """Widths are relative; numeric_columns also covers values with unit text."""
        headers = list(headers)
        rows = [list(row) for row in rows]
        count = len(headers)
        if not count or any(len(row) != count for row in rows):
            raise ValueError("Every table row must match the header length")
        widths = list(widths) if widths is not None else [1] * count
        if len(widths) != count or any(w <= 0 for w in widths):
            raise ValueError("Provide one positive width per column")
        if numeric_columns is not None and any(
            c < 0 or c >= count for c in numeric_columns
        ):
            raise ValueError("Numeric column index is outside the table")

        def numeric(column, value):
            return is_number(value) or (
                numeric_columns is not None and column in numeric_columns
            )

        if title:
            self.word.add_paragraph(str(title), "Table Title")
            self.parts.append(
                '#block(sticky: true)[#set text(size: 11pt, weight: "bold")\n'
                f"#set par(leading: 3pt)\n#{literal(title)}]"
            )
        table = self.word.add_table(rows=1, cols=count)
        table.autofit = False
        section = self.word.sections[-1]
        available = section.page_width - section.left_margin - section.right_margin
        for column, width in zip(table.columns, widths):
            column.width = int(available * width / sum(widths))
        header = table.rows[0]
        header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row_index, values in enumerate([headers] + rows):
            row = header if row_index == 0 else table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for column, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = table.columns[column].width
                paragraph = cell.paragraphs[0]
                paragraph.style = "Table Header" if row_index == 0 else "Table Text"
                paragraph.text = str(value)
                paragraph.paragraph_format.keep_with_next = row_index == 0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT
                    if row_index and numeric(column, value)
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
        cells = [
            "table.header(" + ", ".join(f"strong({literal(h)})" for h in headers) + ")"
        ]
        for row in rows:
            cells.extend(
                f"table.cell(align: {'right' if numeric(c, value) else 'start'}, {literal(value)})"
                for c, value in enumerate(row)
            )
        columns = ", ".join(f"{float(width)}fr" for width in widths)
        self.parts.append(
            f"#table(columns: ({columns},), stroke: none, inset: 4pt,\n"
            + ",\n".join(cells)
            + "\n)"
        )

    def figure(self, image, *, caption=None, source=None, width_inches=None):
        image = Path(image).resolve()
        section = self.word.sections[-1]
        available = (
            section.page_width - section.left_margin - section.right_margin
        ) / Inches(1)
        width = available if width_inches is None else width_inches
        if not 0 < width <= available:
            raise ValueError("Figure width must fit the text area")
        paragraph = self.word.add_paragraph(style="Figure")
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_with_next = bool(caption or source)
        paragraph.add_run().add_picture(str(image), width=Inches(width))
        if caption:
            paragraph = self.word.add_paragraph(str(caption), "Caption")
            paragraph.paragraph_format.keep_with_next = bool(source)
        if source:
            self.word.add_paragraph(str(source), "Source")
        data = image.read_bytes()
        name = f"media/{hashlib.sha256(data).hexdigest()}{image.suffix.lower()}"
        self.media[name] = data
        image_code = f"image({json.dumps(name)}, width: {width}in)"
        # One block keeps the image and its notes together; no forced page break.
        lines = [f"#{image_code}"]
        for note in (caption, source):
            if note:
                lines.append(
                    "#block[#set text(size: 10pt)\n"
                    f"#set par(leading: 3pt)\n#{literal(note)}]"
                )
        self.parts.append("#block(breakable: false)[\n" + "\n\n".join(lines) + "\n]")

    def save(
        self, directory, name="report", formats=("docx", "pdf"), *, baseline_check=True
    ):
        if not formats or set(formats) - {"docx", "pdf"}:
            raise ValueError("Select docx, pdf, or both")
        if Path(name).name != name:
            raise ValueError("Use a filename stem, not a path")
        directory = Path(directory).resolve()
        directory.mkdir(parents=True, exist_ok=True)
        if baseline_check:
            errors = check_document(self.word)
            if errors:
                raise ValueError("Style preflight failed:\n" + "\n".join(errors))
        results = {}
        started = time.monotonic()
        if "docx" in formats:
            target = directory / f"{name}.docx"
            self.word.save(target)
            results["docx"] = str(target)
        if "pdf" in formats:
            import typst

            for media_name, data in self.media.items():
                resource = directory / media_name
                resource.parent.mkdir(parents=True, exist_ok=True)
                resource.write_bytes(data)
            target = directory / f"{name}.typ"
            target.write_text(
                self.preamble + "\n\n" + "\n\n".join(self.parts), encoding="utf-8"
            )
            typst.compile(str(target), output=str(target.with_suffix(".pdf")))
            results["pdf"] = str(target.with_suffix(".pdf"))
        return {
            "files": results,
            "write_compile_seconds": round(time.monotonic() - started, 3),
        }
