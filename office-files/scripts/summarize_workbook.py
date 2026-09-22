#!/usr/bin/env python3
"""Build a bounded tabular summary, independent expectations and optional QA preview.

Use author_workbook.py or Python for arbitrary workbook features. This builder
consumes supplied data; it never creates records, narrative conclusions or QA
approval. Python aggregation never reads XLSX caches or evaluates Excel text.
"""

import argparse
import csv
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
import math
from pathlib import Path
import re
import subprocess
import sys
import time

from openpyxl.utils.cell import get_column_letter

from author_workbook import author


SCRIPT_DIR = Path(__file__).resolve().parent
OPS = {"count", "sum", "difference", "ratio", "change"}


def fields(value, allowed, required, where):
    if not isinstance(value, dict) or set(value) - set(allowed) or not set(required) <= set(value):
        raise ValueError(f"{where}: allowed fields {', '.join(sorted(allowed))}; required {', '.join(sorted(required))}")


def nonempty(value, where):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{where} must be nonempty text")
    return value


def identifier(value, where):
    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z_][A-Za-z_0-9]*", value):
        raise ValueError(f"{where} must be an ASCII identifier, e.g. net_cost")
    return value


def unique(values, where):
    if len({value.casefold() for value in values}) != len(values):
        raise ValueError(f"{where} must be unique, ignoring case")


def column_width(value, where):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 < value <= 255:
        raise ValueError(f"{where} must be a finite number in (0, 255]")


def sheet_name(value):
    nonempty(value, "sheet name")
    if len(value) > 31 or re.search(r"[\\/*?:\[\]]", value) or value != value.strip() or value.startswith("'") or value.endswith("'"):
        raise ValueError(f"Invalid Excel sheet name: {value!r}")
    return value


def validate_spec(spec):
    fields(spec, {"version", "filename", "theme", "source", "summaries", "dashboard", "notes"}, {"version", "source", "summaries"}, "summary spec")
    if isinstance(spec["version"], bool) or spec["version"] != 1:
        raise ValueError("Use summary spec version: 1")
    filename = spec.setdefault("filename", "summary.xlsx")
    if not isinstance(filename, str) or Path(filename).name != filename or Path(filename).suffix.lower() != ".xlsx":
        raise ValueError("filename must be an .xlsx basename, without directories")
    source = spec["source"]
    fields(source, {"sheet", "title", "columns"}, {"columns"}, "source")
    sheet_name(source.setdefault("sheet", "Data"))
    nonempty(source.setdefault("title", "Source data"), "source.title")
    columns = source["columns"]
    if not isinstance(columns, list) or not columns or len(columns) > 16384:
        raise ValueError("source.columns must be a nonempty array within Excel's column limit")
    for column in columns:
        fields(column, {"key", "header", "type", "format", "width"}, {"key", "header", "type"}, "source column")
        nonempty(column["key"], "column.key")
        nonempty(column["header"], "column.header")
        if column["type"] not in {"date", "text", "number"}:
            raise ValueError("Column type must be date, text or number")
        if "format" in column:
            nonempty(column["format"], "column.format")
        column_width(column.get("width", 24 if column["type"] == "text" else 18), "Column width")
    unique([c["key"] for c in columns], "Source column keys")
    unique([c["header"] for c in columns], "Source headers")
    kinds = {c["key"]: c["type"] for c in columns}
    summaries = spec["summaries"]
    if not isinstance(summaries, list) or not summaries:
        raise ValueError("summaries must explicitly configure at least one nonempty grouping")
    shared_widths = {}
    for summary in summaries:
        fields(summary, {"id", "sheet", "title", "group_by", "metrics", "totals", "total_label"}, {"id", "sheet", "title", "group_by", "metrics"}, "summary")
        identifier(summary["id"], "summary.id")
        sheet_name(summary["sheet"])
        nonempty(summary["title"], "summary.title")
        grouping = summary["group_by"]
        fields(grouping, {"column", "kind", "header"}, {"column", "kind"}, "group_by")
        kind = grouping["kind"]
        if kind not in {"category", "month"} or kinds.get(grouping["column"]) != ("text" if kind == "category" else "date"):
            raise ValueError("group_by needs a text column for category, or a date column for month")
        nonempty(grouping.setdefault("header", "Month" if kind == "month" else next(c["header"] for c in columns if c["key"] == grouping["column"])), "group_by.header")
        metrics = summary["metrics"]
        if not isinstance(metrics, list) or not metrics or len(metrics) > 16383:
            raise ValueError("Each summary needs a nonempty metrics array within Excel's column limit")
        earlier = set()
        for index, metric in enumerate(metrics, 2):
            if not isinstance(metric, dict) or metric.get("op") not in OPS:
                raise ValueError(f"Metric op must be one of: {', '.join(sorted(OPS))}; arbitrary expressions are unsupported")
            op = metric["op"]
            operands = {"count": set(), "sum": {"column"}, "difference": {"left", "right"}, "ratio": {"numerator", "denominator"}, "change": {"metric"}}[op]
            fields(metric, {"id", "label", "op", "format", "width"} | operands, {"id", "label", "op"} | operands, "metric")
            identifier(metric["id"], "metric.id")
            if metric["id"].casefold() in {name.casefold() for name in earlier}:
                raise ValueError(f"{summary['id']} metric ids must be unique, ignoring case")
            nonempty(metric["label"], "metric.label")
            if "format" in metric:
                nonempty(metric["format"], "metric.format")
            if "width" in metric:
                column_width(metric["width"], f"Metric {summary['id']}.{metric['id']} width")
                location = (summary["sheet"], get_column_letter(index))
                if location in shared_widths and shared_widths[location] != metric["width"]:
                    raise ValueError(f"Conflicting explicit widths for {location[0]}!{location[1]}: {shared_widths[location]} and {metric['width']}. Stacked summaries share physical columns; use one width for this column or place the summaries on separate sheets.")
                shared_widths[location] = metric["width"]
            if op == "sum" and kinds.get(metric["column"]) != "number":
                raise ValueError(f"sum metric {metric['id']} needs a declared number column")
            if op in {"difference", "ratio", "change"} and any(metric[key] not in earlier for key in operands):
                raise ValueError(f"Metric {metric['id']} must reference earlier metrics by id; cycles and expressions are unsupported")
            if op == "change" and kind != "month":
                raise ValueError("change is a month-to-month percentage change and requires month grouping")
            earlier.add(metric["id"])
        unique([grouping["header"], *[m["label"] for m in metrics]], f"{summary['id']} headers")
        if not isinstance(summary.setdefault("totals", True), bool):
            raise ValueError("totals must be a JSON boolean")
        nonempty(summary.setdefault("total_label", "Total"), "total_label")
    unique([s["id"] for s in summaries], "Summary ids")
    summary_map = {s["id"]: s for s in summaries}
    # Exact repeated summary sheet names mean stacked sections; case-only aliases do not.
    summary_sheets = list(dict.fromkeys(s["sheet"] for s in summaries))
    names = [source["sheet"], *summary_sheets]
    dashboard = spec.get("dashboard")
    if dashboard is not None:
        fields(dashboard, {"sheet", "title", "kpis", "charts"}, {"sheet", "title"}, "dashboard")
        names.append(sheet_name(dashboard["sheet"]))
        nonempty(dashboard["title"], "dashboard.title")
        if not isinstance(dashboard.get("kpis", []), list) or not isinstance(dashboard.get("charts", []), list) or not (dashboard.get("kpis") or dashboard.get("charts")):
            raise ValueError("A dashboard needs explicitly configured kpis or charts")
        for kpi in dashboard.get("kpis", []):
            fields(kpi, {"summary", "metric", "label"}, {"summary", "metric", "label"}, "KPI")
            nonempty(kpi["label"], "KPI label")
            summary = summary_map.get(kpi["summary"])
            metric = next((m for m in summary["metrics"] if m["id"] == kpi["metric"]), None) if summary else None
            if not summary or not summary["totals"] or not metric or metric["op"] == "change":
                raise ValueError("A KPI must reference a configured summary total and a metric other than change")
        for chart in dashboard.get("charts", []):
            fields(chart, {"summary", "metrics", "type", "title", "labels", "number_format"}, {"summary", "metrics", "title"}, "chart")
            nonempty(chart["title"], "chart.title")
            summary = summary_map.get(chart["summary"])
            selected = chart["metrics"]
            ids = [m["id"] for m in summary["metrics"]] if summary else []
            if not isinstance(selected, list) or not selected or any(m not in ids for m in selected):
                raise ValueError("Chart metrics must be ids from its configured summary")
            positions = [ids.index(m) for m in selected]
            if positions != list(range(positions[0], positions[0] + len(positions))):
                raise ValueError("Chart metrics must be adjacent and in summary-column order. Reorder the summary metrics, choose one metric, or use author_workbook.py for other layouts.")
            if chart.get("type", "column") not in {"column", "bar", "line"}:
                raise ValueError("Chart type must be column, bar or line")
            if not isinstance(chart.get("labels", False), bool):
                raise ValueError("Chart labels must be a JSON boolean")
            if "number_format" in chart:
                nonempty(chart["number_format"], "chart.number_format")
    notes = spec.get("notes")
    if notes is not None:
        fields(notes, {"sheet", "title", "rows", "column_widths"}, {"sheet", "title", "rows"}, "notes")
        names.append(sheet_name(notes["sheet"]))
        nonempty(notes["title"], "notes.title")
        if not isinstance(notes["rows"], list) or not notes["rows"] or any(not isinstance(row, list) or len(row) != 2 or any(not isinstance(value, str) for value in row) for row in notes["rows"]):
            raise ValueError("notes.rows must contain two-column arrays of literal strings; the first row is the header")
        for header in notes["rows"][0]:
            nonempty(header, "notes header")
        widths = notes.get("column_widths", {})
        fields(widths, {"A", "B"}, set(), "notes.column_widths")
        for column, width in widths.items():
            column_width(width, f"notes.column_widths.{column}")
    unique(names, "Source, summary and dashboard sheet names")
    return spec


def read_data(path, columns):
    keys = [column["key"] for column in columns]
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if reader.fieldnames is None or len(set(reader.fieldnames)) != len(reader.fieldnames) or set(reader.fieldnames) != set(keys):
                raise ValueError("CSV headers must match the declared column keys exactly, without duplicates")
            raw = list(reader)
    elif path.suffix.lower() == ".json":
        raw = json.loads(path.read_text(encoding="utf-8-sig"), parse_float=Decimal)
    else:
        raise ValueError("--data must be UTF-8 CSV or a JSON record array")
    if not isinstance(raw, list) or not raw or len(raw) > 1048573:
        raise ValueError("Data must be a nonempty record array within Excel's row limit")
    records = []
    for index, row in enumerate(raw, 1):
        if not isinstance(row, dict) or set(row) != set(keys):
            raise ValueError(f"Record {index}: fields must match declared column keys exactly")
        parsed = {}
        for column in columns:
            key, kind = column["key"], column["type"]
            value = row[key]
            where = f"Record {index}, {key}"
            if value is None:
                raise ValueError(f"{where}: missing values are unsupported")
            if kind == "date":
                if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                    raise ValueError(f"{where}: use an ISO YYYY-MM-DD date")
                try:
                    value = date.fromisoformat(value)
                except ValueError as error:
                    raise ValueError(f"{where}: invalid calendar date") from error
                if not date(1900, 1, 1) <= value < date(9999, 12, 1):
                    raise ValueError(f"{where}: date must be between 1900-01-01 and 9999-11-30")
            elif kind == "number":
                if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)) or isinstance(value, str) and value != value.strip():
                    raise ValueError(f"{where}: use a finite number without surrounding whitespace")
                if isinstance(value, str) and not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value):
                    raise ValueError(f"{where}: nonnumeric quantity; use a number without separators or symbols")
                try:
                    value = Decimal(str(value))
                except InvalidOperation as error:
                    raise ValueError(f"{where}: nonnumeric quantity") from error
                if not value.is_finite() or not math.isfinite(float(value)):
                    raise ValueError(f"{where}: number must be finite and representable in Excel")
            elif not isinstance(value, str):
                raise ValueError(f"{where}: text values must be strings")
            parsed[key] = value
        records.append(parsed)
    return records


def next_month(value):
    return date(value.year + (value.month == 12), value.month % 12 + 1, 1)


def group_keys(records, grouping):
    values = [row[grouping["column"]] for row in records]
    if grouping["kind"] == "month":
        current = min(values).replace(day=1)
        end = max(values).replace(day=1)
        keys = []
        while current <= end:
            keys.append(current)
            current = next_month(current)
        return keys
    if any(not value.strip() or value != value.strip() for value in values):
        raise ValueError("Category labels must be nonempty and have no surrounding whitespace")
    labels = sorted(set(values))
    unique(labels, "Category labels (Excel SUMIFS is case-insensitive)")
    return labels


def derived(metric, values, previous=None):
    """Independent arithmetic on raw-data aggregates; no Excel expressions."""
    op = metric["op"]
    if op == "difference":
        left, right = values[metric["left"]], values[metric["right"]]
        return left - right if left is not None and right is not None else None
    if op == "ratio":
        numerator, denominator = values[metric["numerator"]], values[metric["denominator"]]
        return numerator / denominator if numerator is not None and denominator not in (None, 0) else None
    if op == "change" and previous is not None:
        current, before = values[metric["metric"]], previous[metric["metric"]]
        return (current - before) / before if current is not None and before not in (None, 0) else None
    return None


def aggregate(records, summaries):
    results = {}
    for summary in summaries:
        grouping = summary["group_by"]
        keys = group_keys(records, grouping)
        buckets = {key: [] for key in keys}
        for record in records:
            value = record[grouping["column"]]
            key = value.replace(day=1) if grouping["kind"] == "month" else value
            buckets[key].append(record)
        calculated = []
        previous = None
        for key in keys:
            values = {}
            for metric in summary["metrics"]:
                if metric["op"] == "count":
                    value = Decimal(len(buckets[key]))
                elif metric["op"] == "sum":
                    value = sum((r[metric["column"]] for r in buckets[key]), Decimal(0))
                else:
                    value = derived(metric, values, previous)
                values[metric["id"]] = value
            calculated.append({"key": key, "values": values})
            previous = values
        totals = {}
        for metric in summary["metrics"]:
            if metric["op"] == "count":
                value = Decimal(len(records))
            elif metric["op"] == "sum":
                value = sum((r[metric["column"]] for r in records), Decimal(0))
            else:
                value = derived(metric, totals)
            totals[metric["id"]] = value
        results[summary["id"]] = {"groups": calculated, "totals": totals}
    return results


def scalar(value):
    if isinstance(value, Decimal):
        number = float(value)
        if not math.isfinite(number):
            raise ValueError("An aggregate is outside Excel's finite numeric range")
        return int(value) if value == value.to_integral_value() else number
    return value


def quote_sheet(name):
    return "'" + name.replace("'", "''") + "'"


def metric_format(metric):
    return metric.get("format", {"count": "integer", "sum": "decimal", "difference": "decimal", "ratio": "percent", "change": "percent"}[metric["op"]])


def derived_formula(metric, refs, previous=None):
    """Emit native Excel text; independent expectations use derived() instead."""
    op = metric["op"]
    if op == "difference":
        left, right = refs[metric["left"]], refs[metric["right"]]
        return f'=IF(OR({left}="",{right}=""),"",{left}-{right})'
    if op == "ratio":
        numerator, denominator = refs[metric["numerator"]], refs[metric["denominator"]]
        return f'=IF(OR({numerator}="",{denominator}="",{denominator}=0),"",{numerator}/{denominator})'
    if previous is None:
        return '=""'
    current, before = refs[metric["metric"]], previous[metric["metric"]]
    return f'=IF(OR({current}="",{before}="",{before}=0),"",({current}-{before})/{before})'


def build_layout(spec, records, aggregates):
    source = spec["source"]
    last_data_row = len(records) + 3
    last_source_column = get_column_letter(len(source["columns"]))
    workbook = {"version": 1, "theme": spec.get("theme", {}), "sheets": []}
    expected = {"calculation_basis": "", "range_review": "", "cells": [], "tables": [{"sheet": source["sheet"], "name": "InputRecords", "ref": f"A3:{last_source_column}{last_data_row}"}]}
    coverage = {}
    def expect(sheet, address, value, formula=False):
        # Empty literal strings serialize as blank cells; a formula returning
        # an empty string has a distinct, freshly calculated string cache.
        checked_value = "" if formula and value is None else None if not formula and value == "" else scalar(value)
        item = {"sheet": sheet, "cell": address, "value": checked_value}
        if isinstance(item["value"], (int, float)):
            item.update(abs_tol=1e-9, rel_tol=1e-12)
        expected["cells"].append(item)
        coverage[sheet]["expected_cells"] += 1

    def new_sheet(name):
        sheet = {"name": name, "cells": {}, "column_widths": {}, "merges": []}
        coverage[name] = {"expected_cells": 0, "formula_cells": 0, "expected_formula_cells": 0, "typed_date_literals": 0}
        workbook["sheets"].append(sheet)
        return sheet

    def put(sheet, address, value, *, formula=None, style=None, fmt=None):
        entry = {"formula": formula} if formula is not None else {"value": scalar(value), "type": "text"} if isinstance(value, str) else {"value": scalar(value)}
        if style:
            entry["style"] = style
        if fmt:
            entry["format"] = fmt
        sheet["cells"][address] = entry
        expect(sheet["name"], address, value, formula=formula is not None)
        if formula is not None:
            coverage[sheet["name"]]["formula_cells"] += 1
            coverage[sheet["name"]]["expected_formula_cells"] += 1

    data_sheet = new_sheet(source["sheet"])
    data_sheet.update(start_cell="A3", columns=[], rows=[], table="InputRecords", freeze="A4", print={"title_rows": "1:3"})
    put(data_sheet, "A1", source["title"], style="title")
    if len(source["columns"]) > 1:
        data_sheet["merges"].append(f"A1:{last_source_column}1")
    ranges = {}
    for index, column in enumerate(source["columns"], 1):
        letter = get_column_letter(index)
        fmt = column.get("format", {"date": "date", "text": "text", "number": "decimal"}[column["type"]])
        definition = {"header": column["header"], "width": column.get("width", 24 if column["type"] == "text" else 18), "format": fmt}
        if column["type"] == "text" and index > 1 and source["columns"][index - 2]["type"] in {"date", "number"}:
            definition.update(horizontal="left", indent=1)
        data_sheet["columns"].append(definition)
        ranges[column["key"]] = f"{quote_sheet(source['sheet'])}!${letter}$4:${letter}${last_data_row}"
        put(data_sheet, f"{letter}3", column["header"], style="header")
    for row_number, record in enumerate(records, 4):
        row = []
        for index, column in enumerate(source["columns"], 1):
            value = record[column["key"]]
            if column["type"] == "date":
                row.append({"value": value.isoformat(), "type": "date"})
                coverage[source["sheet"]]["typed_date_literals"] += 1
            else:
                row.append({"value": scalar(value), **({"type": "text"} if column["type"] == "text" else {})})
                expect(source["sheet"], f"{get_column_letter(index)}{row_number}", value)
        data_sheet["rows"].append(row)
    layouts, sheets, cursors = {}, {}, {}
    reviews = [f"{source['sheet']}!A3:{last_source_column}{last_data_row}: InputRecords covers all {len(records)} raw records; dates are typed ISO date literals."]
    for summary in spec["summaries"]:
        name = summary["sheet"]
        sheet = sheets.setdefault(name, None)
        if sheet is None:
            sheet = sheets[name] = new_sheet(name)
            sheet["freeze"] = "B4"
            sheet["column_widths"]["A"] = 26
        title_row = cursors.get(name, 1)
        header, first = title_row + 2, title_row + 3
        groups = aggregates[summary["id"]]["groups"]
        last = first + len(groups) - 1
        total = last + 1 if summary["totals"] else None
        end_column = get_column_letter(len(summary["metrics"]) + 1)
        put(sheet, f"A{title_row}", summary["title"], style="title")
        sheet["merges"].append(f"A{title_row}:{end_column}{title_row}")
        put(sheet, f"A{header}", summary["group_by"]["header"], style="header")
        metric_columns = {}
        for index, metric in enumerate(summary["metrics"], 2):
            letter = get_column_letter(index)
            metric_columns[metric["id"]] = letter
            if "width" in metric:
                sheet["column_widths"][letter] = metric["width"]
            else:
                sheet["column_widths"].setdefault(letter, 20)
            put(sheet, f"{letter}{header}", metric["label"], style="header")
        previous = None
        for row, group in enumerate(groups, first):
            key = group["key"]
            put(sheet, f"A{row}", key.strftime("%Y-%m") if isinstance(key, date) else key)
            group_range = ranges[summary["group_by"]["column"]]
            if isinstance(key, date):
                after = next_month(key)
                criteria = f'{group_range},">="&DATE({key.year},{key.month},1),{group_range},"<"&DATE({after.year},{after.month},1)'
            else:
                # SUMIFS interprets criteria metacharacters even through a cell reference.
                escaped = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($A{row},"~","~~"),"*","~*"),"?","~?")'
                criteria = f'{group_range},"="&{escaped}'
            refs = {key: f"{column}{row}" for key, column in metric_columns.items()}
            for metric in summary["metrics"]:
                op = metric["op"]
                if op == "count":
                    formula = f"=COUNTIFS({criteria})"
                elif op == "sum":
                    formula = f"=SUMIFS({ranges[metric['column']]},{criteria})"
                else:
                    formula = derived_formula(metric, refs, previous)
                put(sheet, refs[metric["id"]], group["values"][metric["id"]], formula=formula, fmt=metric_format(metric))
            previous = refs
        if total:
            put(sheet, f"A{total}", summary["total_label"], style="total")
            refs = {key: f"{column}{total}" for key, column in metric_columns.items()}
            for metric in summary["metrics"]:
                column = metric_columns[metric["id"]]
                formula = f"=SUM({column}{first}:{column}{last})" if metric["op"] in {"count", "sum"} else derived_formula(metric, refs)
                put(sheet, refs[metric["id"]], aggregates[summary["id"]]["totals"][metric["id"]], formula=formula, style="total", fmt=metric_format(metric))
        cursors[name] = (total or last) + 4
        layouts[summary["id"]] = {"sheet": name, "header": header, "first": first, "last": last, "total": total, "columns": metric_columns}
        reviews.append(f"{summary['id']}: {name}!A{header}:{end_column}{last} covers every {summary['group_by']['kind']} group; total row {total if total else 'omitted'}. Every grouped metric, derived value and emitted total has an independent expectation.")
    dashboard = spec.get("dashboard")
    if dashboard:
        sheet = new_sheet(dashboard["sheet"])
        sheet["column_widths"] = {get_column_letter(i): 12 for i in range(1, 12)}
        sheet["column_widths"]["D"] = 20
        sheet["column_widths"]["F"] = 3
        sheet["merges"].append("A1:K1")
        sheet["print"] = {"orientation": "landscape", "title_rows": "1:1"}
        put(sheet, "A1", dashboard["title"], style="title")
        summary_map = {s["id"]: s for s in spec["summaries"]}
        for row, kpi in enumerate(dashboard.get("kpis", []), 3):
            layout = layouts[kpi["summary"]]
            metric = next(m for m in summary_map[kpi["summary"]]["metrics"] if m["id"] == kpi["metric"])
            sheet["merges"].append(f"A{row}:C{row}")
            put(sheet, f"A{row}", kpi["label"])
            target = f"{quote_sheet(layout['sheet'])}!{layout['columns'][kpi['metric']]}{layout['total']}"
            put(sheet, f"D{row}", aggregates[kpi["summary"]]["totals"][kpi["metric"]], formula=f'=IF({target}="","",{target})', style="total", fmt=metric_format(metric))
        chart_start = max(7, len(dashboard.get("kpis", [])) + 5)
        for index, chart in enumerate(dashboard.get("charts", [])):
            layout = layouts[chart["summary"]]
            top = chart_start + (index // 2) * 17
            anchor = f"{'A' if index % 2 == 0 else 'G'}{top}:{'E' if index % 2 == 0 else 'K'}{top + 14}"
            data = f"{layout['columns'][chart['metrics'][0]]}{layout['header']}:{layout['columns'][chart['metrics'][-1]]}{layout['last']}"
            categories = f"A{layout['first']}:A{layout['last']}"
            definition = {"type": chart.get("type", "column"), "title": chart["title"], "data_sheet": layout["sheet"], "data": data, "categories": categories, "anchor": anchor, "labels": chart.get("labels", False)}
            if "number_format" in chart:
                definition["number_format"] = chart["number_format"]
            sheet.setdefault("charts", []).append(definition)
            reviews.append(f"Chart {index + 1}: {layout['sheet']}!{data}; categories {categories}; excludes totals; dashboard slot {anchor}.")
            if index % 2 == 0 and index >= 2:
                sheet["print"].setdefault("break_rows", []).append(top - 1)
    notes = spec.get("notes")
    if notes:
        sheet = new_sheet(notes["sheet"])
        sheet["column_widths"] = {"A": 26, "B": 80, **notes.get("column_widths", {})}
        sheet["merges"] = ["A1:B1"]
        sheet["print"] = {"orientation": "landscape", "title_rows": "1:3"}
        put(sheet, "A1", notes["title"], style="title")
        for row_number, row in enumerate(notes["rows"], 3):
            for index, value in enumerate(row, 1):
                put(sheet, f"{get_column_letter(index)}{row_number}", value, style="header" if row_number == 3 else "normal")
        reviews.append(f"{notes['sheet']}!A3:B{len(notes['rows']) + 2}: every supplied literal note/header has an expectation; no claims were generated.")
    expected["calculation_basis"] = ("Python Decimal aggregation of the supplied raw records: counts and sums are computed directly per group and again over all records for totals. Differences and ratios use those aggregates; total ratios divide total operands (never average group ratios). Month changes compare consecutive calendar-month aggregates; first month, missing/zero denominators and total change are blank. This code path does not read workbook caches or evaluate generated Excel formulas.")
    expected["range_review"] = " ".join(reviews)
    if any(c["formula_cells"] != c["expected_formula_cells"] for c in coverage.values()):
        raise ValueError("Internal error: generated formula coverage is incomplete")
    return workbook, expected, coverage


def fingerprint(path):
    path = Path(path).resolve(strict=True)
    return {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def write_json(path, content):
    path.write_text(json.dumps(content, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def run(data_path, spec_path, out, prepare_review=False):
    started = time.perf_counter()
    data_path, spec_path = Path(data_path).resolve(strict=True), Path(spec_path).resolve(strict=True)
    out = Path(out).resolve()
    if out.exists():
        raise ValueError("--out must be a new directory; retain previous source/QA evidence")
    bindings = [fingerprint(p) for p in (data_path, spec_path, __file__, SCRIPT_DIR / "author_workbook.py")]
    spec = validate_spec(json.loads(spec_path.read_text(encoding="utf-8")))
    records = read_data(data_path, spec["source"]["columns"])
    timings = {"read_validate_seconds": round(time.perf_counter() - started, 6)}
    phase = time.perf_counter()
    results = aggregate(records, spec["summaries"])
    timings["aggregate_seconds"] = round(time.perf_counter() - phase, 6)
    phase = time.perf_counter()
    workbook_spec, expectations, coverage = build_layout(spec, records, results)
    timings["layout_seconds"] = round(time.perf_counter() - phase, 6)
    out.mkdir(parents=True, exist_ok=False)
    generated_spec, expectation_path = out / "workbook-spec.json", out / "expectations.json"
    output = out / spec["filename"]
    write_json(generated_spec, workbook_spec)
    write_json(expectation_path, expectations)
    phase = time.perf_counter()
    author(workbook_spec, output)
    timings["author_seconds"] = round(time.perf_counter() - phase, 6)
    manifest = {"schema_version": 1, "kind": "office-tabular-summary", "created_at": datetime.now(timezone.utc).isoformat(), "status": "AUTHORED_NOT_VERIFIED", "inputs": bindings, "outputs": [fingerprint(p) for p in (output, generated_spec, expectation_path)], "source_record_count": len(records), "summary_ids": [s["id"] for s in spec["summaries"]], "sheet_coverage": coverage, "timings": timings, "timing_scope": "Local Python/subprocess wall time only; excludes model configuration/composition, tool dispatch, human review and upload. It does not establish an end-to-end model speedup.", "limitations": ["Only the explicitly configured summaries/dashboard are created; custom formulas, grouping expressions, pivots and native Excel preservation require the lower-level workflow.", "All emitted summary and KPI formulas have raw-data-derived expected values. Raw date cells are strictly parsed and typed; the numeric checker does not compare date literals individually.", "Groups and source/chart ranges are a snapshot of the supplied rows. Rebuild and repeat QA after adding records or categories.", "Generated expectations and range descriptions are not visual review or independent human endorsement. Review all preview pages and run accept before delivery."]}
    response = {"status": manifest["status"], "workbook": str(output), "spec": str(generated_spec), "expectations": str(expectation_path), "manifest": str(out / "manifest.json")}
    try:
        if prepare_review:
            qa = out / "qa"
            check = [sys.executable, str(SCRIPT_DIR / "check_workbook.py")]
            common = ["--input", str(output), "--expectations", str(expectation_path), "--data", str(data_path)]
            for resource in (spec_path, SCRIPT_DIR / "summarize_workbook.py", SCRIPT_DIR / "author_workbook.py", generated_spec):
                common += ["--resource", str(resource)]
            stages = [("quick", check + ["quick", *common, "--out", str(out / "quick.json")]), ("verify", check + ["verify", *common, "--out", str(qa)]), ("preview", [sys.executable, str(SCRIPT_DIR / "render_workbook.py"), "--qa", str(qa)])]
            for name, command in stages:
                phase = time.perf_counter()
                try:
                    process = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
                except (OSError, subprocess.TimeoutExpired) as error:
                    manifest["status"] = f"FAILED_{name.upper()}"
                    manifest["failed_stage"] = {"name": name, "error": str(error)}
                    raise
                finally:
                    timings[f"{name}_seconds"] = round(time.perf_counter() - phase, 6)
                if process.returncode:
                    manifest["status"] = f"FAILED_{name.upper()}"
                    manifest["failed_stage"] = {"name": name, "returncode": process.returncode, "stdout": process.stdout, "stderr": process.stderr}
                    raise ValueError(f"{name} failed; inspect {out / 'manifest.json'}. {process.stdout.strip()} {process.stderr.strip()}")
            verification = json.loads((qa / "verification.json").read_text(encoding="utf-8"))
            formulas = set(verification["structure"]["formulas"])
            for name, sheet_coverage in coverage.items():
                sheet_coverage["verified_formula_cells"] = sum(item["matches"] and f"{name}!{item['cell']}" in formulas for item in verification["cell_results"] if item["sheet"] == name)
            manifest["status"] = response["status"] = "VERIFIED_NEEDS_VISUAL_REVIEW"
            response.update(candidate=verification["candidate"]["path"], qa=str(qa), gallery=str(qa / "preview/gallery/index.html"), review=str(qa / "preview/review.json"))
        for binding in bindings:
            if fingerprint(binding["path"]) != binding:
                manifest["status"] = "INPUT_CHANGED"
                raise ValueError(f"Input/resource changed during generation: {binding['path']}; rebuild into a new directory")
    finally:
        timings["total_seconds"] = round(time.perf_counter() - started, 6)
        write_json(out / "manifest.json", manifest)
    return response


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", required=True, help="UTF-8 CSV or JSON array; supplied source data only")
    parser.add_argument("--spec", required=True, help="Summary JSON; start with assets/summary-spec.json")
    parser.add_argument("--out", required=True, help="New output directory; no source or prior evidence is overwritten")
    parser.add_argument("--prepare-review", action="store_true", help="Also run fixed quick, fresh Calc verification and page preview; never approve or accept")
    args = parser.parse_args()
    try:
        print(json.dumps(run(args.data, args.spec, args.out, args.prepare_review), ensure_ascii=False))
        return 0
    except (ValueError, TypeError, KeyError, OSError, OverflowError, subprocess.TimeoutExpired) as error:
        print(json.dumps({"status": "SUMMARY_FAILED", "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
