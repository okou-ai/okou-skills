"""Regression checks for workbook results and artifact-bound acceptance.

Run: python3 -m unittest discover -s office-files/tests -p 'test_check_workbook.py'
The Calc cases are real conversions and skip only when Calc is unavailable.
"""

import argparse
import contextlib
import importlib.util
import io
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest import mock
import zipfile

from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "check_workbook.py"
SPEC = importlib.util.spec_from_file_location("check_workbook", SCRIPT)
checker = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(checker)
HAS_CALC = shutil.which("soffice") and Path("/etc/libreoffice/registry/calc.xcd").exists()


class WorkbookChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "sales.xlsx"
        self.data = self.root / "raw.json"
        self.expectations = self.root / "expectations.json"
        self.author = self.root / "author.py"
        self.author.write_text("# fixture author\n", encoding="utf-8")
        self.data.write_text(json.dumps([3000, 48, 380]), encoding="utf-8")

    def fixture(self, wrong_reference=False, formula=True, structures=False):
        raw = json.loads(self.data.read_text())
        workbook = Workbook()
        data = workbook.active
        data.title = "Data"
        data.append(["Month", "Quantity"])
        for month, value in enumerate(raw, 1):
            data.append([month, value])
        data.append(["Total", "=SUM(B2:B4)" if formula else sum(raw)])
        summary = workbook.create_sheet("Summary")
        summary["A2"] = "Total"
        summary["B2"] = "=Data!B4" if wrong_reference else ("=Data!B5" if formula else sum(raw))
        if structures:
            data.add_table(Table(displayName="Sales", ref="A1:B4"))
            workbook.defined_names.add(DefinedName("Quantities", attr_text="Data!$B$2:$B$4"))
            chart = BarChart()
            chart.add_data(Reference(data, min_col=2, min_row=1, max_row=4), titles_from_data=True)
            chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=4))
            summary.add_chart(chart, "D2")
        workbook.save(self.source)
        expected = {
            "calculation_basis": "Python sum of raw.json quantities; no workbook values are used.",
            "range_review": "Reviewed Data rows 2:4, total B5 and Summary B2 after adding month 3; summary must target B5.",
            "cells": [{"sheet": "Summary", "cell": "B2", "value": sum(raw), "abs_tol": 0}],
        }
        if structures:
            expected["tables"] = [{"sheet": "Data", "name": "Sales", "ref": "A1:B4"}]
        self.expectations.write_text(json.dumps(expected), encoding="utf-8")

    def verify(self, name="qa", data=True):
        args = argparse.Namespace(input=str(self.source), expectations=str(self.expectations), data=[str(self.data)] if data else [], resource=[str(self.author)], out=str(self.root / name))
        with contextlib.redirect_stdout(io.StringIO()):
            code = checker.verify(args)
        return code, json.loads((self.root / name / "verification.json").read_text())

    def accept(self, name="qa"):
        with contextlib.redirect_stdout(io.StringIO()):
            return checker.accept(argparse.Namespace(qa=str(self.root / name)))

    @unittest.skipUnless(HAS_CALC, "libreoffice-calc not installed")
    def test_real_calc_retains_formulas_tables_chart_and_name(self):
        self.fixture(structures=True)
        original = self.source.read_bytes()
        code, report = self.verify()
        self.assertEqual(code, 0, report["failures"])
        self.assertEqual(report["cell_results"][0]["actual"], 3428)
        self.assertEqual(report["mode"], "recalculated")
        self.assertEqual(self.source.read_bytes(), original)
        self.assertNotEqual(report["candidate"]["path"], str(self.source))
        self.assertEqual(self.accept(), 0)
        candidate = Path(report["candidate"]["path"])
        candidate.write_bytes(candidate.read_bytes() + b"\n")
        with self.assertRaisesRegex(ValueError, "changed after verification"):
            self.accept()

    @unittest.skipUnless(HAS_CALC, "libreoffice-calc not installed")
    def test_wrong_cross_sheet_reference_380_cannot_pass_expected_3428(self):
        self.fixture()
        # Start before the new month, then reproduce openpyxl's stale reference.
        workbook = load_workbook(self.source)
        data = workbook["Data"]
        data.delete_rows(4)
        data["B4"] = "=SUM(B2:B3)"
        workbook["Summary"]["B2"] = "=Data!B4"
        data.insert_rows(4)
        data["A4"] = 3
        data["B4"] = 380
        data["B5"] = "=SUM(B2:B4)"
        workbook.save(self.source)
        code, report = self.verify()
        self.assertEqual(code, 1)
        self.assertEqual(report["status"], "UNVERIFIED")
        self.assertEqual(report["cell_results"][0]["actual"], 380)
        self.assertIn("expected 3428, got 380", report["failures"][0])
        with self.assertRaisesRegex(ValueError, "incomplete or failed"):
            self.accept()

    @unittest.skipUnless(HAS_CALC, "libreoffice-calc not installed")
    def test_stale_cached_values_are_removed_and_recalculated(self):
        self.fixture()
        temporary = self.root / "stale.xlsx"
        with zipfile.ZipFile(self.source) as original, zipfile.ZipFile(temporary, "w") as output:
            for item in original.infolist():
                content = original.read(item.filename)
                if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                    root = ET.fromstring(content)
                    for cell in root.iter(f"{{{checker.NS}}}c"):
                        if cell.find(f"{{{checker.NS}}}f") is not None:
                            cached = cell.find(f"{{{checker.NS}}}v")
                            if cached is None:
                                cached = ET.SubElement(cell, f"{{{checker.NS}}}v")
                            cached.text = "999999"
                    content = ET.tostring(root)
                output.writestr(item, content)
        temporary.replace(self.source)
        self.assertEqual(load_workbook(self.source, data_only=True)["Summary"]["B2"].value, 999999)
        code, report = self.verify()
        self.assertEqual(code, 0, report["failures"])
        self.assertEqual(report["cell_results"][0]["actual"], 3428)

    @unittest.skipUnless(HAS_CALC, "libreoffice-calc not installed")
    def test_fresh_empty_string_formula_is_not_a_missing_cache(self):
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["Summary"]["C2"] = '=""'
        workbook.save(self.source)
        expected = json.loads(self.expectations.read_text())
        expected["cells"].append({"sheet": "Summary", "cell": "C2", "value": ""})
        self.expectations.write_text(json.dumps(expected))
        code, report = self.verify()
        self.assertEqual(code, 0, report["failures"])
        self.assertEqual(report["cell_results"][1]["actual"], "")
        self.assertEqual(self.accept(), 0)

    @unittest.skipUnless(HAS_CALC, "libreoffice-calc not installed")
    def test_calculated_error_cannot_pass_even_when_not_an_expected_cell(self):
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["Summary"]["C2"] = "=1/0"
        workbook.save(self.source)
        code, report = self.verify()
        self.assertEqual(code, 1)
        self.assertIn("#DIV/0!", report["failures"][0])

    @unittest.skipUnless(HAS_CALC, "libreoffice-calc not installed")
    def test_chinese_sheet_names_survive_calc_quote_normalization(self):
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["Data"].title = "数据"
        workbook["Summary"]["B2"] = "=数据!B5"
        workbook.save(self.source)
        code, report = self.verify()
        self.assertEqual(code, 0, report["failures"])
        self.assertEqual(report["cell_results"][0]["actual"], 3428)

    def test_value_only_checks_and_changed_bindings_invalidate_acceptance(self):
        self.fixture(formula=False)
        for binding in (self.source, self.data, self.author, self.expectations):
            with self.subTest(binding=binding.name):
                name = "qa-" + binding.suffix[1:] + binding.stem
                code, report = self.verify(name)
                self.assertEqual(code, 0, report["failures"])
                self.assertEqual(report["mode"], "values_only_no_formulas")
                self.assertEqual(self.accept(name), 0)
                original = binding.read_bytes()
                binding.write_bytes(original + b"\n")
                with self.assertRaisesRegex(ValueError, "changed after verification"):
                    self.accept(name)
                self.assertFalse((self.root / name / "acceptance.json").exists())
                binding.write_bytes(original)

    def test_empty_expectations_and_undeclared_raw_data_fail(self):
        self.fixture()
        code, report = self.verify("qa-no-data", data=False)
        self.assertEqual(code, 1)
        self.assertIn("--data", report["failures"][0])
        expected = json.loads(self.expectations.read_text())
        expected["cells"] = []
        self.expectations.write_text(json.dumps(expected))
        code, report = self.verify("qa-empty")
        self.assertEqual(code, 1)
        self.assertIn("nonempty independently derived", report["failures"][0])

    def test_missing_engine_never_accepts_cached_values_or_old_output(self):
        self.fixture()
        with mock.patch.object(checker.shutil, "which", return_value=None):
            code, report = self.verify()
        self.assertEqual(code, 1)
        self.assertIn("libreoffice-calc", report["failures"][0])
        with self.assertRaises(FileExistsError):
            self.verify()

    def test_advanced_parts_are_not_silently_resaved_by_calc(self):
        self.fixture()
        with zipfile.ZipFile(self.source, "a") as package:
            package.writestr("xl/pivotTables/pivotTable1.xml", "<pivotTableDefinition/>")
        original = self.source.read_bytes()
        with self.assertRaisesRegex(ValueError, "native Excel"):
            checker.fresh_calculation_copy(self.source, self.root / "staged.xlsx")
        self.assertEqual(self.source.read_bytes(), original)

    def test_error_and_missing_sheet_references_fail(self):
        for formula in ("=#REF!", "=Missing!B5"):
            with self.subTest(formula=formula):
                self.fixture()
                workbook = load_workbook(self.source)
                workbook["Summary"]["B2"] = formula
                workbook.save(self.source)
                code, report = self.verify("qa-" + str(len(list(self.root.glob("qa-*")))))
                self.assertEqual(code, 1)
                self.assertTrue(report["failures"])

    def test_numeric_tolerance_and_false_value_do_not_silently_match(self):
        self.fixture(formula=False)
        workbook = load_workbook(self.source)
        workbook["Summary"]["B2"] = 3428.0001
        workbook.save(self.source)
        expected = json.loads(self.expectations.read_text())
        expected["cells"][0]["abs_tol"] = 0.001
        self.expectations.write_text(json.dumps(expected))
        self.assertEqual(self.verify("qa-tolerance")[0], 0)
        expected["cells"][0]["value"] = False
        self.expectations.write_text(json.dumps(expected))
        self.assertEqual(self.verify("qa-false")[0], 1)


if __name__ == "__main__":
    unittest.main()
