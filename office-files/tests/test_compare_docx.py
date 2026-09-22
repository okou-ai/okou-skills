"""Regression tests for DOCX preservation; requires the documented python-docx.

Run: python3 -m unittest discover -s office-files/tests -p 'test_compare_docx.py'
"""

import base64
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "compare_docx.py"
SPEC = importlib.util.spec_from_file_location("compare_docx", SCRIPT)
comparison = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(comparison)
NS = {"w": comparison.W, "r": comparison.R, "rel": comparison.REL}
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")


def rewrite(source, target, mutation):
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    mutation(parts)
    with ZipFile(target, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def modify_xml(parts, name, mutation):
    root = etree.fromstring(parts[name])
    mutation(root)
    parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def remove_related_part(parts, name):
    parts.pop(name)
    relative = name.removeprefix("word/")
    modify_xml(parts, "word/_rels/document.xml.rels", lambda root: [root.remove(node) for node in list(root) if node.get("Target") == relative])
    modify_xml(parts, "[Content_Types].xml", lambda root: [root.remove(node) for node in list(root) if node.get("PartName") == "/" + name])
    parts.pop("word/_rels/" + Path(name).name + ".rels", None)


class DocxPreservationTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.original, self.edited = self.directory / "original.docx", self.directory / "edited.docx"
        self.policy = self.directory / "changes.json"
        document = Document()
        document.styles.add_style("Comment Reference", WD_STYLE_TYPE.CHARACTER)
        document.add_paragraph("Revenue: 1000")
        paragraph = document.add_paragraph()
        run = paragraph.add_run("Contract terms")
        run.bold = True
        document.add_comment(run, text="Retain the contract wording.", author="Reviewer", initials="R")
        paragraph.add_run("    ")
        paragraph = document.add_paragraph()
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to("https://example.com/contract", RT.HYPERLINK, is_external=True))
        link_run, link_text = OxmlElement("w:r"), OxmlElement("w:t")
        link_text.text = "Read contract"
        link_run.append(link_text)
        hyperlink.append(link_run)
        paragraph._p.append(hyperlink)
        paragraph = document.add_paragraph()
        revision = OxmlElement("w:ins")
        for name, value in (("id", "41"), ("author", "Reviewer"), ("date", "2026-09-22T00:00:00Z")):
            revision.set(qn("w:" + name), value)
        revision_run, revision_text = OxmlElement("w:r"), OxmlElement("w:t")
        revision_text.text = "Pending addition"
        revision_run.append(revision_text)
        revision.append(revision_run)
        paragraph._p.append(revision)
        document.sections[0].header.paragraphs[0].text = "Company header"
        footer = document.sections[0].footer.paragraphs[0]
        for kind in ("begin", "instruction", "separate", "result", "end"):
            field_run = footer.add_run()
            if kind == "instruction":
                node = OxmlElement("w:instrText")
                node.text = " PAGE "
            elif kind == "result":
                node = OxmlElement("w:t")
                node.text = "1"
            else:
                node = OxmlElement("w:fldChar")
                node.set(qn("w:fldCharType"), kind)
            field_run._r.append(node)
        document.add_picture(BytesIO(PNG))
        initial = self.directory / "initial.docx"
        document.save(initial)
        rewrite(initial, self.original, lambda parts: parts.update({
            "word/_rels/header1.xml.rels": f'<Relationships xmlns="{comparison.REL}"/>'.encode()
        }))

    def edit_revenue(self):
        document = Document(self.original)
        document.paragraphs[0].runs[0].text = "Revenue: 1200"
        document.save(self.edited)

    def approve_text(self):
        report = comparison.compare(self.original, self.edited)
        changes = [{**{key: item[key] for key in comparison.CHANGE_KEYS},
                    "reason": "Change the requested revenue figure from 1000 to 1200."}
                   for item in report["findings"]
                   if item["code"] == "text_changed" and item["part"] == "word/document.xml"]
        self.assertEqual(len(changes), 1)
        self.policy.write_text(json.dumps({"schema_version": 1, "changes": changes}))
        return comparison.compare(self.original, self.edited, self.policy)

    def test_text_edit_preserves_actual_features(self):
        self.edit_revenue()
        report = self.approve_text()
        self.assertEqual(report["status"], "PASS", report["findings"])
        self.assertTrue(comparison.validate_binding(report))
        scopes = {item["scope"] for item in report["features"]}
        self.assertTrue({"header", "footer", "fields", "comments", "comment_anchors", "hyperlinks", "revisions", "media", "drawings"} <= scopes)
        self.assertTrue(all(item["before_sha256"] == item["after_sha256"] for item in report["features"]))

    def test_text_change_without_intent_is_blocked(self):
        self.edit_revenue()
        report = comparison.compare(self.original, self.edited)
        self.assertEqual(report["status"], "BLOCKED")
        self.assertEqual([item["code"] for item in report["findings"] if item["severity"] == "blocker"], ["text_changed"])

    def test_empty_relationship_part_removal_is_informational(self):
        Document(self.original).save(self.edited)
        report = comparison.compare(self.original, self.edited)
        self.assertEqual(report["status"], "PASS", report["findings"])
        findings = [item for item in report["findings"] if item["code"] == "empty_relationship_part"]
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["severity"], "info")

    def test_unrequested_feature_losses_block_even_with_approved_body_edit(self):
        self.edit_revenue()
        self.assertEqual(self.approve_text()["status"], "PASS")

        def drop_comments(parts):
            remove_related_part(parts, "word/comments.xml")
            modify_xml(parts, "word/document.xml", lambda root: [node.getparent().remove(node) for node in root.xpath("//w:commentRangeStart | //w:commentRangeEnd | //w:commentReference", namespaces=NS)])

        def accept_revision(parts):
            def unwrap(root):
                for node in root.xpath("//w:ins", namespaces=NS):
                    parent, position = node.getparent(), node.getparent().index(node)
                    for child in list(node):
                        parent.insert(position, child)
                        position += 1
                    parent.remove(node)
            modify_xml(parts, "word/document.xml", unwrap)

        def drop_header(parts):
            remove_related_part(parts, "word/header1.xml")
            modify_xml(parts, "word/document.xml", lambda root: [node.getparent().remove(node) for node in root.xpath("//w:headerReference", namespaces=NS)])

        def drop_link(parts):
            def unwrap(root):
                for node in root.xpath("//w:hyperlink", namespaces=NS):
                    parent, position = node.getparent(), node.getparent().index(node)
                    for child in list(node):
                        parent.insert(position, child)
                        position += 1
                    parent.remove(node)
            modify_xml(parts, "word/document.xml", unwrap)

        def drop_field(parts):
            modify_xml(parts, "word/footer1.xml", lambda root: [node.getparent().remove(node) for node in root.xpath("//w:fldChar | //w:instrText", namespaces=NS)])

        for scope, mutation in (("comments", drop_comments), ("revisions", accept_revision),
                                ("header", drop_header), ("hyperlinks", drop_link), ("fields", drop_field)):
            with self.subTest(feature=scope):
                broken = self.directory / (scope + ".docx")
                rewrite(self.edited, broken, mutation)
                report = comparison.compare(self.original, broken, self.policy)
                self.assertEqual(report["status"], "BLOCKED")
                self.assertTrue(any(item["code"] == "feature_changed" and item["scope"] == scope and item["severity"] == "blocker" for item in report["findings"]), report["findings"])

    def test_namespace_and_relationship_id_normalization_is_not_corruption(self):
        def mutation(parts):
            def renumber(root):
                for node in root:
                    if node.get("Type", "").endswith("/hyperlink"):
                        node.set("Id", "rIdRenumbered")
            modify_xml(parts, "word/_rels/document.xml.rels", renumber)
            root = etree.fromstring(parts["word/document.xml"])
            root.xpath("//w:hyperlink", namespaces=NS)[0].set(qn("r:id"), "rIdRenumbered")
            parts["word/document.xml"] = etree.tostring(root, pretty_print=True)
        rewrite(self.original, self.edited, mutation)
        report = comparison.compare(self.original, self.edited)
        self.assertEqual(report["status"], "PASS", report["findings"])
        self.assertEqual(comparison.compare(self.original, self.edited, strict_bytes=True)["status"], "BLOCKED")

    def test_dangling_relationship_is_an_unapprovable_integrity_failure(self):
        rewrite(self.original, self.edited, lambda parts: parts.pop("word/header1.xml"))
        report = comparison.compare(self.original, self.edited)
        invalid = [item for item in report["findings"] if item["code"] == "invalid_reference"]
        self.assertTrue(invalid)
        changes = [{**{key: item[key] for key in comparison.CHANGE_KEYS}, "reason": "This test attempts to waive an invalid header reference."} for item in invalid]
        self.policy.write_text(json.dumps({"schema_version": 1, "changes": changes}))
        report = comparison.compare(self.original, self.edited, self.policy)
        self.assertTrue(all(item["severity"] == "blocker" for item in report["findings"] if item["code"] == "invalid_reference"))

    def test_changed_original_edited_or_policy_invalidates_binding(self):
        self.edit_revenue()
        report = self.approve_text()
        for path in (self.original, self.edited, self.policy):
            with self.subTest(input=path.name):
                original_bytes = path.read_bytes()
                path.write_bytes(original_bytes + b" ")
                with self.assertRaisesRegex(ValueError, "changed"):
                    comparison.validate_binding(report)
                path.write_bytes(original_bytes)
        self.assertTrue(comparison.validate_binding(report))

    def test_old_policy_does_not_approve_a_different_text_edit(self):
        self.edit_revenue()
        self.approve_text()
        document = Document(self.edited)
        document.paragraphs[0].runs[0].text = "Revenue: 9999"
        document.save(self.edited)
        report = comparison.compare(self.original, self.edited, self.policy)
        self.assertEqual(report["status"], "BLOCKED")
        self.assertTrue(any(item["code"] == "unused_policy_change" for item in report["findings"]))

    def test_wildcard_or_blanket_policy_is_rejected(self):
        self.edit_revenue()
        self.approve_text()
        policy = json.loads(self.policy.read_text())
        for key, value in (("part", "word/*"), ("reason", "all changes approved")):
            with self.subTest(key=key):
                edited_policy = json.loads(json.dumps(policy))
                edited_policy["changes"][0][key] = value
                self.policy.write_text(json.dumps(edited_policy))
                with self.assertRaises(ValueError):
                    comparison.compare(self.original, self.edited, self.policy)

    def test_concise_task_specific_chinese_reason_is_accepted(self):
        self.edit_revenue()
        self.approve_text()
        policy = json.loads(self.policy.read_text())
        policy["changes"][0]["reason"] = "收入改为1200"
        self.policy.write_text(json.dumps(policy))
        self.assertEqual(comparison.compare(self.original, self.edited, self.policy)["status"], "PASS")

    def test_cli_error_removes_an_old_pass_report(self):
        self.edit_revenue()
        report = self.approve_text()
        output = self.directory / "comparison.json"
        output.write_text(json.dumps(report))
        self.policy.write_text("not valid JSON")
        process = subprocess.run([sys.executable, str(SCRIPT), str(self.original), str(self.edited),
                                  "--out", str(output), "--changes", str(self.policy)], capture_output=True, text=True)
        self.assertEqual(process.returncode, 2)
        self.assertFalse(output.exists())

    def test_mixed_content_text_loss_is_detected(self):
        source = self.directory / "mixed-original.docx"
        rewrite(self.original, source, lambda parts: parts.update({
            "customXml/item1.xml": b"<mixed>Hello <bold>important</bold> remaining content</mixed>"
        }))
        rewrite(source, self.edited, lambda parts: parts.update({
            "customXml/item1.xml": b"<mixed>Hello <bold>important</bold></mixed>"
        }))
        report = comparison.compare(source, self.edited)
        self.assertEqual(report["status"], "BLOCKED")
        self.assertTrue(any(item["code"] == "text_changed" and item["part"] == "customXml/item1.xml" for item in report["findings"]))


if __name__ == "__main__":
    unittest.main()
