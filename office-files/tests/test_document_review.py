"""Regression checks for required expectations, batch review and native metadata."""

import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import zipfile

import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import check_document as checks
from compare_docx import compare
from page_gallery import build_gallery
from render_document import docx_language, render


class DocumentReviewTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.pdf = self.root / "sample.pdf"
        self.qa = self.root / "qa"
        self.expected = self.root / "expectations.json"

    def write_pdf(self, count=1):
        with pymupdf.open() as document:
            for number in range(1, count + 1):
                page = document.new_page(width=360 if number % 2 else 520, height=480)
                page.insert_text((30, 50), f"Heading {number}", fontsize=20)
                page.insert_text((30, 80), f"Opening paragraph {number} with a known total of 3428.")
            document.save(self.pdf)

    def write_expectations(self, value=None):
        value = value if value is not None else {
            "required_text": ["known total of 3428"],
            "same_page": [{"first": "Heading 1", "second": "Opening paragraph 1", "reason": "Heading introduces opening text"}],
        }
        self.expected.write_text(json.dumps(value), encoding="utf-8")

    def inspect(self):
        return checks.inspect(self.pdf, self.qa, expectations=self.expected)

    def synthetic_review(self):
        # These fixture acknowledgements only exercise acceptance conditions.
        review_path = self.qa / "review.json"
        review = json.loads(review_path.read_text())
        for page in review["pages"]:
            for item in page["criteria"].values():
                item.update(status="pass", observations="Synthetic fixture review for gate regression testing.")
        for finding in review["warning_acknowledgements"]:
            finding["observations"] = "Synthetic fixture warning acknowledgement."
        checks.write_json(review_path, review)

    def test_empty_expectations_cannot_pass_even_with_complete_review(self):
        self.write_pdf()
        self.write_expectations({"required_text": [], "same_page": []})
        report = self.inspect()
        self.assertEqual({"empty_required_text", "empty_same_page"},
                         {item["code"] for item in report["findings"] if item["severity"] == "blocker"})
        self.synthetic_review()
        with self.assertRaisesRegex(ValueError, "Unresolved machine blockers"):
            checks.accept(self.qa)

    def test_missing_expectation_file_cannot_pass(self):
        self.write_pdf()
        report = checks.inspect(self.pdf, self.qa)
        self.assertIn("missing_expectations", [finding["code"] for finding in report["findings"]])
        self.synthetic_review()
        with self.assertRaisesRegex(ValueError, "no bound expectations"):
            checks.accept(self.qa)

    def test_explicit_nonapplicability_and_known_text_can_pass(self):
        self.write_pdf()
        self.write_expectations({"required_text": ["3428"], "same_page": [],
                                 "not_applicable": {"same_page": "Single-page fixture; no pagination requirement"}})
        self.inspect()
        self.synthetic_review()
        self.assertEqual("READY_TO_DELIVER", checks.accept(self.qa)["status"])

    def test_wrong_expected_number_is_blocked(self):
        self.write_pdf()
        self.write_expectations({"required_text": ["total of 380"], "same_page": [],
                                 "not_applicable": {"same_page": "One-page test fixture"}})
        report = self.inspect()
        self.assertIn("required_text_missing", [item["code"] for item in report["findings"]])

    def test_gallery_batches_all_pages_without_marking_review_passed(self):
        self.write_pdf(7)
        self.write_expectations()
        self.inspect()
        before = checks.sha256(self.qa / "review.json")
        build_gallery(self.qa, columns=2, per_sheet=4)
        self.assertEqual(before, checks.sha256(self.qa / "review.json"))
        manifest = json.loads((self.qa / "gallery/manifest.json").read_text())
        self.assertEqual(list(range(1, 8)), [page["number"] for page in manifest["pages"]])
        self.assertEqual([(1, 4), (5, 7)], [(sheet["first_page"], sheet["last_page"]) for sheet in manifest["contact_sheets"]])
        for page in manifest["pages"]:
            self.assertEqual(page["sha256"], checks.sha256(self.qa / "gallery" / page["path"]))
        for sheet in manifest["contact_sheets"]:
            self.assertGreater(pymupdf.Pixmap(str(self.qa / "gallery" / sheet["path"])).width, 0)
        with self.assertRaisesRegex(ValueError, "needs a passed review"):
            checks.accept(self.qa)

    def test_changed_page_is_rejected_by_gallery_and_acceptance(self):
        self.write_pdf()
        self.write_expectations()
        self.inspect()
        self.synthetic_review()
        checks.accept(self.qa)
        with (self.qa / "page-001.png").open("ab") as image:
            image.write(b"changed")
        with self.assertRaisesRegex(ValueError, "changed after inspection"):
            build_gallery(self.qa)
        with self.assertRaisesRegex(ValueError, "image changed"):
            checks.accept(self.qa)
        self.assertFalse((self.qa / "acceptance.json").exists())

    def test_gallery_regeneration_removes_old_pages(self):
        self.write_pdf(2)
        self.write_expectations()
        self.inspect()
        self.pdf.unlink()
        self.write_pdf(1)
        self.inspect()
        self.assertFalse((self.qa / "gallery/page-002.png").exists())

    def test_native_word_uses_regional_language_metadata(self):
        path = self.root / "regional.docx"
        xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
               '<w:body><w:p><w:r><w:rPr><w:lang w:val="en-US" w:eastAsia="zh-TW"/></w:rPr>'
               '<w:t>中文內容</w:t></w:r></w:p></w:body></w:document>')
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("word/document.xml", xml)
        self.assertEqual(("zh-TW", "document-metadata"), docx_language(path))

    def test_native_pdf_language_and_override_are_recorded_without_rewriting(self):
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((40, 50), "Declared document language")
            document.xref_set_key(document.pdf_catalog(), "Lang", "(fr-FR)")
            document.save(self.pdf)
        result = render(self.pdf, self.root / "rendered")
        self.assertEqual(("fr-FR", "document-metadata"), (result["language"], result["language_source"]))
        self.assertEqual(checks.sha256(self.pdf), checks.sha256(result["outputs"]["pdf"]["path"]))
        result = render(self.pdf, self.root / "override", lang="zh-CN")
        self.assertEqual(("zh-CN", "override"), (result["language"], result["language_source"]))

    def test_word_comparison_binds_original_and_exact_edited_candidate(self):
        from docx import Document

        original = self.root / "original.docx"
        edited = self.root / "edited.docx"
        document = Document()
        document.add_paragraph("An unchanged document for comparison binding.")
        document.save(original)
        shutil.copyfile(original, edited)
        comparison = self.root / "comparison.json"
        checks.write_json(comparison, compare(original, edited))
        report = {"inputs": {"docx": checks.fingerprint(edited)}, "findings": []}
        checks.bind_docx_comparison(report, comparison)
        self.assertFalse(report["findings"])
        self.assertEqual(checks.sha256(original), report["inputs"]["comparison_original"]["sha256"])

        different = self.root / "different.docx"
        document.add_paragraph("This candidate was never compared.")
        document.save(different)
        report = {"inputs": {"docx": checks.fingerprint(different)}, "findings": []}
        checks.bind_docx_comparison(report, comparison)
        self.assertEqual("invalid_docx_comparison", report["findings"][0]["code"])

        with original.open("ab") as stream:
            stream.write(b"changed")
        report = {"inputs": {"docx": checks.fingerprint(edited)}, "findings": []}
        checks.bind_docx_comparison(report, comparison)
        self.assertEqual("invalid_docx_comparison", report["findings"][0]["code"])


if __name__ == "__main__":
    unittest.main()
