"""Round-trip tests against real Pandoc DOCX files, not generated XML snapshots."""

from __future__ import annotations

import base64
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml.ns import qn
import pypandoc

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from document_style import build_reference, polish_document


class DocumentStyleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pandoc = pypandoc.get_pandoc_path()

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def render(self, source: str, lang="en", input_format="markdown", name="document"):
        reference = self.root / f"{name}-reference.docx"
        output = self.root / f"{name}.docx"
        build_reference(self.pandoc, reference, lang)
        subprocess.run(
            [self.pandoc, "-f", input_format, "-t", "docx", "--reference-doc", str(reference),
             "--resource-path", str(self.root), "-o", str(output)],
            input=source.encode(), check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        )
        return output

    def test_reference_keeps_pandoc_styles_and_readable_page_defaults(self):
        path = self.root / "reference.docx"
        build_reference(self.pandoc, path, "en-US")
        document = Document(path)
        default = Document(BytesIO(subprocess.run(
            [self.pandoc, "--print-default-data-file", "reference.docx"],
            check=True, stdout=subprocess.PIPE,
        ).stdout))
        self.assertTrue({s.style_id for s in default.styles} <= {s.style_id for s in document.styles})
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, delta=.1)
        self.assertAlmostEqual(section.page_height.mm, 297, delta=.1)
        self.assertAlmostEqual(section.left_margin.mm, 28, delta=.1)
        self.assertAlmostEqual(section.right_margin.mm, 28, delta=.1)
        self.assertEqual(document.styles["Normal"].font.size.pt, 11)
        self.assertEqual(document.styles["Compact"].font.size.pt, 11)
        self.assertGreater(document.styles["Heading 1"].font.size, document.styles["Heading 2"].font.size)
        self.assertGreater(document.styles["Heading 2"].font.size, document.styles["Heading 3"].font.size)
        for name in ("Caption", "Image Caption", "Footnote Text", "Footnote Block Text"):
            self.assertGreaterEqual(document.styles[name].font.size.pt, 9)
        self.assertTrue(document.styles["Keep with Next"].paragraph_format.keep_with_next)
        fields = section.footer._element.xpath(".//w:fldSimple")
        self.assertEqual([field.get(qn("w:instr")) for field in fields], ["PAGE"])
        self.assertEqual(document.styles["Footer"].style_id, "Footer")
        self.assertEqual(document.styles["Header"].style_id, "Header")
        self.assertFalse(section.header.paragraphs[0].text)

    def test_polish_preserves_text_code_links_and_images_and_pairs_caption(self):
        image = self.root / "pixel.png"
        image.write_bytes(base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a0xkAAAAASUVORK5CYII="
        ))
        output = self.render(
            "# Overview\n\nExact values: 003.40, 42.5%, AB-123. 中文。\n\n"
            "- A list with `run --id=001` and **emphasis**.\n\n"
            "Read [the source](https://example.com/source?q=001).\n\n"
            "The next figure illustrates the observation.\n\n"
            "![A caption with 42.5%.](pixel.png){width=4cm}\n\n"
            "| Category | Value |\n|---|---:|\n| One | 001 |\n",
            lang="zh-CN",
        )
        original = Document(output)
        before = [node.text for node in original._element.iter(qn("w:t"))]
        before_code = [r.text for p in original.paragraphs for r in p.runs if r.style.name == "Verbatim Char"]
        with zipfile.ZipFile(output) as archive:
            media = {name: archive.read(name) for name in archive.namelist() if name.startswith("word/media/")}
        polish_document(output, "zh-CN")
        document = Document(output)
        self.assertEqual([node.text for node in document._element.iter(qn("w:t"))], before)
        self.assertEqual([r.text for p in document.paragraphs for r in p.runs if r.style.name == "Verbatim Char"], before_code)
        self.assertTrue(any(rel.target_ref == "https://example.com/source?q=001" for rel in document.part.rels.values()))
        with zipfile.ZipFile(output) as archive:
            self.assertEqual({name: archive.read(name) for name in media}, media)
        drawing_index = next(i for i, paragraph in enumerate(document.paragraphs) if paragraph._p.xpath(".//w:drawing"))
        figure = document.paragraphs[drawing_index]
        caption = document.paragraphs[drawing_index + 1]
        self.assertTrue(figure.paragraph_format.keep_with_next)
        self.assertTrue(caption.paragraph_format.keep_together)
        self.assertFalse(caption.paragraph_format.keep_with_next)
        self.assertTrue(document.paragraphs[drawing_index - 1].paragraph_format.keep_with_next)
        polish_document(output, "zh-CN")
        self.assertEqual([node.text for node in Document(output)._element.iter(qn("w:t"))], before)

    def test_default_table_has_balanced_columns_repeating_header_and_separate_font(self):
        output = self.render(
            "# Values\n\n- Body-size list item.\n\nHere are the values.\n\n"
            "| Category | Count | Change |\n|---|---:|---:|\n"
            "| A particularly long descriptive label that should wrap | 120 | 5.3% |\n"
            "| Short label | 80 | 0.2% |\n"
        )
        polish_document(output, "en")
        document = Document(output)
        table = document.tables[0]
        section = document.sections[0]
        width = section.page_width - section.left_margin - section.right_margin
        columns = [column.width for column in table.columns]
        self.assertAlmostEqual(sum(columns), width, delta=1905)
        self.assertLessEqual(max(columns) / sum(columns), .701)
        self.assertFalse(table.autofit)
        self.assertEqual(table.rows[0]._tr.trPr.find(qn("w:tblHeader")).get(qn("w:val")), "1")
        self.assertTrue(all(p.paragraph_format.keep_with_next for cell in table.rows[0].cells for p in cell.paragraphs))
        self.assertEqual(table.rows[1]._tr.trPr.find(qn("w:cantSplit")).get(qn("w:val")), "1")
        self.assertEqual(table.rows[1].cells[0].paragraphs[0].style.font.size.pt, 10)
        list_paragraph = next(p for p in document.paragraphs if "Body-size list" in p.text)
        self.assertEqual(list_paragraph.style.font.size.pt, 11)
        self.assertEqual(table.rows[1].cells[1].paragraphs[0].alignment, 2)  # right alignment survives

    def test_short_numeric_column_does_not_take_label_or_description_space(self):
        output = self.render(
            "| 项目 | 数量 | 说明 |\n|---|---:|---|\n"
            "| 数据核对 | 24 | 检查原始值与最终交付内容是否一致。 |\n"
            "| 文档导出 | 12 | 同一份可编辑文档生成固定版面 PDF。 |\n"
            "| 页面检查 | 36 | 逐页检查层级、分页、表格和可读性。 |\n", lang="zh-CN"
        )
        polish_document(output, "zh-CN")
        columns = [column.width for column in Document(output).tables[0].columns]
        self.assertLess(columns[1], columns[0] * .7)
        self.assertLess(columns[1] / sum(columns), .16)
        self.assertGreater(columns[2] / sum(columns), .6)
        self.assertGreater(columns[1] / 12700, 2 * 10 + 10)  # CJK header plus cell padding fits.

    def test_real_multiblock_figure_has_defined_borderless_layout_style(self):
        image = self.root / "pixel.png"
        image.write_bytes(base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a0xkAAAAASUVORK5CYII="
        ))
        ast = json.loads(subprocess.run(
            [self.pandoc, "-f", "markdown", "-t", "json"],
            input=b"![Figure caption](pixel.png){width=4cm}\n", stdout=subprocess.PIPE, check=True,
        ).stdout)
        figure = ast["blocks"][0]
        self.assertEqual(figure["t"], "Figure")
        figure["c"][2].append({"t": "Para", "c": [{"t": "Str", "c": "Supplemental figure text."}]})
        output = self.render(json.dumps(ast), input_format="json")
        original = Document(output)
        table = original.tables[0]
        self.assertEqual(table._tbl.tblPr.tblStyle.val, "FigureTable")
        widths = [column.width for column in table.columns]
        text = [element.text for element in original._element.iter(qn("w:t"))]
        polish_document(output, "en")
        document = Document(output)
        table = document.tables[0]
        self.assertEqual(document.styles["FigureTable"].base_style.style_id, "Table")
        self.assertEqual([column.width for column in table.columns], widths)
        self.assertIsNone(table._tbl.tblPr.find(qn("w:tblBorders")))
        self.assertEqual([element.text for element in document._element.iter(qn("w:t"))], text)
        self.assertTrue(table._tbl.xpath(".//w:drawing"))

    def test_explicit_column_widths_are_preserved(self):
        source = "| Label | Details |\n|---|---|\n| A | Short |\n"
        ast = json.loads(subprocess.run(
            [self.pandoc, "-f", "markdown", "-t", "json"], input=source.encode(),
            stdout=subprocess.PIPE, check=True,
        ).stdout)
        ast["blocks"][0]["c"][2] = [
            [{"t": "AlignLeft"}, {"t": "ColWidth", "c": .27}],
            [{"t": "AlignLeft"}, {"t": "ColWidth", "c": .73}],
        ]
        output = self.render(json.dumps(ast), input_format="json")
        original = Document(output).tables[0]
        widths = [column.width for column in original.columns]
        table_width = dict(original._tbl.tblPr.find(qn("w:tblW")).attrib)
        polish_document(output, "en")
        table = Document(output).tables[0]
        self.assertEqual([column.width for column in table.columns], widths)
        self.assertEqual(dict(table._tbl.tblPr.find(qn("w:tblW")).attrib), table_width)

    def test_very_tall_rows_can_split_and_headerless_tables_stay_headerless(self):
        output = self.render("| Label | Detail |\n|---|---|\n| A | " + "A long explanation. " * 700 + " |\n")
        document = Document(output)
        table = document.tables[0]
        header = table.rows[0]._tr.trPr.find(qn("w:tblHeader"))
        header.getparent().remove(header)
        document.save(output)
        polish_document(output, "en")
        table = Document(output).tables[0]
        self.assertIsNone(table.rows[0]._tr.trPr.find(qn("w:tblHeader")))
        self.assertEqual(table.rows[1]._tr.trPr.find(qn("w:cantSplit")).get(qn("w:val")), "0")
        self.assertFalse(table.rows[1].cells[0].paragraphs[0].style.paragraph_format.keep_with_next)

    def test_language_fonts_and_rtl_preserve_ltr_code_and_numbers(self):
        for lang, east_asia, complex_font, rtl in (
            ("zh-CN", "Noto Sans CJK SC", "Noto Sans", False),
            ("zh-Hant-TW", "Noto Sans CJK TC", "Noto Sans", False),
            ("ja-JP", "Noto Sans CJK JP", "Noto Sans", False),
            ("ko-KR", "Noto Sans CJK KR", "Noto Sans", False),
            ("ar-SA", "Noto Sans CJK SC", "Noto Sans Arabic", True),
            ("he-IL", "Noto Sans CJK SC", "Noto Sans Hebrew", True),
            ("hi-IN", "Noto Sans CJK SC", "Noto Sans Devanagari", False),
        ):
            with self.subTest(lang=lang):
                output = self.render("# عنوان\n\nمرحبا `run --id=001` 123.45\n", lang=lang, name=lang)
                polish_document(output, lang)
                document = Document(output)
                normal = document.styles["Normal"]
                fonts = normal.element.rPr.find(qn("w:rFonts"))
                self.assertEqual(fonts.get(qn("w:eastAsia")), east_asia)
                self.assertEqual(fonts.get(qn("w:cs")), complex_font)
                language_property = normal.element.rPr.find(qn("w:lang"))
                expected_latin_language = "en-US" if lang.split("-")[0] in {"zh", "ja", "ko"} else lang
                self.assertEqual(language_property.get(qn("w:val")), expected_latin_language)
                self.assertEqual(normal.element.pPr.find(qn("w:bidi")).get(qn("w:val")), "1" if rtl else "0")
                self.assertEqual(document.paragraphs[1].text, "مرحبا run --id=001 123.45")
                if rtl:
                    runs = document.paragraphs[1].runs
                    code = next(run for run in runs if run.style.name == "Verbatim Char")
                    self.assertFalse(code.font.rtl)
                    self.assertTrue(runs[0].font.rtl)
                    self.assertFalse(runs[-1].font.rtl)
                    # Bidi must precede spacing/alignment in the property list.
                    names = [node.tag for node in normal.element.pPr]
                    self.assertLess(names.index(qn("w:bidi")), names.index(qn("w:spacing")))
                    self.assertLess(names.index(qn("w:bidi")), names.index(qn("w:jc")))

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice Writer is required for the real wrapping regression")
    def test_cjk_document_keeps_english_words_whole_in_real_pdf(self):
        import pymupdf

        output = self.render(
            "---\nlang: zh-CN\n---\n\n"
            "# Appendix A — English terms and handover checklist\n\n"
            "All data in this document are illustrative. No task status, date, amount, or project reference "
            "describes an actual business event.\n\n"
            "- **Owner:** The person handing over the work and explaining the current operating context.\n"
            "- **Receiver:** The person accepting the work and verifying that it can be performed independently.\n"
            "- **Approver:** The person authorized to accept remaining risks or require corrective work.\n"
            "- **Evidence:** A record that supports a task status, such as a verified link, a completed trial, or an explicit decision.\n",
            lang="zh-CN",
        )
        original = [node.text for node in Document(output)._element.iter(qn("w:t"))]
        polish_document(output, "zh-CN")
        self.assertEqual([node.text for node in Document(output)._element.iter(qn("w:t"))], original)
        profile = (self.root / "lo-profile").as_uri()
        subprocess.run([shutil.which("soffice"), f"-env:UserInstallation={profile}",
                        "--headless", "--convert-to", "pdf", "--outdir", str(self.root), str(output)],
                       check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90)
        with pymupdf.open(output.with_suffix(".pdf")) as pdf:
            self.assertEqual(len(pdf), 1)
            text = pdf[0].get_text()
        for word in ("checklist", "reference", "operating", "performed", "corrective", "completed"):
            self.assertIn(word, text)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice Writer is required for real RTL geometry")
    def test_rtl_title_and_headings_align_right_without_excess_title_gap(self):
        import pymupdf

        source = (Path(__file__).parent / "fixtures" / "rtl.md").read_text(encoding="utf-8")
        output = self.render(source, lang="ar")
        original = [node.text for node in Document(output)._element.iter(qn("w:t"))]
        polish_document(output, "ar")
        document = Document(output)
        self.assertEqual([node.text for node in document._element.iter(qn("w:t"))], original)
        spacing = document.styles["Normal"].element.pPr.find(qn("w:spacing"))
        self.assertEqual(spacing.get(qn("w:lineRule")), "atLeast")
        profile = (self.root / "rtl-lo-profile").as_uri()
        subprocess.run([shutil.which("soffice"), f"-env:UserInstallation={profile}",
                        "--headless", "--convert-to", "pdf", "--outdir", str(self.root), str(output)],
                       check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90)
        with pymupdf.open(output.with_suffix(".pdf")) as pdf:
            self.assertEqual(len(pdf), 1)
            spans = [span for block in pdf[0].get_text("dict")["blocks"] if "lines" in block
                     for line in block["lines"] for span in line["spans"]]
            right_edge = pdf[0].rect.width - document.sections[0].right_margin.pt
            headings = [span for span in spans if span["size"] >= 15]
            self.assertGreaterEqual(len(headings), 4)
            for span in headings:
                self.assertAlmostEqual(span["bbox"][2], right_edge, delta=2)
            self.assertLess(headings[1]["origin"][1] - headings[0]["origin"][1], 55)


if __name__ == "__main__":
    unittest.main()
