"""Native-file regressions and cache behavior through the shipped entry points."""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt
from PIL import Image
from pypdf import PdfReader

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
from office import OfficeDocument


class OfficeFilesTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)

    def command(self, script, *args, env=None):
        return subprocess.run(
            [sys.executable, str(SCRIPTS / script), *map(str, args)],
            check=False,
            text=True,
            capture_output=True,
            env=env,
            timeout=90,
        )

    def create_report(self, text="Shared report"):
        report = OfficeDocument()
        report.title(text)
        report.paragraph('Literal #panic("must not execute") and [brackets].')
        report.list(["First action", "Second action"], numbered=True)
        report.table(
            ["Supplier", "Cost", "Availability"],
            [["A", "590,000", "99.93%"], ["B", "602,000", "99.96%"]],
        )
        image = self.root / "chart.png"
        Image.new("RGB", (800, 320), "white").save(image)
        report.figure(image, caption="Figure caption", source="Source: test fixture")
        report.save(self.root)
        return self.root / "report.docx", self.root / "report.pdf"

    def test_shared_native_content_and_baseline(self):
        docx, pdf = self.create_report()
        word = Document(docx)
        self.assertEqual(len(word.tables), 1)
        self.assertEqual(len(word.inline_shapes), 1)
        self.assertAlmostEqual(
            word.inline_shapes[0].width / word.inline_shapes[0].height, 2.5
        )
        pdf_text = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
        word_text = "\n".join(p.text for p in word.paragraphs)
        for text in (
            "Shared report",
            "First action",
            "Figure caption",
            '#panic("must not execute")',
        ):
            self.assertIn(text, word_text)
            self.assertIn(text, pdf_text)
        for text in ("590,000", "602,000", "99.93%", "99.96%"):
            self.assertIn(text, pdf_text)
            self.assertIn(
                text,
                " ".join(
                    cell.text for row in word.tables[0].rows for cell in row.cells
                ),
            )
        check = self.command("check_docx.py", docx)
        self.assertEqual(check.returncode, 0, check.stdout + check.stderr)

    def test_known_style_regressions_are_rejected_without_repair(self):
        docx, _ = self.create_report()
        word = Document(docx)
        numbered = next(p for p in word.paragraphs if p.style.name == "List Number")
        numbered.runs[0].font.size = Pt(11)
        word.tables[0].cell(0, 0).paragraphs[0].runs[0].bold = False
        word.tables[0].cell(1, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        word.tables[0].cell(2, 1).paragraphs[0].runs[0].font.size = Pt(9)
        picture = next(p for p in word.paragraphs if p._p.xpath(".//w:drawing"))
        picture.paragraph_format.line_spacing = Pt(12)
        picture.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        word.save(docx)
        original = docx.read_bytes()
        check = self.command("check_docx.py", docx)
        self.assertEqual(check.returncode, 1)
        errors = "\n".join(json.loads(check.stdout)["errors"])
        for message in (
            "must use 12 pt",
            "must use 11 pt",
            "header is not bold",
            "numeric cell is not right aligned",
            "image paragraph uses exact",
        ):
            self.assertIn(message, errors)
        self.assertEqual(docx.read_bytes(), original)

    def test_final_file_render_reuse_and_invalidation(self):
        docx, pdf = self.create_report()
        cache = self.root / "review"
        first = self.command("review.py", docx, pdf, "--cache", cache)
        self.assertEqual(first.returncode, 0, first.stdout + first.stderr)
        records = json.loads(first.stdout)
        self.assertTrue(
            all(r["pages"] > 0 and r["visual_review"] == "required" for r in records)
        )
        before = {p: Path(p).stat().st_mtime_ns for r in records for p in r["images"]}
        second = self.command("review.py", docx, pdf, "--cache", cache)
        self.assertEqual(second.returncode, 0, second.stdout + second.stderr)
        self.assertTrue(all(r["cached"] for r in json.loads(second.stdout)))
        self.assertTrue(
            all(Path(p).stat().st_mtime_ns == value for p, value in before.items())
        )
        self.create_report("Changed content")
        third = self.command("review.py", docx, pdf, "--cache", cache)
        self.assertEqual(third.returncode, 0, third.stdout + third.stderr)
        self.assertTrue(all(not r["cached"] for r in json.loads(third.stdout)))
        self.assertTrue(
            all(
                "Changed content"
                in PdfReader(r["rendered_pdf"]).pages[0].extract_text()
                for r in json.loads(third.stdout)
            )
        )

    def test_identical_failed_word_conversion_is_not_retried(self):
        docx, _ = self.create_report()
        binary = self.root / "bin"
        binary.mkdir()
        attempts = self.root / "attempts"
        converter = binary / "soffice"
        converter.write_text(
            '#!/bin/sh\nif [ "$1" = "--version" ]; then echo fixture; else echo attempted >> "$OFFICE_TEST_ATTEMPTS"; echo "load failed"; fi\n'
        )
        converter.chmod(0o755)
        env = {
            **os.environ,
            "PATH": str(binary) + os.pathsep + os.environ["PATH"],
            "OFFICE_TEST_ATTEMPTS": str(attempts),
        }
        for index in range(2):
            outcome = self.command(
                "review.py", docx, "--cache", self.root / "review", env=env
            )
            self.assertEqual(outcome.returncode, 1)
            result = json.loads(outcome.stdout)[0]
            self.assertEqual(result["visual_review"], "unverified")
            self.assertEqual(result["cached"], index == 1)
        self.assertEqual(attempts.read_text().splitlines(), ["attempted"])

    def test_equal_files_share_cache_when_batched(self):
        _, pdf = self.create_report()
        copy = self.root / "copy.pdf"
        shutil.copyfile(pdf, copy)
        output = self.command("review.py", pdf, copy, "--cache", self.root / "review")
        self.assertEqual(output.returncode, 0, output.stdout + output.stderr)
        first, second = json.loads(output.stdout)
        self.assertEqual(first["images"], second["images"])
        self.assertEqual(sorted([first["cached"], second["cached"]]), [False, True])

    def test_environment_returns_usable_python_and_reuses_writer_probe(self):
        for index in range(2):
            result = self.command(
                "office_env.py",
                "--formats",
                "docx",
                "pdf",
                "--cache",
                self.root / "environment",
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            capabilities = json.loads(result.stdout)
            self.assertEqual(capabilities["missing"], [])
            self.assertTrue(capabilities["writer"]["available"])
            self.assertEqual(capabilities["writer"]["cached"], index == 1)
            probe = subprocess.run(
                [capabilities["python"], "-c", "import docx, typst, pypdf"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(probe.returncode, 0, probe.stderr)


if __name__ == "__main__":
    unittest.main()
