"""Real exports cover content, source preservation and final-file relationships.

Run: python3 -m unittest discover -s office-files/tests -p 'test_*.py'
Requires the skill requirements, fonts and LibreOffice Writer.
"""
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

import docx
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from render_document import (apply_design_components, find_pandoc, inline_text,
                             language_of, render, source_expectations,
                             validate_docx, validate_source)
from document_style import build_reference

FIXTURES = Path(__file__).parent / 'fixtures'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


class SourceTests(unittest.TestCase):
    def test_pinned_bundled_pandoc_precedes_an_unrelated_system_binary(self):
        import pypandoc
        expected = Path(pypandoc.__file__).parent / 'files' / 'pandoc'
        with patch('render_document.shutil.which', return_value='/unrelated/pandoc'):
            self.assertEqual(find_pandoc(), str(expected))

    def test_language_explicit_overrides_and_regional_metadata(self):
        ast = {'meta': {'lang': {'t': 'MetaString', 'c': 'zh-Hant'}}, 'blocks': []}
        self.assertEqual(language_of(ast, None), 'zh-Hant')
        self.assertEqual(language_of(ast, 'ja-JP'), 'ja-JP')

    def test_inline_links_do_not_add_url_to_heading_expectations(self):
        text = [{'t': 'Link', 'c': [['', [], []], [{'t': 'Str', 'c': 'Values'}], ['https://example.com', '']]}]
        self.assertEqual(inline_text(text), 'Values')
        ast = {'blocks': [{'t': 'Header', 'c': [1, ['', [], []], text]},
                          {'t': 'Para', 'c': [{'t': 'Str', 'c': 'Opening text'}]}]}
        self.assertEqual(source_expectations(ast)['same_page'][0]['second'], 'Opening text')

    def test_silent_raw_content_loss_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'would be lost'):
            validate_source({'blocks': [{'t': 'RawBlock', 'c': ['html', '<table>Important</table>']}]})
        validate_source({'blocks': [{'t': 'RawBlock', 'c': ['openxml', '<w:p/>']}]})

    def test_design_classes_map_to_native_components_without_rewriting_text(self):
        ast = {'blocks': [
            {'t': 'Table', 'c': [['summary', ['metric-grid'], []], {}, [], {}, [], {}]},
            {'t': 'Header', 'c': [1, ['', ['chapter'], []], [{'t': 'Str', 'c': 'Details'}]]},
        ]}
        apply_design_components(ast)
        self.assertIn(['custom-style', 'MetricGrid'], ast['blocks'][0]['c'][0][2])
        self.assertEqual(ast['blocks'][1]['t'], 'RawBlock')
        self.assertIn('w:type="page"', ast['blocks'][1]['c'][1])
        self.assertEqual(ast['blocks'][2]['t'], 'Header')
        self.assertEqual(inline_text(ast['blocks'][2]['c'][2]), 'Details')

    def test_math_source_is_not_required_as_literal_pdf_text(self):
        ast = {'blocks': [
            {'t': 'Header', 'c': [1, ['', [], []], [{'t': 'Str', 'c': 'Comparing'},
                                                  {'t': 'Math', 'c': [{'t': 'InlineMath'}, r'\alpha']}]]},
            {'t': 'Para', 'c': [{'t': 'Str', 'c': 'The value'},
                               {'t': 'Math', 'c': [{'t': 'InlineMath'}, 'x^2']}]},
        ]}
        self.assertEqual(source_expectations(ast), {'required_text': [], 'same_page': []})


@unittest.skipUnless(shutil.which('soffice'), 'LibreOffice Writer is required')
class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        cls.exports = {}
        for fixture in ('mixed-zh', 'english', 'rtl'):
            cls.exports[fixture] = render(FIXTURES / (fixture + '.md'), cls.root / fixture)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_multilingual_content_survives_real_export(self):
        for name, expected in [('mixed-zh', ['2026-09-21', '1,234.50', '阅读与执行']),
                               ('english', ['ACME-42', '1,234.50 USD', 'Comparative values']),
                               ('rtl', ['ACME-42', '1234.50'])]:
            with self.subTest(name=name):
                artifact = self.exports[name]
                with pymupdf.open(artifact['outputs']['pdf']['path']) as pdf:
                    text = ''.join(page.get_text() for page in pdf)
                    self.assertGreater(len(pdf), 0)
                    for phrase in expected:
                        self.assertIn(re.sub(r'\s+', '', phrase), re.sub(r'\s+', '', text))
                    self.assertNotIn('\ufffd', text)
                self.assertEqual(artifact['status'], 'needs-inspection')
                for output in artifact['outputs'].values():
                    self.assertEqual(hashlib.sha256(Path(output['path']).read_bytes()).hexdigest(), output['sha256'])

    def test_default_export_has_no_missing_styles(self):
        for artifact in self.exports.values():
            validate_docx(artifact['outputs']['docx']['path'])

    def test_editorial_components_render_as_word_styles_and_real_pages(self):
        source = self.root / 'editorial.md'
        source.write_text(
            '---\nlang: zh-CN\ntitle: 版式样例\nsubtitle: 同一份内容，更清楚的层级\n---\n\n'
            '::: {custom-style="Deck"}\n这是一段用于建立阅读节奏的导语。\n:::\n\n'
            '| 88.3% | 20.2× | 121× |\n|---:|---:|---:|\n'
            '| 内部运行占比 | 机器放大倍数差 | 模型成本差 |\n\n'
            ': {#summary .metric-grid}\n\n'
            '::: {custom-style="Key Takeaway"}\n工具不会自动创造构图；设计参照和逐页复核才会。\n:::\n\n'
            '# 01 详细内容 {.chapter}\n\n正文从新页开始，并保持为可编辑文本。\n',
            encoding='utf-8')
        result = render(source, self.root / 'editorial-out')
        document = docx.Document(result['outputs']['docx']['path'])
        self.assertEqual(document._element.find(W + 'background').get(W + 'color'), 'F5F4ED')
        self.assertEqual(document.tables[0]._tbl.tblPr.tblStyle.val, 'MetricGrid')
        self.assertEqual(document.tables[0].rows[0].cells[0].paragraphs[0].style.name, 'Metric Value')
        self.assertTrue(any(p.style.name == 'Deck' for p in document.paragraphs))
        self.assertTrue(any(p.style.name == 'Key Takeaway' for p in document.paragraphs))
        validate_docx(result['outputs']['docx']['path'])
        with pymupdf.open(result['outputs']['pdf']['path']) as pdf:
            self.assertGreaterEqual(len(pdf), 2)
            self.assertNotIn('详细内容', pdf[0].get_text())
            self.assertIn('详细内容', pdf[1].get_text())

    def test_long_table_keeps_rows_and_repeats_header(self):
        source = self.root / 'long-table.md'
        rows = [f'| Record {i:03} | {i * 17} | A longer explanation for this record, which should wrap within its column. |'
                for i in range(1, 61)]
        source.write_text('# A table spanning pages\n\nThe values below must survive pagination.\n\n'
                          '| Record | Amount | Explanation |\n|--------|-------:|-------------|\n' + '\n'.join(rows), encoding='utf-8')
        result = render(source, self.root / 'long-table-out')
        with pymupdf.open(result['outputs']['pdf']['path']) as pdf:
            text = ''.join(page.get_text() for page in pdf)
            self.assertGreater(len(pdf), 1)
            for number in range(1, 61):
                self.assertIn(f'Record {number:03}', text)
            self.assertGreaterEqual(text.count('Explanation'), len(pdf))

    def test_editable_source_is_copied_byte_for_byte(self):
        source = Path(self.exports['english']['outputs']['docx']['path'])
        before = source.read_bytes()
        result = render(source, self.root / 'existing', output_format='docx')
        self.assertEqual(result['theme'], 'source-preserved')
        self.assertEqual(result['deliver'], ['docx'])
        self.assertEqual(Path(result['outputs']['docx']['path']).read_bytes(), before)
        self.assertEqual(source.read_bytes(), before)

    def test_reference_preserves_page_setup_header_and_body_size(self):
        reference = self.root / 'custom.docx'
        build_reference(find_pandoc(), reference, 'en-US')
        theme = docx.Document(reference)
        theme.sections[0].page_width = docx.shared.Inches(8.5)
        theme.sections[0].page_height = docx.shared.Inches(11)
        theme.sections[0].header.paragraphs[0].text = 'A supplied running header'
        theme.styles['Body Text'].font.size = docx.shared.Pt(14)
        theme.save(reference)
        before = reference.read_bytes()
        result = render(FIXTURES / 'english.md', self.root / 'reference', reference=reference)
        rendered = docx.Document(result['outputs']['docx']['path'])
        self.assertEqual(result['theme'], 'reference-preserved')
        self.assertEqual(reference.read_bytes(), before)
        self.assertEqual(rendered.styles['Body Text'].font.size.pt, 14)
        self.assertIn('A supplied running header', rendered.sections[0].header.paragraphs[0].text)
        with pymupdf.open(result['outputs']['pdf']['path']) as pdf:
            self.assertAlmostEqual(pdf[0].rect.width, 612, delta=1)
            self.assertAlmostEqual(pdf[0].rect.height, 792, delta=1)

    def test_failed_export_cannot_claim_old_pdf_as_new_output(self):
        destination = self.root / 'failed'
        destination.mkdir()
        old_pdf = destination / 'english.pdf'
        old_pdf.write_bytes(b'old output')
        with patch('render_document.export_pdf', side_effect=RuntimeError('Writer unavailable')):
            with self.assertRaisesRegex(RuntimeError, 'Writer unavailable'):
                render(FIXTURES / 'english.md', destination)
        self.assertEqual(old_pdf.read_bytes(), b'old output')
        self.assertEqual(json.loads((destination / 'render.json').read_text())['status'], 'failed')

    def test_missing_image_is_an_export_error(self):
        source = self.root / 'missing-image.md'
        source.write_text('# Figure\n\n![Required diagram](missing-plot.png)\n', encoding='utf-8')
        with self.assertRaisesRegex(RuntimeError, 'Could not fetch resource'):
            render(source, self.root / 'missing-image')
        self.assertFalse((self.root / 'missing-image' / 'missing-image.pdf').exists())

    def test_source_without_styles_part_can_still_be_exported(self):
        source = self.root / 'unstyled.docx'
        with zipfile.ZipFile(source, 'w') as archive:
            archive.writestr('[Content_Types].xml',
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                '</Types>')
            archive.writestr('_rels/.rels',
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                '</Relationships>')
            archive.writestr('word/document.xml',
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body><w:p><w:r><w:t>Preserved simple source</w:t></w:r></w:p></w:body></w:document>')
        validate_docx(source)
        result = render(source, self.root / 'unstyled')
        self.assertEqual(Path(result['outputs']['docx']['path']).read_bytes(), source.read_bytes())
        with pymupdf.open(result['outputs']['pdf']['path']) as pdf:
            self.assertIn('Preserved simple source', pdf[0].get_text())

    def test_source_docx_is_never_overwritten(self):
        source = Path(self.exports['english']['outputs']['docx']['path'])
        with self.assertRaisesRegex(ValueError, 'overwrite an input'):
            render(source, source.parent)

    def test_output_aliases_cannot_overwrite_the_source(self):
        source = Path(self.exports['english']['outputs']['docx']['path'])
        original = source.read_bytes()
        for name in ('english.pdf', 'expectations.json', 'render.json'):
            with self.subTest(name=name):
                destination = self.root / ('alias-' + name)
                destination.mkdir()
                (destination / name).symlink_to(source)
                with self.assertRaisesRegex(ValueError, 'overwrite an input'):
                    render(source, destination)
                self.assertEqual(source.read_bytes(), original)

    def test_dangling_style_is_rejected(self):
        source = Path(self.exports['english']['outputs']['docx']['path'])
        broken = self.root / 'broken.docx'
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(broken, 'w') as target:
            for item in original.infolist():
                content = original.read(item.filename)
                if item.filename == 'word/document.xml':
                    root = ET.fromstring(content)
                    next(root.iter(W + 'pStyle')).set(W + 'val', 'AbsentStyle')
                    content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                target.writestr(item, content)
        with self.assertRaisesRegex(RuntimeError, 'AbsentStyle'):
            validate_docx(broken)


if __name__ == '__main__':
    unittest.main()
